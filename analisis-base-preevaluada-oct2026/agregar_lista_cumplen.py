# -*- coding: utf-8 -*-
"""Agrega las hojas de asociados que cumplen los criterios verificables del memorando.

Uso: python3 -I agregar_lista_cumplen.py <libro_recalculado.xlsx>
Crea Cumplen_Memorando (todos), Cumplen_Hipotecario, Cumplen_Vehiculo, Cumplen_TC y Consumo.
La lista de ASOCIADO se toma del libro ya recalculado; el resto de columnas son formulas INDEX/MATCH
contra BD_CLASIFICADA, y la columna de control verifica que cada asociado siga cumpliendo.
Despues de ejecutarlo hay que recalcular el libro otra vez.
"""
import sys
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
CUMPLE = 'Cumple criterios verificables'
x = pd.read_excel(OUT, sheet_name='BD_CLASIFICADA', dtype={'ASOCIADO': str})
num = lambda c: pd.to_numeric(x[c], errors='coerce')
ok = x.ELEGIBILIDAD_GENERAL == CUMPLE
ok_tc = x.ELEGIBILIDAD_TC == CUMPLE
sel = lambda m: sorted(x.loc[m, 'ASOCIADO'].astype(str), key=int)

wb = openpyxl.load_workbook(OUT)
bd = wb['BD_CLASIFICADA']
hdr = {bd.cell(1, j).value: L(j) for j in range(1, bd.max_column + 1)}
last = bd.max_row
KEY = f"BD_CLASIFICADA!$A$2:$A${last}"
RNG = lambda c: f"BD_CLASIFICADA!${hdr[c]}$2:${hdr[c]}${last}"

F = 'Arial'
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
USD, INT = '"$"#,##0.00', '#,##0'
LEVEL_FILLS = {'Alto': 'F8CBAD', 'Medio': 'FFE699', 'Bajo': 'C6E0B4'}

# bloques de columnas: (titulo del bloque, color, [(encabezado, columna BD_CLASIFICADA, formato, ancho)])
B_ASOC = ('ASOCIADO', '1F3864', [('NOMBRE', 'NOMBRE', None, 36), ('DUI', 'DUI', '@', 11), ('AGENCIA', 'AGENCIA', None, 16)])
B_FIN = ('SITUACIÓN FINANCIERA', '2E75B6', [('INGRESO MENSUAL', 'INGRESO', USD, 13),
                                            ('FONDO COMPENSATORIO', 'FONDO_COMPENSATORIO', USD, 14),
                                            ('DEUDA SSF', 'SALDO_SSF', USD, 14)])
B_NIV = ('NIVELES DE RIESGO', 'C55A11', [('ANTIGÜEDAD', 'NIVEL_ANTIGUEDAD', None, 11), ('INGRESO', 'NIVEL_INGRESO', None, 11),
                                         ('DEUDA', 'NIVEL_DEUDA', None, 11), ('DEUDA / INGRESO', 'NIVEL_DEUDA_INGRESO', None, 11)])
B_OBS = ('OBSERVACIONES', '7F7F7F', [('ALERTAS DE RIESGO', 'TOTAL_ALERTAS_RIESGO', '0', 9),
                                     ('CONDICIONES DEL MEMORANDO', 'CONDICIONES_MEMO', None, 46)])
M_HIPO = ('HIPOTECARIO', 'PRECA HIPO', INT, 13)
M_VEH = ('VEHÍCULO', 'PRECA VEHICULO', INT, 12)
M_TC = ('TARJETA', 'PRECA TC', INT, 11)
B_MONTO = lambda cols: ('MONTO PRE EVALUADO' if len(cols) > 1 else 'MONTO', '548235', cols)
# columnas que pueden venir vacias o "N/A" en la base: se muestran en blanco
BLANKABLE = {'PRECA HIPO', 'PRECA VEHICULO', 'PRECA TC', 'SEGMENTO', 'SALDO_SSF', 'CONDICIONES_MEMO', 'TC_LIMITE_SIN_SOLVENCIA'}


def make_sheet(name, pos, ids, title, note, elig_col, count_formula, blocks, product=None):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name, pos)
    ws.sheet_view.showGridLines = False
    cols = [c for _, _, cs in blocks for c in cs]
    n, end = len(ids), 6 + len(ids)
    flag = L(len(cols) + 2)
    ws['A1'] = f'{title}  —  {n} asociados'
    ws['A1'].font = Font(name=F, size=14, bold=True, color='1F3864')
    ws['A2'] = note
    ws['A2'].font = Font(name=F, size=9, italic=True, color='C00000')
    link = ws['A3']
    link.value = '◄ INICIO'
    link.hyperlink = "#'INICIO'!A1"
    link.font = Font(name=F, size=10, bold=True, color='0563C1', underline='single')
    ws['B3'] = (f'=IF(AND(COUNTA(A7:A{end})={count_formula},COUNTIF({flag}7:{flag}{end},"Sí")=COUNTA(A7:A{end})),'
                f'"Control: OK, "&COUNTA(A7:A{end})&" asociados y todos siguen cumpliendo",'
                f'"Control: REVISAR, la base cambió; regenerar la lista")')
    ws['B3'].font = Font(name=F, size=9, bold=True, color='006100')
    # fila 5: bloques de color; fila 6: encabezados
    j = 2
    ws.cell(5, 1).fill = PatternFill('solid', fgColor='1F3864')
    for btitle, color, bcols in blocks:
        c = ws.cell(5, j, btitle)
        c.font = Font(name=F, size=9, bold=True, color='FFFFFF')
        c.alignment = Alignment(horizontal='center')
        for k in range(j, j + len(bcols)):
            ws.cell(5, k).fill = PatternFill('solid', fgColor=color)
            ws.cell(6, k).fill = PatternFill('solid', fgColor=color)
        if len(bcols) > 1:
            ws.merge_cells(start_row=5, start_column=j, end_row=5, end_column=j + len(bcols) - 1)
        j += len(bcols)
    ws.cell(5, j).fill = PatternFill('solid', fgColor='404040')
    heads = ['N° ASOCIADO'] + [h for h, *_ in cols] + ['¿SIGUE CUMPLIENDO?']
    for k, h in enumerate(heads, 1):
        c = ws.cell(6, k, h)
        c.font = Font(name=F, size=9, bold=True, color='FFFFFF')
        if k == 1:
            c.fill = PatternFill('solid', fgColor='1F3864')
        if k == len(heads):
            c.fill = PatternFill('solid', fgColor='404040')
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = box
    ws.row_dimensions[6].height = 32
    elig = RNG(elig_col)
    stripe = PatternFill('solid', fgColor='F7F9FC')
    for i, a in enumerate(ids):
        r = 7 + i
        c = ws.cell(r, 1, a)
        c.font = Font(name=F, size=9, bold=True)
        c.number_format = '@'
        c.border = box
        for k, (h, src, fmt, w) in enumerate(cols, 2):
            look = f'INDEX({RNG(src)},MATCH($A{r},{KEY},0))'
            f = f'={look}'
            if src in BLANKABLE:
                f = f'=IF(ISNUMBER({look})+(LEN({look}&"")>0)*({look}<>"N/A"),{look},"")'
            c = ws.cell(r, k, f)
            c.font = Font(name=F, size=9)
            c.border = box
            if fmt:
                c.number_format = fmt
            if src.startswith('NIVEL_'):
                c.alignment = Alignment(horizontal='center')
            if i % 2:
                c.fill = stripe
        cond = f'INDEX({elig},MATCH($A{r},{KEY},0))="{CUMPLE}"'
        if product:
            cond = f'AND({cond},ISNUMBER(INDEX({RNG(product)},MATCH($A{r},{KEY},0))))'
        c = ws.cell(r, len(cols) + 2, f'=IF({cond},"Sí","No")')
        c.font = Font(name=F, size=9)
        c.alignment = Alignment(horizontal='center')
        c.border = box
    # colores de nivel
    for k, (h, src, fmt, w) in enumerate(cols, 2):
        if src.startswith('NIVEL_'):
            rng = f'{L(k)}7:{L(k)}{end}'
            for lvl, color in LEVEL_FILLS.items():
                ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=[f'"{lvl}"'],
                                                              fill=PatternFill('solid', fgColor=color, bgColor=color)))
    ws.column_dimensions['A'].width = 11
    for k, (h, src, fmt, w) in enumerate(cols, 2):
        ws.column_dimensions[L(k)].width = w
    ws.column_dimensions[flag].width = 11
    ws.freeze_panes = 'C7'
    ws.auto_filter.ref = f'A6:{flag}{end}'
    # resumen de montos de los productos que aparecen en la hoja
    r = end + 2
    ws.cell(r, 1, 'RESUMEN').font = Font(name=F, size=11, bold=True, color='1F3864')
    pos_of = {src: k for k, (h, src, fmt, w) in enumerate(cols, 2)}
    rr = r + 1
    for lab, src in [('Hipotecario', 'PRECA HIPO'), ('Vehículo', 'PRECA VEHICULO'), ('Tarjeta de crédito', 'PRECA TC')]:
        if src not in pos_of:
            continue
        col = L(pos_of[src])
        ws.cell(rr, 2, lab).font = Font(name=F, size=10, bold=True)
        ws.cell(rr, 3, f'=COUNT({col}7:{col}{end})&" asociados"').font = Font(name=F, size=10)
        c = ws.cell(rr, 4, f'=SUM({col}7:{col}{end})')
        c.number_format = '"$"#,##0'
        c.font = Font(name=F, size=10, bold=True)
        rr += 1
    return ws


EXT = ('Falta la consulta externa de lista negra, créditos castigados, personas relacionadas o empleados '
       'y solvencia en Hacienda. No es una aprobación de crédito.')
make_sheet('Cumplen_Memorando', 2, sel(ok), 'CUMPLEN TODOS LOS CRITERIOS DEL MEMORANDO',
           'Cumplen los criterios 1 a 7 que se pueden verificar con la base. ' + EXT, 'ELEGIBILIDAD_GENERAL',
           f'COUNTIF({RNG("ELEGIBILIDAD_GENERAL")},"{CUMPLE}")',
           [B_ASOC, B_FIN, B_MONTO([M_HIPO, M_VEH, M_TC]), B_NIV, B_OBS])
make_sheet('Cumplen_Hipotecario', 3, sel(ok & num('PRECA HIPO').notna()), 'CRÉDITO HIPOTECARIO',
           'Cumplen los criterios y tienen monto hipotecario. Todos deben verificar solvencia en Hacienda (monto > $30,000). ' + EXT,
           'ELEGIBILIDAD_GENERAL', f'COUNTIFS({RNG("ELEGIBILIDAD_GENERAL")},"{CUMPLE}",{RNG("PRECA HIPO")},">0")',
           [B_ASOC, B_FIN, B_MONTO([M_HIPO]), B_NIV, B_OBS], 'PRECA HIPO')
make_sheet('Cumplen_Vehiculo', 4, sel(ok & num('PRECA VEHICULO').notna()), 'CRÉDITO DE VEHÍCULO',
           'Cumplen los criterios y tienen monto de vehículo (tabla RANGOS VEHICULO: ingreso ≥ $3,000). ' + EXT,
           'ELEGIBILIDAD_GENERAL', f'COUNTIFS({RNG("ELEGIBILIDAD_GENERAL")},"{CUMPLE}",{RNG("PRECA VEHICULO")},">0")',
           [B_ASOC, B_FIN, B_MONTO([M_VEH]), B_NIV, B_OBS], 'PRECA VEHICULO')
make_sheet('Cumplen_TC', 5, sel(ok_tc & num('PRECA TC').notna()), 'TARJETA DE CRÉDITO',
           'Cumplen los criterios, no tienen tarjeta y tienen límite pre evaluado. Límite > $4,000 sin solvencia en Hacienda: '
           'se otorga con $4,000. Si no está al día en aportaciones debe ponerse al día. ' + EXT,
           'ELEGIBILIDAD_TC', f'COUNTIFS({RNG("ELEGIBILIDAD_TC")},"{CUMPLE}",{RNG("PRECA TC")},">0")',
           [B_ASOC, B_FIN,
            B_MONTO([M_TC, ('SEGMENTO', 'SEGMENTO', None, 20), ('LÍMITE SI NO ES SOLVENTE', 'TC_LIMITE_SIN_SOLVENCIA', INT, 12),
                     ('APORTACIONES AL DÍA', 'APORTACIONES_AL_DIA', None, 11)]),
            B_NIV, B_OBS], 'PRECA TC')

# ---- Consumo: no existe en la base ni en el memorando
if 'Consumo' in wb.sheetnames:
    del wb['Consumo']
wc = wb.create_sheet('Consumo', 6)
wc.sheet_view.showGridLines = False
wc.column_dimensions['A'].width = 36
wc.column_dimensions['B'].width = 100
wc['A1'] = 'CRÉDITO DE CONSUMO: NO SE PUEDE EVALUAR CON ESTA BASE'
wc['A1'].font = Font(name=F, size=14, bold=True, color='C00000')
wc['A2'] = '◄ INICIO'
wc['A2'].hyperlink = "#'INICIO'!A1"
wc['A2'].font = Font(name=F, size=10, bold=True, color='0563C1', underline='single')
rows = [
    ('Situación', 'El memorando solo pide validar Crédito Hipotecario, Vehículo y Tarjeta de Crédito. La base no tiene columna '
                  'de monto pre evaluado de consumo (solo PRECA HIPO, PRECA VEHICULO y PRECA TC) ni tabla de montos de consumo.'),
    ('Qué NO se hizo', 'No se asignaron montos de consumo: inventar una tabla o un monto sería presentar como dato algo que la '
                       'institución no ha definido.'),
    ('Universo que sí cumple los criterios generales',
     f'=COUNTIF({RNG("ELEGIBILIDAD_GENERAL")},"{CUMPLE}")&" asociados (hoja Cumplen_Memorando) cumplen los criterios verificables y '
     f'podrían ser el universo de partida para consumo, si Riesgos lo aprueba."'),
    ('Para evaluar consumo se necesita', '1) Tabla institucional de montos de consumo por ingreso (y/o fondo). '
                                          '2) Plazo y tasa para calcular la cuota. '
                                          '3) Cuotas actuales de todas las obligaciones para medir la relación cuota-ingreso. '
                                          '4) Límite institucional de cuota/ingreso para consumo.'),
    ('Siguiente paso', 'Con esa tabla se agrega una columna PRECA CONSUMO a la base y se genera la hoja Cumplen_Consumo con el mismo '
                       'formato que las de hipotecario, vehículo y tarjeta.'),
]
for i, (a, b) in enumerate(rows, 4):
    c = wc.cell(i, 1, a)
    c.font = Font(name=F, size=10, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='548235')
    c.alignment = Alignment(wrap_text=True, vertical='center')
    c.border = box
    c = wc.cell(i, 2, b)
    c.font = Font(name=F, size=10)
    c.alignment = Alignment(wrap_text=True, vertical='center')
    c.border = box
    wc.row_dimensions[i].height = 45

wb.save(OUT)
print({k: len(v) for k, v in [('todos', sel(ok)), ('hipo', sel(ok & num('PRECA HIPO').notna())),
                              ('veh', sel(ok & num('PRECA VEHICULO').notna())), ('tc', sel(ok_tc & num('PRECA TC').notna()))]})
