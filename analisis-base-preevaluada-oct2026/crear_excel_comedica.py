# -*- coding: utf-8 -*-
"""Genera Recomendacion_Base_PreEvaluada_COMEDICA_CORREGIDA.xlsx (libro con fórmulas, automatizado).

Uso: python3 -I crear_excel_comedica.py <archivo_auditado.xlsx> <salida.xlsx>

Toma la hoja BD FINAL del archivo auditado, la copia sin cambios en 6_Base_Original y construye:
  7_Calculo                 motor: una fila por registro, todo con fórmulas que leen 6_Base_Original y los
                            parámetros de 2_Criterios_y_Rangos (si se pega una base nueva, todo se recalcula)
  1_Resumen_Ejecutivo       conteos y montos con COUNTIFS / SUMIFS sobre 7_Calculo
  2_Criterios_y_Rangos      parámetros (celdas con nombre), rangos, pesos y estado de validación
  3_Recomendacion_Preliminar y 4_Casos_por_Revisar
                            listas dinámicas con INDEX / MATCH / SMALL (compatibles con Excel 2010+)
  5_Distribucion_Controles  distribuciones, alternativas de umbral y pruebas de integridad (OK / REVISAR)
  8_Hallazgos, 9_Registro_Cambios
No usa macros ni funciones de matriz dinámica (FILTER, SORT, XLOOKUP).
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: F401  (reservado)

SRC, OUT = sys.argv[1:3]
CAP = 2600          # filas que soporta el motor (la base actual tiene 2,262)
LCAP = 2600         # filas que soportan las listas dinámicas
BASE, CALC = '6_Base_Original', '7_Calculo'
E1 = 'Cumple criterios iniciales'
E2, E3 = '2. Recomendado preliminarmente', '3. Requiere revisión adicional'
E4, E5 = '4. Pendiente por información insuficiente', '5. No elegible (restricción del memorando)'

F = 'Arial'
NAVY, BLUE, LBLUE, WHITE, GRAY = '1F3864', '2E75B6', 'DDEBF7', 'FFFFFF', '595959'
GREEN_F, YEL_F, RED_F, GRAY_F = 'C6E0B4', 'FFE699', 'F8CBAD', 'EDEDED'
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = lambda c: PatternFill('solid', fgColor=c)
font = lambda size=10, bold=False, color='000000', italic=False: Font(name=F, size=size, bold=bold, color=color, italic=italic)
USD, USD2, PCT, INT = '$#,##0', '$#,##0.00', '0.0%', '#,##0'

src = openpyxl.load_workbook(SRC)['BD FINAL']
NREG = src.max_row - 1
wb = openpyxl.Workbook()
wb.remove(wb.active)


def head(ws, row, headers, widths=None, color=NAVY, height=34):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row, j, h)
        c.font = font(10, True, WHITE)
        c.fill = fill(color)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = box
        if widths:
            ws.column_dimensions[L(j)].width = widths[j - 1]
    ws.row_dimensions[row].height = height


def title(ws, text, sub=None):
    ws['A1'] = text
    ws['A1'].font = font(16, True, NAVY)
    if sub:
        ws['A2'] = sub
        ws['A2'].font = font(10, italic=True, color=GRAY)
    ws.sheet_view.showGridLines = False


def put(ws, cell, v, fmt=None, **k):
    c = ws[cell]
    c.value = v
    c.font = font(**{kk: vv for kk, vv in k.items() if kk in ('size', 'bold', 'color', 'italic')})
    if fmt:
        c.number_format = fmt
    if 'bg' in k:
        c.fill = fill(k['bg'])
    c.alignment = Alignment(vertical='center', wrap_text=k.get('wrap', False), horizontal=k.get('h'))
    if k.get('border'):
        c.border = box
    return c


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


# ============================================================== 6_Base_Original (copia exacta)
wbase = wb.create_sheet(BASE)
for row in src.iter_rows():
    for c in row:
        if c.value is not None:
            d = wbase.cell(c.row, c.column, c.value)
            d.number_format = c.number_format
            d.font = font(9, c.row == 1, WHITE if c.row == 1 else '000000')
            if c.row == 1:
                d.fill = fill(NAVY)
for j in range(1, src.max_column + 1):
    wbase.column_dimensions[L(j)].width = 14
wbase.freeze_panes = 'C2'
wbase.auto_filter.ref = f'A1:{L(src.max_column)}{src.max_row}'
hdr = [c.value for c in src[1]]
assert hdr[:31] == ['ASOCIADO', 'NOMBRE', 'EDAD', 'SEXO', 'CLI_FECHA_NACIMIENTO', 'ESTADO', 'DESCRIPCION', 'ESPECIALIDAD', 'DUI',
                    'CONCATENADO', 'CORREO', 'TELEFONO_CELULAR', 'DIRECCION', 'LUGAR_DE_TRABAJO', 'ANTIGUEDAD_ANIOS', 'INGRESO',
                    'FONDO_COMPENSATORIO', 'PRECA HIPO', 'PRECA VEHICULO', 'PRECA TC', 'SEGMENTO', 'CUOTAS_VENC_APORT', 'AGENCIA',
                    'TIENE_CREDITO_VIGENTE', 'MAX_CAT_RIESGO', 'APORTACIONES_AL_DIA', 'MONTO_TOTAL_APORTACIONES', 'TIENE_TARJETA',
                    'CALIFICACION_SSF', 'SALDO_SSF', 'INSTITUCION'], hdr
B = {k: L(hdr.index(h) + 1) for k, h in [('ASOC', 'ASOCIADO'), ('NOM', 'NOMBRE'), ('EDAD', 'EDAD'), ('FNAC', 'CLI_FECHA_NACIMIENTO'),
                                         ('DUI', 'DUI'), ('ANT', 'ANTIGUEDAD_ANIOS'), ('ING', 'INGRESO'), ('FON', 'FONDO_COMPENSATORIO'),
                                         ('HIPO', 'PRECA HIPO'), ('VEH', 'PRECA VEHICULO'), ('TC', 'PRECA TC'), ('CVA', 'CUOTAS_VENC_APORT'),
                                         ('AGE', 'AGENCIA'), ('CATI', 'MAX_CAT_RIESGO'), ('APD', 'APORTACIONES_AL_DIA'),
                                         ('TTC', 'TIENE_TARJETA'), ('CATE', 'CALIFICACION_SSF'), ('SSF', 'SALDO_SSF')]}

# ============================================================== 2_Criterios_y_Rangos (parámetros con nombre)
wcr = wb.create_sheet('2_Criterios_y_Rangos')
title(wcr, 'CRITERIOS, RANGOS Y PARÁMETROS', 'Las celdas AZULES son parámetros: si se cambian, todo el libro se recalcula. '
      'Estado: "Memorando" = documentado; "Propuesto" = propuesta de Riesgo pendiente de aprobación.')
PARAMS = [  # código, descripción, valor, formato, fuente, estado
    ('P_FECHA_CORTE', 'Fecha de corte de la base', '=DATE(2026,9,30)', 'dd/mm/yyyy', 'Memorando 27/10/2026', 'Memorando'),
    ('P_EDAD_MIN', 'Edad mínima (años cumplidos al corte)', 30, '0', 'Memorando, criterio 7', 'Memorando'),
    ('P_EDAD_MAX', 'Edad máxima (años cumplidos al corte)', 55, '0', 'Memorando, criterio 7', 'Memorando'),
    ('P_ANT_ALTO_MAX', 'Antigüedad: nivel Alto hasta (años, incluido)', 1, '0', 'Hoja "Variables" de la base (0-1 Alto)', 'Base del usuario'),
    ('P_ANT_BAJO_MIN', 'Antigüedad: nivel Bajo desde (años, incluido)', 5, '0', 'Hoja "Variables" de la base (5+ Bajo)', 'Base del usuario'),
    ('P_ING_L1', 'Ingreso: límite Alto/Medio (USD/mes)', 1500, USD, 'Propuesta (corte de la tabla de tarjeta Dorada)', 'Propuesto'),
    ('P_ING_L2', 'Ingreso: límite Medio/Bajo (USD/mes)', 3000, USD, 'Propuesta (corte "NO APLICA" de la tabla de vehículo)', 'Propuesto'),
    ('P_ING_ATIP', 'Ingreso atípico a confirmar (USD/mes)', 12172.65, USD2, 'Estadístico de la base (P75 + 3 × RIC)', 'Propuesto'),
    ('P_DEU_L1', 'Deuda SSF: límite Bajo/Medio (USD)', 1000, USD, 'Propuesta (distribución de la base)', 'Propuesto'),
    ('P_DEU_L2', 'Deuda SSF: límite Medio/Alto (USD)', 10000, USD, 'Propuesta (tercio de mayor deuda)', 'Propuesto'),
    ('P_DTI_L1', 'Deuda ÷ ingreso anual: límite Bajo/Medio', 0.25, PCT, 'Propuesta (= 3 meses de ingreso)', 'Propuesto'),
    ('P_DTI_L2', 'Deuda ÷ ingreso anual: límite Medio/Alto', 1.0, PCT, 'Propuesta (= 12 meses de ingreso)', 'Propuesto'),
    ('P_MORA_L1', 'Cuotas de aportación vencidas: Medio desde', 1, '0', 'Propuesta', 'Propuesto'),
    ('P_MORA_L2', 'Cuotas de aportación vencidas: Alto desde', 3, '0', 'Propuesta', 'Propuesto'),
    ('P_W_ANT', 'Peso antigüedad (puntaje piloto)', 0.13, '0%', 'Indicado por el usuario', 'Piloto'),
    ('P_W_ING', 'Peso ingreso (puntaje piloto)', 0.30, '0%', 'Indicado por el usuario', 'Piloto'),
    ('P_W_DEU', 'Peso deuda SSF (puntaje piloto)', 0.20, '0%', 'Indicado por el usuario', 'Piloto'),
    ('P_W_MORA', 'Peso cuotas de aportación vencidas (puntaje piloto)', 0.12, '0%', 'Indicado por el usuario', 'Piloto'),
    ('P_W_CATE', 'Peso categoría SSF (puntaje piloto)', 0.25, '0%', 'Indicado por el usuario', 'Piloto'),
    ('P_HIPO_SOLV', 'Hipotecario: solvencia en Hacienda si el monto supera (USD)', 30000, USD, 'Memorando, otros criterios 6', 'Memorando'),
    ('P_VEH_SOLV', 'Vehículo: solvencia en Hacienda si el monto supera (USD)', 25000, USD, 'Propuesta del usuario', 'Propuesto'),
    ('P_TC_TOPE', 'Tarjeta sin solvencia en Hacienda: límite máximo (USD)', 4000, USD, 'Memorando, otros criterios 5', 'Memorando'),
]
r = 4
head(wcr, r, ['Código', 'Parámetro', 'Valor', 'Fuente', 'Estado de validación'], [16, 52, 14, 46, 20])
P = {}
for code, desc, val, fmt, fuente, est in PARAMS:
    r += 1
    for j, v in enumerate([code, desc, val, fuente, est], 1):
        c = wcr.cell(r, j, v)
        c.font = font(10, j == 3, '0000FF' if j == 3 else '000000')
        c.border = box
        if j == 3:
            c.number_format = fmt
            c.fill = fill(LBLUE)
    name(code, f"'2_Criterios_y_Rangos'!$C${r}")
    P[code] = f"'2_Criterios_y_Rangos'!$C${r}"
r_w_total = r + 1
put(wcr, f'B{r_w_total}', 'Suma de pesos (debe ser 100%)', bold=True)
put(wcr, f'C{r_w_total}', '=P_W_ANT+P_W_ING+P_W_DEU+P_W_MORA+P_W_CATE', '0%', bold=True, border=True)

# tablas de producto (copiadas de la base del usuario: hojas RANGOS VEHICULO y RANGOS TC)
r = r_w_total + 3
put(wcr, f'A{r}', 'TABLA DE VEHÍCULO (hoja "RANGOS VEHICULO" de la base, tal cual)', size=12, bold=True, color=NAVY)
r += 1
head(wcr, r, ['Ingreso desde (USD/mes)', 'Ingreso hasta', 'Monto según tabla'], None, BLUE, 22)
veh_rows = [(0, 2999.99, 'NO APLICA'), (3000, 4999.99, 20000), (5000, 9999.99, 25000), (10000, 14999.99, 30000), (15000, None, 40000)]
v0 = r + 1
for lo, hi, m in veh_rows:
    r += 1
    for j, v in enumerate([lo, hi if hi is not None else 'o más', m], 1):
        c = wcr.cell(r, j, v)
        c.border = box
        c.font = font(10)
        c.number_format = USD2 if j < 3 else USD
name('VEH_ING_MIN', f"'2_Criterios_y_Rangos'!$A${v0}:$A${r}")
name('VEH_MONTO', f"'2_Criterios_y_Rangos'!$C${v0}:$C${r}")
r += 2
put(wcr, f'A{r}', 'TABLA DE TARJETA (hoja "RANGOS TC" de la base, tal cual)', size=12, bold=True, color=NAVY)
r += 1
head(wcr, r, ['Ingreso desde (USD/mes)', 'Ingreso hasta', 'Límite', 'Segmento'], None, BLUE, 22)
tc_rows = [(500, 749.99, 500, 'CLASICA LOCAL'), (750, 999.99, 1500, 'CLASICA INTERNACIONAL'), (1000, 1499.99, 2000, 'CLASICA INTERNACIONAL'),
           (1500, 1699.99, 2500, 'DORADA'), (1700, 1999.99, 3000, 'DORADA'), (2000, 2499.99, 3500, 'PLATINUM'),
           (2500, 2999.99, 4500, 'PLATINUM'), (3000, 3499.99, 5000, 'PLATINUM'), (3500, 7499.99, 7000, 'BLACK'),
           (7500, 7999.99, 8000, 'BLACK'), (8000, None, 10000, 'BLACK')]
t0 = r + 1
for lo, hi, m, seg in tc_rows:
    r += 1
    for j, v in enumerate([lo, hi if hi is not None else 'o más', m, seg], 1):
        c = wcr.cell(r, j, v)
        c.border = box
        c.font = font(10)
        if j < 4:
            c.number_format = USD2 if j < 3 else USD
name('TC_ING_MIN', f"'2_Criterios_y_Rangos'!$A${t0}:$A${r}")
name('TC_LIM', f"'2_Criterios_y_Rangos'!$C${t0}:$C${r}")
put(wcr, f'A{r + 1}', 'Hipotecario: la base no trae tabla; el monto se toma como viene en PRECA HIPO (N/A = no aplica). Pendiente: tabla oficial.',
    italic=True, color=GRAY)

# tabla de criterios (Variable | Indicador | Rango mínimo | Rango máximo | Nivel | Puntaje | Justificación | Estado)
r += 3
put(wcr, f'A{r}', 'RANGOS POR VARIABLE (mínimo INCLUIDO, máximo EXCLUIDO; no hay vacíos ni solapamientos)', size=12, bold=True, color=NAVY)
r += 1
crit_head = r
head(wcr, r, ['Variable', 'Indicador', 'Rango mínimo (incluido)', 'Rango máximo (excluido)', 'Nivel', 'Puntaje (piloto)',
              'Justificación', 'Estado de validación'], None, NAVY)
for col, w in zip('ABCDEFGH', [19, 52, 14, 46, 20, 12, 52, 20]):
    wcr.column_dimensions[col].width = w
CRIT = [
    ('Antigüedad', 'Años como asociado (ANTIGUEDAD_ANIOS)', 0, '=P_ANT_ALTO_MAX+1', 'Alto', 1, 'Hoja "Variables": 0-1 Alto. Años enteros.', 'Base del usuario'),
    ('Antigüedad', 'Años como asociado', '=P_ANT_ALTO_MAX+1', '=P_ANT_BAJO_MIN', 'Medio', 2, 'Hoja "Variables": 2-4 Medio (más de 1 y menos de 5).', 'Base del usuario'),
    ('Antigüedad', 'Años como asociado', '=P_ANT_BAJO_MIN', 'sin límite', 'Bajo', 3, 'Hoja "Variables": 5 o más Bajo.', 'Base del usuario'),
    ('Antigüedad', 'Vacío / negativo / mayor que la edad', '—', '—', 'Sin información', '—', 'No se asume riesgo bajo: queda pendiente.', 'Regla de calidad'),
    ('Ingreso', 'Ingreso mensual (INGRESO), USD', 0.01, '=P_ING_L1', 'Alto', 1, 'Por debajo del corte de tarjeta Dorada.', 'Propuesto'),
    ('Ingreso', 'Ingreso mensual, USD', '=P_ING_L1', '=P_ING_L2', 'Medio', 2, 'Tramo intermedio.', 'Propuesto'),
    ('Ingreso', 'Ingreso mensual, USD', '=P_ING_L2', 'sin límite', 'Bajo', 3, 'Desde aquí aplica la tabla de vehículo.', 'Propuesto'),
    ('Ingreso', 'Vacío = no informado; 0 = ingreso cero', '—', '—', 'Sin información / Ingreso cero', '—', 'No se convierten en cero ni en riesgo bajo.', 'Regla de calidad'),
    ('Deuda SSF', 'Saldo de deuda en la central de riesgo (SALDO_SSF), USD', 0.01, '=P_DEU_L1', 'Bajo', 3, 'Saldo pequeño.', 'Propuesto'),
    ('Deuda SSF', 'Saldo de deuda, USD', '=P_DEU_L1', '=P_DEU_L2', 'Medio', 2, 'Tramo intermedio.', 'Propuesto'),
    ('Deuda SSF', 'Saldo de deuda, USD', '=P_DEU_L2', 'sin límite', 'Alto', 1, 'Tercio de mayor deuda de la base.', 'Propuesto'),
    ('Deuda SSF', 'Saldo 0 o negativo = sin deuda; vacío = sin información', '—', '—', 'Sin deuda / Sin información', '3 / —', 'Saldo ≠ cuota.', 'Regla de calidad'),
    ('Deuda ÷ ingreso anual', 'SALDO_SSF ÷ (INGRESO × 12)', 0, '=P_DTI_L1', 'Bajo', '—', 'Deuda menor a 3 meses de ingreso.', 'Propuesto'),
    ('Deuda ÷ ingreso anual', 'SALDO_SSF ÷ (INGRESO × 12)', '=P_DTI_L1', '=P_DTI_L2', 'Medio', '—', 'Entre 3 y 12 meses de ingreso.', 'Propuesto'),
    ('Deuda ÷ ingreso anual', 'SALDO_SSF ÷ (INGRESO × 12)', '=P_DTI_L2', 'sin límite', 'Alto', '—', 'Deuda mayor o igual al ingreso de un año.', 'Propuesto'),
    ('Cuotas aportación vencidas', 'CUOTAS_VENC_APORT', 0, '=P_MORA_L1', 'Bajo', 3, 'Al día.', 'Propuesto'),
    ('Cuotas aportación vencidas', 'CUOTAS_VENC_APORT', '=P_MORA_L1', '=P_MORA_L2', 'Medio', 2, 'Atraso leve.', 'Propuesto'),
    ('Cuotas aportación vencidas', 'CUOTAS_VENC_APORT', '=P_MORA_L2', 'sin límite', 'Alto', 1, 'Atraso.', 'Propuesto'),
    ('Categoría SSF', 'CALIFICACION_SSF', 'A1', 'A1', 'Bajo', 3, 'Mejor categoría.', 'Memorando (A1/A2)'),
    ('Categoría SSF', 'CALIFICACION_SSF', 'A2', 'A2', 'Medio', 2, 'Permitida por el memorando.', 'Memorando (A1/A2)'),
    ('Categoría SSF', 'CALIFICACION_SSF', 'B, C, D, E', '—', 'Alto', 1, 'Fuera del memorando.', 'Memorando'),
]
for row in CRIT:
    r += 1
    for j, v in enumerate(row, 1):
        c = wcr.cell(r, j, v)
        c.font = font(9)
        c.border = box
        c.alignment = Alignment(wrap_text=True, vertical='center')
        if j in (3, 4) and row[0] in ('Ingreso', 'Deuda SSF'):
            c.number_format = USD2
        if j in (3, 4) and row[0] == 'Deuda ÷ ingreso anual':
            c.number_format = PCT
    lv = row[4]
    wcr.cell(r, 5).fill = fill({'Alto': RED_F, 'Medio': YEL_F, 'Bajo': GREEN_F}.get(lv, GRAY_F))
r += 2
NOTAS = [
    'NIVEL DE RIESGO DEL ASOCIADO (propuesto): se cuentan los datos en nivel Alto entre antigüedad, ingreso, deuda SSF, cuotas vencidas y '
    'categoría SSF. 0 = BAJO · 1 = MEDIO · 2 o más = ALTO. Si alguno no tiene información = "Sin información suficiente".',
    'PUNTAJE (piloto, solo informativo y para ordenar): Σ (puntos × peso); Bajo = 3, Medio = 2, Alto = 1; escala de 1 a 3. No decide.',
    'CAPACIDAD (indicador aproximado): la base trae el SALDO de la deuda, NO la cuota mensual. La relación cuota/ingreso y la capacidad de '
    'pago real NO se pueden calcular con esta base: quedan pendientes de la evaluación crediticia formal.',
    'Antigüedad: estos niveles describen SOLO la antigüedad; no son el riesgo crediticio integral del asociado.',
]
for t in NOTAS:
    put(wcr, f'A{r}', t, italic=True, color=GRAY, size=9)
    r += 1
wcr.freeze_panes = 'A5'

# ============================================================== 7_Calculo (motor con fórmulas)
wc = wb.create_sheet(CALC)
COLS = []  # (clave, encabezado, fórmula con {r}, formato, ancho)


def add(k, h, f, fmt=None, w=12):
    COLS.append((k, h, f, fmt, w))


b = lambda k: f"'{BASE}'!{B[k]}{{r}}"
add('ID', 'ID ASOCIADO', f'=IF({b("ASOC")}="","",{b("ASOC")}&"")', '@', 11)
add('NOM', 'NOMBRE', f'=IF({{ID}}="","",{b("NOM")})', None, 30)
add('AGE', 'AGENCIA', f'=IF({{ID}}="","",{b("AGE")}&"")', None, 14)
add('ING', 'INGRESO (USD/mes)', f'=IF({{ID}}="","",IF(ISNUMBER({b("ING")}),{b("ING")},""))', USD2, 12)
add('SSF', 'DEUDA SSF (USD)', f'=IF({{ID}}="","",IF(ISNUMBER({b("SSF")}),{b("SSF")},""))', USD2, 12)
add('ANT', 'ANTIGÜEDAD (años)', f'=IF({{ID}}="","",IF(ISNUMBER({b("ANT")}),{b("ANT")},""))', '0', 10)
add('FON', 'FONDO (USD)', f'=IF({{ID}}="","",IF(ISNUMBER({b("FON")}),{b("FON")},""))', USD2, 12)
add('EDADC', 'EDAD AL CORTE', f'=IF({{ID}}="","",DATEDIF(DATE(VALUE(LEFT({b("FNAC")},4)),VALUE(MID({b("FNAC")},6,2)),'
    f'VALUE(RIGHT({b("FNAC")},2))),P_FECHA_CORTE,"y"))', '0', 9)
add('CATI', 'CATEGORÍA INTERNA', f'=IF({{ID}}="","",TRIM({b("CATI")}&""))', None, 9)
add('CATE', 'CATEGORÍA SSF', f'=IF({{ID}}="","",TRIM({b("CATE")}&""))', None, 9)
# criterios del memorando (contiguos)
add('C_EDAD', 'C1 EDAD 30-55', '=IF({ID}="","",IF(AND({EDADC}>=P_EDAD_MIN,{EDADC}<=P_EDAD_MAX),"Cumple","No cumple"))', None, 11)
add('C_ING', 'C2 INGRESO', '=IF({ID}="","",IF({ING}="","Sin información",IF({ING}>0,"Cumple","No cumple")))', None, 11)
add('C_FON', 'C3 FONDO', '=IF({ID}="","",IF({FON}="","Sin información",IF({FON}>0,"Cumple","No cumple")))', None, 11)
add('C_ANT', 'C4 ANTIGÜEDAD', f'=IF({{ID}}="","",IF({{ANT}}="","Sin información",IF(OR({{ANT}}<0,{{ANT}}>{b("EDAD")}),"Sin información","Cumple")))',
    None, 11)
add('C_CATI', 'C5 CAT. INTERNA A1/A2', '=IF({ID}="","",IF(OR({CATI}="A1",{CATI}="A2"),"Cumple",IF({CATI}="","Sin información","No cumple")))',
    None, 11)
add('C_CATE', 'C6 CAT. SSF A1/A2', '=IF({ID}="","",IF(OR({CATE}="A1",{CATE}="A2"),"Cumple",IF({CATE}="","Sin información","No cumple")))',
    None, 11)
add('ELEG', 'CRITERIOS INICIALES', '=IF({ID}="","",IF(COUNTIF({C_EDAD}:{C_CATE},"No cumple")>0,"No elegible",'
    'IF(COUNTIF({C_EDAD}:{C_CATE},"Sin información")>0,"Información insuficiente","' + E1 + '")))', None, 16)
add('M_ELEG', 'QUÉ FALTA O NO CUMPLE (memorando)', '=IF(OR({ID}="",{ELEG}="' + E1 + '"),"",IF({C_EDAD}="No cumple","Edad fuera de 30-55 años; ","")'
    '&IF({C_ING}="No cumple","Ingreso igual a 0; ",IF({C_ING}="Sin información","Ingreso no informado; ",""))'
    '&IF({C_FON}="No cumple","Fondo compensatorio igual a 0; ",IF({C_FON}="Sin información","Fondo no informado; ",""))'
    '&IF({C_ANT}="Sin información",IF({ANT}="","Antigüedad no informada; ","Antigüedad inconsistente; "),"")'
    '&IF({C_CATI}="Sin información","Sin categoría interna; ",IF({C_CATI}="No cumple","Categoría interna no es A1/A2; ",""))'
    '&IF({C_CATE}="Sin información","Sin categoría SSF; ",IF({C_CATE}="No cumple","Categoría SSF no es A1/A2; ","")))', None, 34)
# niveles (los 5 del nivel de riesgo, contiguos)
add('N_ANT', 'NIVEL ANTIGÜEDAD', f'=IF({{ID}}="","",IF({{ANT}}="","Sin información",IF(OR({{ANT}}<0,{{ANT}}>{b("EDAD")}),"Sin información",'
    'IF({ANT}<=P_ANT_ALTO_MAX,"Alto",IF({ANT}<P_ANT_BAJO_MIN,"Medio","Bajo")))))', None, 10)
add('N_ING', 'NIVEL INGRESO', '=IF({ID}="","",IF({ING}="","Sin información",IF({ING}<=0,"Sin información",'
    'IF({ING}<P_ING_L1,"Alto",IF({ING}<P_ING_L2,"Medio","Bajo")))))', None, 10)
add('N_DEU', 'NIVEL DEUDA SSF', '=IF({ID}="","",IF({SSF}="","Sin información",IF({SSF}<=0,"Sin deuda",'
    'IF({SSF}<P_DEU_L1,"Bajo",IF({SSF}<P_DEU_L2,"Medio","Alto")))))', None, 10)
add('N_MORA', 'NIVEL CUOTAS VENCIDAS', f'=IF({{ID}}="","",IF(NOT(ISNUMBER({b("CVA")})),"Sin información",'
    f'IF({b("CVA")}<P_MORA_L1,"Bajo",IF({b("CVA")}<P_MORA_L2,"Medio","Alto"))))', None, 10)
add('N_CATE', 'NIVEL CATEGORÍA SSF', '=IF({ID}="","",IF({CATE}="A1","Bajo",IF({CATE}="A2","Medio",IF({CATE}="","Sin información","Alto"))))',
    None, 10)
add('DTI', 'DEUDA ÷ INGRESO ANUAL', '=IF(OR({ID}="",{ING}="",{SSF}=""),"",IF({ING}<=0,"",MAX({SSF},0)/({ING}*12)))', PCT, 10)
add('N_DTI', 'NIVEL DEUDA ÷ INGRESO', '=IF({ID}="","",IF({DTI}="","Sin información",IF({SSF}<=0,"Sin deuda",'
    'IF({DTI}<P_DTI_L1,"Bajo",IF({DTI}<P_DTI_L2,"Medio","Alto")))))', None, 10)
add('AL_IB', 'ALERTA INGRESO BAJO Y DEUDA ALTA', '=IF({ID}="","",IF(AND({N_ING}="Alto",{N_DEU}="Alto"),1,0))', '0', 9)
add('ROJTXT', 'DATOS EN NIVEL ALTO', '=IF({ID}="","",MID(IF({N_ANT}="Alto",", antigüedad","")&IF({N_ING}="Alto",", ingreso","")'
    '&IF({N_DEU}="Alto",", deuda SSF","")&IF({N_MORA}="Alto",", cuotas vencidas","")&IF({N_CATE}="Alto",", categoría SSF",""),3,90))', None, 22)
add('ROJOS', 'DATOS EN NIVEL ALTO (de 5)', '=IF({ID}="","",COUNTIF({N_ANT}:{N_CATE},"Alto"))', '0', 9)
add('SININF', 'DATOS SIN INFORMACIÓN (de 5)', '=IF({ID}="","",COUNTIF({N_ANT}:{N_CATE},"Sin información"))', '0', 9)
pts = lambda k: f'IF({{{k}}}="Alto",1,IF({{{k}}}="Medio",2,3))'
add('PUNT', 'PUNTAJE PILOTO (1 a 3)', '=IF(OR({ID}="",{SININF}>0),"",ROUND(P_W_ANT*' + pts('N_ANT') + '+P_W_ING*' + pts('N_ING')
    + '+P_W_DEU*' + pts('N_DEU') + '+P_W_MORA*' + pts('N_MORA') + '+P_W_CATE*' + pts('N_CATE') + ',2))', '0.00', 9)
add('NIVEL', 'NIVEL DE RIESGO', '=IF({ID}="","",IF({ELEG}<>"' + E1 + '","No aplica",IF({SININF}>0,"Sin información suficiente",'
    'IF({ROJOS}=0,"BAJO",IF({ROJOS}=1,"MEDIO","ALTO")))))', None, 11)
add('CAPAC', 'DEUDA FRENTE AL INGRESO', '=IF(OR({ID}="",{ELEG}<>"' + E1 + '"),"",IF({N_DTI}="Sin información","Sin información",'
    'IF(OR({N_DTI}="Alto",{AL_IB}=1),"Deuda alta para su ingreso","Adecuada")))', None, 14)
# productos de la base
add('HIPO', 'PRECA HIPO (base)', f'=IF({{ID}}="","",IF(ISNUMBER({b("HIPO")}),{b("HIPO")},""))', USD, 11)
add('VEH', 'PRECA VEHICULO (base)', f'=IF({{ID}}="","",IF(ISNUMBER({b("VEH")}),{b("VEH")},""))', USD, 11)
add('TC', 'PRECA TC (base)', f'=IF({{ID}}="","",IF(ISNUMBER({b("TC")}),{b("TC")},""))', USD, 10)
add('TTC', 'YA TIENE TARJETA', f'=IF({{ID}}="","",{b("TTC")}&"")', None, 8)
add('VEH_T', 'VEHÍCULO SEGÚN TABLA', '=IF(OR({ID}="",{ING}=""),"",IFERROR(LOOKUP({ING},VEH_ING_MIN,VEH_MONTO),""))', USD, 11)
add('TC_T', 'TARJETA SEGÚN TABLA', '=IF(OR({ID}="",{ING}=""),"",IFERROR(LOOKUP({ING},TC_ING_MIN,TC_LIM),""))', USD, 10)
add('VEH_OK', 'VEHÍCULO DENTRO DE TABLA', '=IF(OR({ID}="",{VEH}=""),"",IF({VEH}={VEH_T},"Sí","No"))', None, 9)
add('TC_OK', 'TARJETA DENTRO DE TABLA', '=IF(OR({ID}="",{TC}=""),"",IF({TC}={TC_T},"Sí","No"))', None, 9)
add('CABE', 'DEUDA CABE EN EL HIPOTECARIO', '=IF(OR({ID}="",{HIPO}="",{SSF}=""),0,IF({SSF}<={HIPO},1,0))', '0', 9)
# estado de preevaluación
add('ESTADO', 'ESTADO DE PREEVALUACIÓN', '=IF({ID}="","",IF({ELEG}="No elegible","' + E5 + '",IF({ELEG}="Información insuficiente","' + E4 + '",'
    'IF(OR({NIVEL}="Sin información suficiente",{CAPAC}="Sin información"),"' + E4 + '",'
    'IF(AND({NIVEL}<>"ALTO",{CAPAC}="Adecuada"),"' + E2 + '",IF(AND({NIVEL}<>"ALTO",{CABE}=1),"' + E2 + '","' + E3 + '"))))))', None, 24)
add('P_HIPO', 'MONTO PROPUESTO HIPOTECARIO', '=IF({ESTADO}="' + E2 + '",{HIPO},"")', USD, 12)
add('P_VEH', 'MONTO PROPUESTO VEHÍCULO', '=IF(AND({ESTADO}="' + E2 + '",{CAPAC}="Adecuada"),{VEH},"")', USD, 12)
add('P_TC', 'MONTO PROPUESTO TARJETA', '=IF(AND({ESTADO}="' + E2 + '",{TTC}="NO"),{TC},"")', USD, 11)
add('E_TC', 'TARJETA A EVALUAR (revisión)', '=IF(AND({ESTADO}="' + E3 + '",{TC}<>"",{TTC}="NO",{TC_OK}="Sí"),{TC},"")', USD, 11)
add('PROD', 'PRODUCTO PROPUESTO', '=IF({ESTADO}="' + E2 + '",MID(IF({P_HIPO}<>"",", Hipotecario","")&IF({P_VEH}<>"",", Vehículo","")'
    '&IF({P_TC}<>"",", Tarjeta",""),3,60),IF({ESTADO}="' + E3 + '",IF({E_TC}<>"","Tarjeta a evaluar; hipotecario y vehículo: '
    'Pendiente de evaluación","Pendiente de evaluación"),""))', None, 22)
add('MOTIVO', 'MOTIVO', '=IF({ESTADO}="' + E2 + '",IF({CAPAC}="Adecuada","Cumple criterios iniciales; deuda adecuada para su ingreso; riesgo "&{NIVEL},'
    '"Deuda SSF alta para su ingreso, pero menor que el hipotecario ofrecido (condicionado a pagarla); riesgo "&{NIVEL}),'
    'IF({ESTADO}="' + E3 + '",IF({NIVEL}="ALTO","Riesgo ALTO: "&{ROJOS}&" de 5 datos en nivel Alto ("&{ROJTXT}&"). ","")'
    '&IF({CAPAC}<>"Adecuada","Deuda SSF de $"&TEXT({SSF},"#,##0")&" alta para su ingreso de $"&TEXT({ING},"#,##0")&" al mes"'
    '&IF({HIPO}="","; no tiene hipotecario con el cual pagarla",IF({CABE}=0,"; es mayor que el hipotecario ofrecido",""))&".",""),'
    'IF({ESTADO}="' + E4 + '",IF({M_ELEG}<>"",{M_ELEG},"Algún dato de riesgo sin información"),IF({ESTADO}="' + E5 + '",{M_ELEG},""))))', None, 40)
add('VALID', 'VALIDACIÓN REQUERIDA', '=IF({ESTADO}="' + E3 + '","Confirmar en la central de riesgo las cuotas mensuales actuales y la capacidad '
    'de pago real (la base solo trae saldos)",IF({ESTADO}="' + E4 + '","Completar o corregir: "&IF({M_ELEG}<>"",{M_ELEG},"datos de riesgo"),""))',
    None, 34)
add('ACCION', 'ACCIÓN SUGERIDA', '=IF({ESTADO}="' + E3 + '",IF({E_TC}<>"","Evaluar solo la tarjeta (límite según su tabla); hipotecario y '
    'vehículo después de verificar","Evaluación crediticia formal antes de ofrecer"),IF({ESTADO}="' + E4 + '",'
    '"Solicitar o actualizar los datos y volver a evaluar",""))', None, 30)
add('PEND', 'VALIDACIONES PENDIENTES ANTES DE OTORGAR', '=IF(OR({ESTADO}="' + E2 + '",{ESTADO}="' + E3 + '"),'
    '"Lista negra, créditos castigados, relacionados/empleados"'
    '&IF({CAPAC}="Deuda alta para su ingreso",IF({ESTADO}="' + E2 + '","; el hipotecario debe pagar primero su deuda SSF",""),"")'
    '&IF(AND({P_HIPO}<>"",N({P_HIPO})>P_HIPO_SOLV),"; solvencia en Hacienda (hipotecario)","")'
    '&IF(AND({P_VEH}<>"",N({P_VEH})>P_VEH_SOLV),"; solvencia en Hacienda (vehículo, propuesto)","")'
    '&IF(AND({P_TC}<>"",N({P_TC})>P_TC_TOPE),"; tarjeta mayor de $4,000: solvencia en Hacienda o límite $4,000","")'
    '&IF(AND({P_VEH}<>"",{VEH_OK}="No"),"; monto de vehículo fuera de su tabla: validar con Créditos","")'
    f'&IF(OR({b("APD")}="NO",N({b("CVA")})>0),"; ponerse al día en aportaciones","")'
    f'&IF({b("EDAD")}<>{{EDADC}},"; corregir la edad (no coincide con la fecha de nacimiento)","")'
    f'&IF(LEN({b("DUI")}&"")<>9,"; corregir el DUI","")'
    '&IF(AND({ING}<>"",N({ING})>P_ING_ATIP),"; confirmar el ingreso (muy alto)",""),"")', None, 46)
add('PASO', 'SIGUIENTE PASO', '=IF(OR({ESTADO}="' + E2 + '",{ESTADO}="' + E3 + '"),"6. Sujeto a evaluación crediticia formal","")', None, 20)
add('K_REC', 'ORDEN LISTA RECOMENDADOS', '=IF({ESTADO}="' + E2 + '",IF({NIVEL}="BAJO",1,2)*100000+ROW(),"")', '0', 9)
add('K_CAS', 'ORDEN LISTA CASOS', '=IF({ESTADO}="' + E3 + '",100000+ROW(),IF({ESTADO}="' + E4 + '",200000+ROW(),""))', '0', 9)

COL = {k: L(i + 1) for i, (k, *_r) in enumerate(COLS)}
title(wc, '')
head(wc, 1, [h for _k, h, *_r in COLS], [w for *_r, w in COLS], NAVY, 46)
for rr in range(2, CAP + 2):
    for j, (k, h, f, fmt, w) in enumerate(COLS, 1):
        txt = f.replace('{r}', str(rr))
        for kk, ll in COL.items():
            txt = txt.replace('{' + kk + '}', f'{ll}{rr}')
        c = wc.cell(rr, j, txt)
        if fmt:
            c.number_format = fmt
        c.font = font(9)
wc.freeze_panes = 'B2'
wc.auto_filter.ref = f'A1:{L(len(COLS))}{CAP + 1}'
rng = lambda k: f"'{CALC}'!${COL[k]}$2:${COL[k]}${CAP + 1}"
for k in COL:
    name('C_' + k, rng(k))
lvl_rules = [('"Alto"', RED_F), ('"Medio"', YEL_F), ('"Bajo"', GREEN_F), ('"ALTO"', RED_F), ('"MEDIO"', YEL_F), ('"BAJO"', GREEN_F)]
for k in ('N_ANT', 'N_ING', 'N_DEU', 'N_MORA', 'N_CATE', 'N_DTI', 'NIVEL'):
    for v, col in lvl_rules:
        wc.conditional_formatting.add(f'{COL[k]}2:{COL[k]}{CAP + 1}', CellIsRule(operator='equal', formula=[v], fill=fill(col)))


# ============================================================== listas dinámicas
def lista(sh_name, ttl, sub, key, cols, color):
    w = wb.create_sheet(sh_name)
    title(w, ttl, sub)
    w['A3'] = f'=COUNT(C_{key})&" asociados en esta lista (se actualiza sola al cambiar la base)"'
    w['A3'].font = font(11, True, color)
    head(w, 5, ['N.º'] + [h for h, *_ in cols] + ['fila'], [6] + [wd for *_, wd in cols] + [6], color, 46)
    last = len(cols) + 2
    for i in range(LCAP):
        rr = 6 + i
        c = w.cell(rr, last, f'=IFERROR(MATCH(SMALL(C_{key},{i + 1}),C_{key},0),"")')
        c.font = font(8, color='BFBFBF')
        w.cell(rr, 1, f'=IF({L(last)}{rr}="","",{i + 1})').font = font(9)
        for j, (h, k, fmt, wd) in enumerate(cols, 2):
            c = w.cell(rr, j, f'=IF(${L(last)}{rr}="","",IF(INDEX(C_{k},${L(last)}{rr})="","",INDEX(C_{k},${L(last)}{rr})))')
            c.font = font(9)
            if fmt:
                c.number_format = fmt
            if wd >= 30:
                c.alignment = Alignment(wrap_text=True, vertical='top')
    end = 5 + LCAP
    w.column_dimensions[L(last)].hidden = True
    w.freeze_panes = 'D6'
    w.auto_filter.ref = f'A5:{L(last - 1)}{end}'
    return w, end


REC_COLS = [('ID ASOCIADO', 'ID', '@', 11), ('NOMBRE', 'NOM', None, 30), ('AGENCIA', 'AGE', None, 14),
            ('INGRESO (USD/mes)', 'ING', USD2, 12), ('ANTIGÜEDAD (años)', 'ANT', '0', 10), ('DEUDA SSF (USD)', 'SSF', USD2, 12),
            ('DEUDA ÷ INGRESO ANUAL', 'DTI', PCT, 10), ('PUNTAJE PILOTO (1 a 3)', 'PUNT', '0.00', 9), ('NIVEL DE RIESGO', 'NIVEL', None, 10),
            ('DATOS EN NIVEL ALTO', 'ROJTXT', None, 18), ('PRODUCTO PROPUESTO', 'PROD', None, 22), ('MONTO HIPOTECARIO', 'P_HIPO', USD, 12), ('MONTO VEHÍCULO', 'P_VEH', USD, 11),
            ('MONTO TARJETA', 'P_TC', USD, 10), ('ESTADO DE PREEVALUACIÓN', 'ESTADO', None, 22), ('MOTIVO DE RECOMENDACIÓN', 'MOTIVO', None, 36),
            ('VALIDACIONES PENDIENTES', 'PEND', None, 46), ('SIGUIENTE PASO', 'PASO', None, 18)]
wr, end_r = lista('3_Recomendacion_Preliminar', 'RECOMENDACIÓN PRELIMINAR (no es aprobación ni desembolso)',
                  'Cumplen los criterios iniciales del memorando, su deuda es adecuada para su ingreso (o cabe en el hipotecario ofrecido) y su '
                  'riesgo NO es ALTO. Ordenados: riesgo BAJO primero, luego MEDIO. Montos = columnas PRECA de la base.', 'K_REC', REC_COLS, NAVY)
CAS_COLS = [('ID ASOCIADO', 'ID', '@', 11), ('NOMBRE', 'NOM', None, 30), ('AGENCIA', 'AGE', None, 14),
            ('INGRESO (USD/mes)', 'ING', USD2, 12), ('DEUDA SSF (USD)', 'SSF', USD2, 12), ('NIVEL DE RIESGO', 'NIVEL', None, 12), ('DATOS EN NIVEL ALTO', 'ROJTXT', None, 18),
            ('ESTADO', 'ESTADO', None, 24), ('MOTIVO ESPECÍFICO', 'MOTIVO', None, 44), ('VALIDACIÓN REQUERIDA', 'VALID', None, 40),
            ('ACCIÓN SUGERIDA', 'ACCION', None, 34), ('PRODUCTO', 'PROD', None, 24), ('TARJETA A EVALUAR', 'E_TC', USD, 11),
            ('PRECA HIPO (base)', 'HIPO', USD, 11), ('PRECA VEHICULO (base)', 'VEH', USD, 11), ('PRECA TC (base)', 'TC', USD, 10)]
wcs, end_c = lista('4_Casos_por_Revisar', 'CASOS POR REVISAR', 'Estado 3: cumplen criterios iniciales pero tienen riesgo ALTO o deuda alta '
                   'para su ingreso. Estado 4: les falta información para evaluarlos. Cada fila dice qué revisar y qué hacer.', 'K_CAS', CAS_COLS, BLUE)
for w, end, k in ((wr, end_r, 'J'), (wcs, end_c, 'G')):
    for v, col in [('"BAJO"', GREEN_F), ('"MEDIO"', YEL_F), ('"ALTO"', RED_F)]:
        w.conditional_formatting.add(f'{k}6:{k}{end}', CellIsRule(operator='equal', formula=[v], fill=fill(col)))

# ============================================================== 1_Resumen_Ejecutivo
ws = wb.create_sheet('1_Resumen_Ejecutivo', 0)
title(ws, 'COMÉDICA – PREEVALUACIÓN CREDITICIA: RESUMEN EJECUTIVO',
      'Preevaluación preliminar. NO es aprobación de crédito ni autorización de desembolso. Todo se calcula con fórmulas sobre 6_Base_Original.')
for col, w in zip('ABCDEF', [60, 15, 15, 12, 3, 3]):
    ws.column_dimensions[col].width = w
cnt_est = lambda e: f'=COUNTIF(C_ESTADO,"{e}")'
rows = [
    ('Fecha de corte de la información', '=P_FECHA_CORTE', 'dd/mm/yyyy'),
    ('Total de registros analizados (base)', '=COUNTIF(C_ID,"?*")', INT),
    ('1. Cumplen criterios iniciales del memorando', f'=COUNTIF(C_ELEG,"{E1}")', INT),
    ('2. Recomendados preliminarmente', cnt_est(E2), INT),
    ('3. Requieren revisión adicional (con observaciones)', cnt_est(E3), INT),
    ('4. Pendientes por información insuficiente', cnt_est(E4), INT),
    ('5. No elegibles según el memorando', cnt_est(E5), INT),
    ('Control: 2 + 3 + 4 + 5 = total', None, None),
]
r = 4
put(ws, f'A{r}', 'INDICADORES PRINCIPALES', size=12, bold=True, color=WHITE, bg=NAVY)
put(ws, f'B{r}', 'Valor', bold=True, color=WHITE, bg=NAVY, h='center')
RES = {}
for lab, f, fmt in rows:
    r += 1
    put(ws, f'A{r}', lab, border=True, bold=lab.startswith(('2.', 'Total')))
    RES[lab[:2]] = r
    if f:
        put(ws, f'B{r}', f, fmt, border=True, bold=True, h='right')
r_ctrl = r
ws[f'B{r_ctrl}'] = f'=IF(B{r_ctrl - 4}+B{r_ctrl - 3}+B{r_ctrl - 2}+B{r_ctrl - 1}=B{r_ctrl - 6},"OK","REVISAR")'
ws[f'B{r_ctrl}'].font = font(10, True)
ws[f'B{r_ctrl}'].border = box
r_rec = r_ctrl - 4
r_rev = r_ctrl - 3

r += 2
put(ws, f'A{r}', 'DISTRIBUCIÓN POR NIVEL DE RIESGO', size=12, bold=True, color=WHITE, bg=NAVY)
for j, h in enumerate(['Recomendados', 'Por revisar', 'Total'], 2):
    put(ws, f'{L(j)}{r}', h, bold=True, color=WHITE, bg=NAVY, h='center')
r_niv0 = r + 1
for lv, col, desc in [('BAJO', GREEN_F, 'ningún dato en nivel Alto'), ('MEDIO', YEL_F, '1 dato en nivel Alto'),
                      ('ALTO', RED_F, '2 o más datos en nivel Alto')]:
    r += 1
    put(ws, f'A{r}', f'Riesgo {lv} ({desc})', border=True, bg=col)
    put(ws, f'B{r}', f'=COUNTIFS(C_ESTADO,"{E2}",C_NIVEL,"{lv}")', INT, border=True, h='right')
    put(ws, f'C{r}', f'=COUNTIFS(C_ESTADO,"{E3}",C_NIVEL,"{lv}")', INT, border=True, h='right')
    put(ws, f'D{r}', f'=B{r}+C{r}', INT, border=True, h='right', bold=True)
r_niv1 = r
r += 1
put(ws, f'A{r}', 'Nota: ningún asociado con riesgo ALTO queda en la lista de recomendados; va a "Casos por revisar".', italic=True, color=GRAY, size=9)

r += 2
put(ws, f'A{r}', 'PRODUCTOS (solo recomendados preliminarmente)', size=12, bold=True, color=WHITE, bg=NAVY)
for j, h in enumerate(['Asociados', 'Monto (USD)'], 2):
    put(ws, f'{L(j)}{r}', h, bold=True, color=WHITE, bg=NAVY, h='center')
for lab, k in (('Hipotecario (PRECA HIPO)', 'P_HIPO'), ('Vehículo (PRECA VEHICULO)', 'P_VEH'), ('Tarjeta de crédito (PRECA TC)', 'P_TC')):
    r += 1
    put(ws, f'A{r}', lab, border=True)
    put(ws, f'B{r}', f'=COUNT(C_{k})', INT, border=True, h='right')
    put(ws, f'C{r}', f'=SUM(C_{k})', USD, border=True, h='right')
r += 1
put(ws, f'A{r}', 'Tarjetas a evaluar en casos por revisar (límite según su tabla)', border=True)
put(ws, f'B{r}', '=COUNT(C_E_TC)', INT, border=True, h='right')
put(ws, f'C{r}', '=SUM(C_E_TC)', USD, border=True, h='right')
r += 1
put(ws, f'A{r}', 'Un asociado puede tener más de un producto. Los montos son los de la base: Riesgo no los modifica.', italic=True, color=GRAY, size=9)

r += 2
put(ws, f'A{r}', 'ALERTAS PRINCIPALES', size=12, bold=True, color=WHITE, bg='C00000')
put(ws, f'B{r}', 'Registros', bold=True, color=WHITE, bg='C00000', h='center')
ALERTAS = [
    ('Ingreso no informado (vacío) en la base', '=COUNTIF(C_C_ING,"Sin información")'),
    ('Ingreso igual a 0', '=COUNTIF(C_C_ING,"No cumple")'),
    ('Recomendados con vehículo fuera de su tabla (validar con Créditos)', f'=COUNTIFS(C_ESTADO,"{E2}",C_VEH_OK,"No",C_P_VEH,">0")'),
    ('Recomendados condicionados a pagar su deuda con el hipotecario', f'=COUNTIFS(C_ESTADO,"{E2}",C_CAPAC,"Deuda alta para su ingreso")'),
    ('Edad anotada distinta a la fecha de nacimiento (toda la base)', f"=SUMPRODUCT(('{BASE}'!$C$2:$C${CAP + 1}<>\"\")*(C_EDADC<>\"\")*"
                                                                      f"('{BASE}'!$C$2:$C${CAP + 1}<>C_EDADC))"),
    ('Diferencia de registros: memorando (2,263) menos base', '=2263-COUNTIF(C_ID,"?*")'),
]
for lab, f in ALERTAS:
    r += 1
    put(ws, f'A{r}', lab, border=True)
    put(ws, f'B{r}', f, INT, border=True, h='right', bold=True)
r += 2
put(ws, f'A{r}', 'LIMITACIONES DE LA EVALUACIÓN', size=12, bold=True, color=WHITE, bg=GRAY)
for t in ['• La base trae el SALDO de las deudas, no la CUOTA mensual: no se puede calcular la capacidad de pago real.',
          '• Lista negra, créditos castigados, relacionados/empleados y solvencia en Hacienda no vienen en la base: se verifican aparte.',
          '• Los rangos de ingreso, deuda y cuotas vencidas, y el nivel de riesgo, son PROPUESTAS de Riesgo pendientes de aprobación.',
          '• El puntaje con pesos es PILOTO: solo ordena; no decide.',
          '• El hipotecario no tiene tabla oficial en la base; se toma el monto de PRECA HIPO.',
          '• La base estará vigente 3 meses desde su aprobación por el Comité de Créditos (memorando).']:
    r += 1
    put(ws, f'A{r}', t, size=9)
# gráficos
ch = BarChart()
ch.type = 'bar'
ch.title = 'Estado de preevaluación'
ch.y_axis.title = 'Asociados'
ch.add_data(Reference(ws, min_col=2, min_row=r_rec, max_row=r_rec + 3), titles_from_data=False)
ch.set_categories(Reference(ws, min_col=1, min_row=r_rec, max_row=r_rec + 3))
ch.legend = None
ch.height, ch.width = 7, 14
ws.add_chart(ch, 'F4')
ch2 = BarChart()
ch2.title = 'Nivel de riesgo (recomendados y por revisar)'
ch2.add_data(Reference(ws, min_col=2, min_row=r_niv0 - 1, max_col=3, max_row=r_niv1), titles_from_data=True)
ch2.set_categories(Reference(ws, min_col=1, min_row=r_niv0, max_row=r_niv1))
ch2.height, ch2.width = 7, 14
ws.add_chart(ch2, 'F20')

# ============================================================== 5_Distribucion_Controles
wd = wb.create_sheet('5_Distribucion_Controles', 4)
title(wd, 'DISTRIBUCIÓN Y CONTROLES DE CALIDAD', 'Conteos por variable y pruebas de integridad. Todo con fórmulas: "OK" = la prueba pasa.')
for col, w in zip('ABCDEFG', [44, 14, 14, 14, 14, 14, 30]):
    wd.column_dimensions[col].width = w
r = 4
put(wd, f'A{r}', 'DISTRIBUCIÓN POR VARIABLE (toda la base / cumplen criterios iniciales)', size=12, bold=True, color=WHITE, bg=NAVY)
for j, h in enumerate(['Bajo', 'Medio', 'Alto', 'Sin deuda', 'Sin información'], 2):
    put(wd, f'{L(j)}{r}', h, bold=True, color=WHITE, bg=NAVY, h='center')
for lab, k in (('Antigüedad', 'N_ANT'), ('Ingreso', 'N_ING'), ('Deuda SSF', 'N_DEU'), ('Deuda ÷ ingreso anual', 'N_DTI'),
               ('Cuotas de aportación vencidas', 'N_MORA'), ('Categoría SSF', 'N_CATE')):
    for scope, cond in (('toda la base', ''), ('cumplen criterios', f',C_ELEG,"{E1}"')):
        r += 1
        put(wd, f'A{r}', f'{lab} – {scope}', border=True)
        for j, lv in enumerate(['Bajo', 'Medio', 'Alto', 'Sin deuda', 'Sin información'], 2):
            put(wd, f'{L(j)}{r}', f'=COUNTIFS(C_{k},"{lv}"{cond})', INT, border=True, h='right')
r += 2
put(wd, f'A{r}', 'ESTADO × NIVEL DE RIESGO', size=12, bold=True, color=WHITE, bg=NAVY)
for j, h in enumerate(['BAJO', 'MEDIO', 'ALTO', 'Sin información suficiente', 'No aplica', 'Total'], 2):
    put(wd, f'{L(j)}{r}', h, bold=True, color=WHITE, bg=NAVY, h='center')
r_ex0 = r + 1
for e in (E2, E3, E4, E5):
    r += 1
    put(wd, f'A{r}', e, border=True)
    for j, lv in enumerate(['BAJO', 'MEDIO', 'ALTO', 'Sin información suficiente', 'No aplica'], 2):
        put(wd, f'{L(j)}{r}', f'=COUNTIFS(C_ESTADO,"{e}",C_NIVEL,"{lv}")', INT, border=True, h='right')
    put(wd, f'G{r}', f'=SUM(B{r}:F{r})', INT, border=True, h='right', bold=True)
r_ex1 = r
r += 2
put(wd, f'A{r}', 'ALTERNATIVAS DE UMBRAL: DEUDA ÷ INGRESO ANUAL (entre los que cumplen criterios)', size=12, bold=True, color=WHITE, bg=NAVY)
put(wd, f'B{r}', 'Umbral Alto', bold=True, color=WHITE, bg=NAVY, h='center')
put(wd, f'C{r}', 'Asociados ≥ umbral', bold=True, color=WHITE, bg=NAVY, h='center')
for lab, v in (('Alternativa A: 75% del ingreso anual (más estricta)', 0.75), ('Alternativa B: 100% (vigente en este libro)', '=P_DTI_L2'),
               ('Alternativa C: 150% (más flexible)', 1.5)):
    r += 1
    put(wd, f'A{r}', lab, border=True)
    put(wd, f'B{r}', v, PCT, border=True, h='center')
    put(wd, f'C{r}', f'=COUNTIFS(C_ELEG,"{E1}",C_DTI,">="&B{r})', INT, border=True, h='right')
r += 1
put(wd, f'A{r}', 'Ninguna alternativa es política aprobada. La B es la que usa el libro; cambie P_DTI_L2 en la hoja 2 para probar otra.',
    italic=True, color=GRAY, size=9)

r += 2
put(wd, f'A{r}', 'PRUEBAS DE INTEGRIDAD Y CONCILIACIÓN', size=12, bold=True, color=WHITE, bg=NAVY)
put(wd, f'B{r}', 'Resultado', bold=True, color=WHITE, bg=NAVY, h='center')
put(wd, f'C{r}', 'Dato', bold=True, color=WHITE, bg=NAVY, h='center')
BID = f"'{BASE}'!$A$2:$A${CAP + 1}"
ok = lambda cond: f'=IF({cond},"OK","REVISAR")'
CTRL = [
    ('1. Registros fuente = registros calculados', ok(f'COUNTIF({BID},"?*")=COUNTIF(C_ID,"?*")'), f'=COUNTIF(C_ID,"?*")'),
    ('1b. El motor tiene capacidad para toda la base', ok(f"COUNTA('{BASE}'!$A:$A)-1<={CAP}"), f"=COUNTA('{BASE}'!$A:$A)-1"),
    ('2. Ningún identificador vacío', ok(f'SUMPRODUCT(({BID}="")*(\'{BASE}\'!$B$2:$B${CAP + 1}<>""))=0'), None),
    ('3. Identificadores duplicados', ok(f'SUMPRODUCT(({BID}<>"")*(COUNTIF({BID},{BID})>1))=0'),
     f'=SUMPRODUCT(({BID}<>"")*(COUNTIF({BID},{BID})>1))'),
    ('4. Rangos ordenados sin solapamiento (L1 < L2)', ok('AND(P_ING_L1<P_ING_L2,P_DEU_L1<P_DEU_L2,P_DTI_L1<P_DTI_L2,P_MORA_L1<P_MORA_L2,'
                                                           'P_ANT_ALTO_MAX<P_ANT_BAJO_MIN)'), None),
    ('5. Límites de rango bien clasificados (pruebas de borde, abajo)', None, None),
    ('6. Ingreso vacío NO se convierte en cero', ok(f"COUNTIF(C_C_ING,\"Sin información\")=SUMPRODUCT(({BID}<>\"\")*"
                                                    f"(NOT(ISNUMBER('{BASE}'!$P$2:$P${CAP + 1}))))"), f'=COUNTIF(C_C_ING,"Sin información")'),
    ('7. Nivel de riesgo = conteo de datos en nivel Alto', ok(f'COUNTIFS(C_NIVEL,"BAJO",C_ROJOS,">0")+COUNTIFS(C_NIVEL,"MEDIO",C_ROJOS,"<>1")'
                                                              f'+COUNTIFS(C_NIVEL,"ALTO",C_ROJOS,"<2")=0'), None),
    ('8. Pesos suman 100%', ok('ROUND(P_W_ANT+P_W_ING+P_W_DEU+P_W_MORA+P_W_CATE,6)=1'), '=P_W_ANT+P_W_ING+P_W_DEU+P_W_MORA+P_W_CATE'),
    ('9. Puntaje dentro de la escala 1 a 3', ok('AND(COUNTIF(C_PUNT,"<1")=0,COUNTIF(C_PUNT,">3")=0)'), None),
    ('10. Cada asociado en un solo estado', ok(f'COUNTIF(C_ESTADO,"?*")=COUNTIF(C_ID,"?*")'), None),
    ('11. Resumen concilia con las listas', ok(f"AND('1_Resumen_Ejecutivo'!B{r_rec}=COUNT(C_K_REC),'1_Resumen_Ejecutivo'!B{r_rev}=COUNTIF(C_ESTADO,\"{E3}\"),"
                                               f"COUNT('3_Recomendacion_Preliminar'!{L(len(REC_COLS) + 2)}6:{L(len(REC_COLS) + 2)}{end_r})=COUNT(C_K_REC),"
                                               f"COUNT('4_Casos_por_Revisar'!{L(len(CAS_COLS) + 2)}6:{L(len(CAS_COLS) + 2)}{end_c})=COUNT(C_K_CAS))"),
     None),
    ('12. Las listas tienen capacidad suficiente', ok(f'AND(COUNT(C_K_REC)<={LCAP},COUNT(C_K_CAS)<={LCAP})'), None),
    ('13. Sin errores de fórmula en el motor', ok(f"SUMPRODUCT(--ISERROR('{CALC}'!$A$2:${L(len(COLS))}${CAP + 1}))=0"), None),
    ('14. Montos propuestos de tarjeta dentro de su tabla', ok(f'COUNTIFS(C_P_TC,">0",C_TC_OK,"No")=0'), f'=COUNTIFS(C_P_TC,">0",C_TC_OK,"No")'),
    ('14b. Montos propuestos de vehículo fuera de su tabla (ALERTA, decisión pendiente)',
     f'=IF(COUNTIFS(C_P_VEH,">0",C_VEH_OK,"No")=0,"OK","ALERTA")', '=COUNTIFS(C_P_VEH,">0",C_VEH_OK,"No")'),
    ('15. Ningún riesgo ALTO en la lista de recomendados', ok(f'COUNTIFS(C_ESTADO,"{E2}",C_NIVEL,"ALTO")=0'), f'=COUNTIF(C_NIVEL,"ALTO")'),
    ('16. Todo es fórmula: al pegar una base nueva en 6_Base_Original se recalcula', '="OK (por diseño)"', None),
]
r_c0 = r + 1
for lab, f, dato in CTRL:
    r += 1
    put(wd, f'A{r}', lab, border=True)
    if f:
        put(wd, f'B{r}', f, border=True, bold=True, h='center')
    if dato:
        put(wd, f'C{r}', dato, INT, border=True, h='right')
r_c5 = r_c0 + 5
r += 2
put(wd, f'A{r}', 'PRUEBAS DE BORDE (mismas reglas que el motor)', size=12, bold=True, color=WHITE, bg=NAVY)
for j, h in enumerate(['Valor', 'Esperado', 'Obtenido', 'Resultado'], 2):
    put(wd, f'{L(j)}{r}', h, bold=True, color=WHITE, bg=NAVY, h='center')
TESTS = [('Ingreso', 1499.99, 'Alto'), ('Ingreso', 1500, 'Medio'), ('Ingreso', 2999.99, 'Medio'), ('Ingreso', 3000, 'Bajo'),
         ('Deuda SSF', 999.99, 'Bajo'), ('Deuda SSF', 1000, 'Medio'), ('Deuda SSF', 9999.99, 'Medio'), ('Deuda SSF', 10000, 'Alto'),
         ('Antigüedad', 1, 'Alto'), ('Antigüedad', 2, 'Medio'), ('Antigüedad', 4, 'Medio'), ('Antigüedad', 5, 'Bajo'),
         ('Cuotas vencidas', 0, 'Bajo'), ('Cuotas vencidas', 2, 'Medio'), ('Cuotas vencidas', 3, 'Alto'),
         ('Deuda ÷ ingreso anual', 0.2499, 'Bajo'), ('Deuda ÷ ingreso anual', 0.25, 'Medio'), ('Deuda ÷ ingreso anual', 1, 'Alto')]
FORM = {'Ingreso': 'IF(B{r}<P_ING_L1,"Alto",IF(B{r}<P_ING_L2,"Medio","Bajo"))',
        'Deuda SSF': 'IF(B{r}<P_DEU_L1,"Bajo",IF(B{r}<P_DEU_L2,"Medio","Alto"))',
        'Antigüedad': 'IF(B{r}<=P_ANT_ALTO_MAX,"Alto",IF(B{r}<P_ANT_BAJO_MIN,"Medio","Bajo"))',
        'Cuotas vencidas': 'IF(B{r}<P_MORA_L1,"Bajo",IF(B{r}<P_MORA_L2,"Medio","Alto"))',
        'Deuda ÷ ingreso anual': 'IF(B{r}<P_DTI_L1,"Bajo",IF(B{r}<P_DTI_L2,"Medio","Alto"))'}
t_first = r + 1
for var, v, exp in TESTS:
    r += 1
    put(wd, f'A{r}', var, border=True)
    put(wd, f'B{r}', v, PCT if var.startswith('Deuda ÷') else None, border=True, h='center')
    put(wd, f'C{r}', exp, border=True, h='center')
    put(wd, f'D{r}', '=' + FORM[var].replace('{r}', str(r)), border=True, h='center')
    put(wd, f'E{r}', f'=IF(C{r}=D{r},"OK","REVISAR")', border=True, h='center', bold=True)
wd[f'B{r_c5}'] = f'=IF(COUNTIF(E{t_first}:E{r},"OK")={len(TESTS)},"OK","REVISAR")'
wd[f'B{r_c5}'].font = font(10, True)
wd[f'B{r_c5}'].border = box
wd[f'B{r_c5}'].alignment = Alignment(horizontal='center')
for rng_ in (f'B{r_c0}:B{r_c0 + len(CTRL)}', f'E{t_first}:E{r}'):
    wd.conditional_formatting.add(rng_, CellIsRule(operator='equal', formula=['"OK"'], fill=fill(GREEN_F)))
    wd.conditional_formatting.add(rng_, CellIsRule(operator='equal', formula=['"REVISAR"'], fill=fill(RED_F)))
    wd.conditional_formatting.add(rng_, CellIsRule(operator='equal', formula=['"ALERTA"'], fill=fill(YEL_F)))

# ============================================================== 8_Hallazgos y 9_Registro_Cambios (texto de la auditoría)
from hallazgos_comedica import HALLAZGOS, CAMBIOS  # noqa: E402

wh = wb.create_sheet('8_Hallazgos')
title(wh, 'MATRIZ DE HALLAZGOS DE LA AUDITORÍA', 'Archivo auditado: Recomendacion_Base_PreEvaluada (1).xlsx')
head(wh, 4, ['N.º', 'Hoja', 'Hallazgo', 'Evidencia', 'Nivel de criticidad', 'Impacto', 'Corrección propuesta', 'Estado'],
     [5, 18, 40, 46, 14, 38, 44, 18])
for i, row in enumerate(HALLAZGOS, 1):
    for j, v in enumerate([i] + list(row), 1):
        c = wh.cell(4 + i, j, v)
        c.font = font(9, j == 5)
        c.border = box
        c.alignment = Alignment(wrap_text=True, vertical='top')
    wh.cell(4 + i, 5).fill = fill({'Crítica': 'F4B183', 'Alta': RED_F, 'Media': YEL_F, 'Baja': GREEN_F}[row[3]])
wh.freeze_panes = 'C5'
wh.auto_filter.ref = f'A4:H{4 + len(HALLAZGOS)}'
put(wh, f'A{6 + len(HALLAZGOS)}', 'Criterio de criticidad: Crítica = puede llevar a ofrecer crédito a quien no corresponde o a decidir con cifras '
    'equivocadas; Alta = afecta la clasificación de muchos registros; Media = afecta la interpretación o la trazabilidad; Baja = forma.',
    italic=True, color=GRAY, size=9)
wg = wb.create_sheet('9_Registro_Cambios')
title(wg, 'REGISTRO DE CAMBIOS')
head(wg, 4, ['N.º', 'Cambio', 'Motivo', 'Hoja afectada', 'Impacto en los resultados', 'Validación realizada', 'Estado'],
     [5, 44, 40, 20, 40, 40, 16])
for i, row in enumerate(CAMBIOS, 1):
    for j, v in enumerate([i] + list(row), 1):
        c = wg.cell(4 + i, j, v)
        c.font = font(9)
        c.border = box
        c.alignment = Alignment(wrap_text=True, vertical='top')
wg.freeze_panes = 'C5'

# orden de hojas
order = ['1_Resumen_Ejecutivo', '2_Criterios_y_Rangos', '3_Recomendacion_Preliminar', '4_Casos_por_Revisar', '5_Distribucion_Controles',
         BASE, CALC, '8_Hallazgos', '9_Registro_Cambios']
wb._sheets = [wb[n] for n in order]
for n, colr in zip(order, [NAVY, BLUE, NAVY, BLUE, GRAY, '404040', '404040', 'C00000', GRAY]):
    wb[n].sheet_properties.tabColor = colr
from openpyxl.worksheet.properties import PageSetupProperties  # noqa: E402
for n in order:
    w_ = wb[n]
    w_.page_setup.orientation = 'landscape'
    w_.page_setup.paperSize = w_.PAPERSIZE_LETTER
    w_.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    w_.page_setup.fitToWidth = 1
    w_.page_setup.fitToHeight = 0
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('ok', OUT, 'registros', NREG, 'columnas motor', len(COLS))
