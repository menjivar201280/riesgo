# -*- coding: utf-8 -*-
"""Genera el Excel corto y fácil: Recomendacion_Base_PreEvaluada.xlsx (4 hojas, solo valores).

Uso: python3 -I crear_excel_corto.py <libro_tecnico.xlsx> <salida.xlsx> <base_original.xlsx>
Regla única y estricta: RECOMENDADO = cumple el memorando y pasa TODAS las revisiones (todo ✔).
Si falla aunque sea una, es NO, y la hoja de NO dice cuál falló con el dato exacto.
"""
import sys
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

SRC, OUT, BASE = sys.argv[1:4]
CUMPLE = 'Cumple criterios verificables'
x = pd.read_excel(SRC, sheet_name='BD_CLASIFICADA', dtype={'ASOCIADO': str, 'DUI': str})
for c in ('PRECA HIPO', 'PRECA VEHICULO', 'PRECA TC', 'INGRESO', 'SALDO_SSF', 'FONDO_COMPENSATORIO'):
    x[c] = pd.to_numeric(x[c], errors='coerce')
TOTAL = len(x)
n1 = x.ELEGIBILIDAD_GENERAL == CUMPLE
RISK = ['NIVEL_ANTIGUEDAD', 'NIVEL_INGRESO', 'NIVEL_DEUDA', 'NIVEL_DEUDA_INGRESO', 'NIVEL_MORA_APORTACIONES',
        'NIVEL_CAT_INTERNA', 'NIVEL_CAT_EXTERNA']
# Ganar mucho no es un dato malo: el ingreso muy alto solo se confirma antes de otorgar.
DATO_MAL = ((x.NIVEL_ANTIGUEDAD == 'Inconsistente - revisar') | (x.ALERTA_ANT_EDAD == 1) | (x.CHK_EDAD != 'Coincide')
            | x.ESTADO_DEUDA_SSF.astype(str).str.startswith('Saldo negativo') | (x.DUI.astype(str).str.len() != 9))
VBAD = x.CONSIST_VEHICULO == 'No coincide con tabla'
TBAD = x.CONSIST_TC == 'No coincide con tabla'
SI = (n1 & (x.TOTAL_ALERTAS_RIESGO == 0) & ~(x[RISK] == 'Alto').any(axis=1) & ~DATO_MAL & (x.DQ_PRECA_SIN_INGRESO == 0)
      & (x.NIVEL_CAT_INTERNA == 'Bajo') & (x.NIVEL_CAT_EXTERNA == 'Bajo') & ~VBAD & ~TBAD)
NO = n1 & ~SI
money = lambda v: f'${v:,.0f}'

# Revisiones en palabras simples: (¿falló?, frase con el dato exacto)
REVISIONES = [
    (lambda e: e.NIVEL_INGRESO == 'Alto', lambda e: f'Gana poco: {money(e.INGRESO)} al mes (se pide $1,500 o más)'),
    (lambda e: e.NIVEL_DEUDA == 'Alto', lambda e: f'Debe mucho en otros bancos: {money(e.SALDO_SSF)} (se pide menos de $10,000)'),
    (lambda e: e.NIVEL_DEUDA_INGRESO == 'Alto',
     lambda e: f'Debe más de lo que gana en todo un año: debe {money(e.SALDO_SSF)} y gana {money(e.INGRESO * 12)} al año'),
    (lambda e: e.AL_FONDO_BAJO_DEUDA_ELEVADA == 1,
     lambda e: f'Poco fondo compensatorio ({money(e.FONDO_COMPENSATORIO)}) para lo que debe'),
    (lambda e: e.AL_ING_BAJO_DEUDA_ELEVADA == 1, lambda e: 'Gana poco para lo que debe'),
    (lambda e: e.AL_APORT_VENCIDAS == 1 or e.NIVEL_MORA_APORTACIONES == 'Alto',
     lambda e: f'Atrasado en aportaciones ({int(e.CUOTAS_VENC_APORT)} cuotas vencidas)'),
    (lambda e: e.NIVEL_ANTIGUEDAD == 'Alto' or e.AL_ANTIG_BAJA_SIN_HISTORIAL == 1,
     lambda e: f'Asociado nuevo: {int(e.ANTIGUEDAD_ANIOS)} año(s) (se pide más de 1)'),
    (lambda e: e.AL_CATEGORIA_A2 == 1 or e.NIVEL_CAT_INTERNA != 'Bajo' or e.NIVEL_CAT_EXTERNA != 'Bajo',
     lambda e: 'No es categoría A1'),
    (lambda e: e.CONSIST_VEHICULO == 'No coincide con tabla',
     lambda e: f'El vehículo no le toca: le dieron {money(e["PRECA VEHICULO"])}, su tabla dice '
               + (e.ESPERADO_VEHICULO if isinstance(e.ESPERADO_VEHICULO, str) else money(e.ESPERADO_VEHICULO))),
    (lambda e: e.CONSIST_TC == 'No coincide con tabla', lambda e: 'El límite de tarjeta no coincide con su tabla'),
    (lambda e: e.CHK_EDAD != 'Coincide', lambda e: 'La edad no coincide con su fecha de nacimiento'),
    (lambda e: len(str(e.DUI)) != 9, lambda e: 'El DUI está mal escrito'),
    (lambda e: e.ALERTA_ANT_EDAD == 1 or e.NIVEL_ANTIGUEDAD == 'Inconsistente - revisar',
     lambda e: 'Los años de asociado no cuadran con su edad'),
    (lambda e: str(e.ESTADO_DEUDA_SSF).startswith('Saldo negativo'), lambda e: 'Su deuda aparece en negativo'),
    (lambda e: e.DQ_PRECA_SIN_INGRESO == 1, lambda e: 'Tiene monto pre evaluado sin ingreso válido'),
]
MOTIVO_CORTO = ['Gana menos de $1,500 al mes', 'Debe $10,000 o más en otros bancos', 'Debe más de lo que gana en todo un año', 'Poco fondo compensatorio para lo que debe', 'Gana poco para lo que debe', 'Atrasado en aportaciones', 'Asociado nuevo (0 a 1 año)', 'No es categoría A1', 'El vehículo no le toca según su tabla', 'El límite de tarjeta no coincide con su tabla', 'La edad no coincide con su fecha de nacimiento', 'El DUI está mal escrito', 'Los años de asociado no cuadran con su edad', 'Su deuda aparece en negativo', 'Monto pre evaluado sin ingreso válido']
assert len(MOTIVO_CORTO) == len(REVISIONES)
x['POR_QUE_NO'] = ''
for i in x.index[NO]:
    e = x.loc[i]
    x.at[i, 'POR_QUE_NO'] = '; '.join(f(e) for falla, f in REVISIONES if falla(e))
assert (x.loc[NO, 'POR_QUE_NO'] != '').all()           # todo NO tiene su motivo
assert not any(falla(x.loc[i]) for i in x.index[SI] for falla, _ in REVISIONES)  # y ningún SÍ falla nada

# Productos que se recomiendan a los SÍ (la tarjeta solo si no tiene tarjeta: memorando)
tc_ok = x.ELEGIBILIDAD_TC == CUMPLE
x['HIPO'] = x['PRECA HIPO'].where(SI)
x['VEH'] = x['PRECA VEHICULO'].where(SI)
x['TC'] = x['PRECA TC'].where(SI & tc_ok)
SI_PROD = SI & x[['HIPO', 'VEH', 'TC']].notna().any(axis=1)
assert (SI == SI_PROD).all()

# Puntaje con los pesos del usuario (solo para ordenar a los SÍ)
PESOS = [('NIVEL_ANTIGUEDAD', .13), ('NIVEL_INGRESO', .30), ('NIVEL_DEUDA', .20), ('NIVEL_MORA_APORTACIONES', .12),
         ('NIVEL_CAT_EXTERNA', .25)]
PTS = {'Bajo': 3, 'Sin deuda': 3, 'Medio': 2, 'Alto': 1, 'Sin fondo': 1}
x['PUNTAJE'] = sum(x[c].map(PTS) * w for c, w in PESOS).round(2)
x['CALIFICACION'] = x.PUNTAJE.map(lambda v: 'Muy bueno' if v >= 2.6 else 'Bueno' if v >= 2.2 else 'Regular' if v >= 1.8 else 'Bajo')


def falta(e):
    p = ['Lista negra, castigados y relacionados/empleados']
    if pd.notna(e.HIPO):
        p.append('Hipotecario: solvencia en Hacienda')
    if pd.notna(e.VEH) and e.VEH > 25000:
        p.append('Vehículo mayor de $25,000: solvencia en Hacienda')
    if pd.notna(e.TC) and e.TC > 4000:
        p.append('Tarjeta: solvencia en Hacienda o se da con $4,000')
    if e.ALERTA_ING_ATIPICO == 1:
        p.append(f'Confirmar su ingreso (es muy alto: {money(e.INGRESO)})')
    return '; '.join(p)


x['FALTA'] = x.apply(falta, axis=1)

# ================================================================ estilos
F = 'Arial'
NAVY, GREEN, RED, GRAY = '1F3864', '548235', 'C00000', '595959'
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = lambda c: PatternFill('solid', fgColor=c)
font = lambda size=11, bold=False, color='000000', italic=False, u=None: Font(name=F, size=size, bold=bold, color=color,
                                                                               italic=italic, underline=u)
wb = openpyxl.Workbook()
wb.remove(wb.active)


def put(ws, cell, v, size=11, bold=False, color='000000', bg=None, h=None, italic=False, fmt=None, wrap=True):
    c = ws[cell]
    c.value = v
    c.font = font(size, bold, color, italic)
    c.alignment = Alignment(horizontal=h, vertical='center', wrap_text=wrap)
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    return c


def link(ws, cell, sheet, text, size=11):
    c = put(ws, cell, text, size, True, '0563C1', wrap=False)
    c.font = font(size, True, '0563C1', u='single')
    c.hyperlink = f"#'{sheet}'!A1"


def merge_put(ws, rng, v, **k):
    ws.merge_cells(rng)
    return put(ws, rng.split(':')[0], v, **k)


cnt = lambda m: int(m.sum())

# ================================================================ HOJA 1: RESUMEN
ws = wb.create_sheet('RESUMEN')
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = NAVY
for col, w in zip('ABCDEFGH', [3, 26, 16, 22, 4, 26, 16, 18]):
    ws.column_dimensions[col].width = w
merge_put(ws, 'B2:H2', 'BASE PRE EVALUADA: ¿A QUIÉN SE LE RECOMIENDA EL CRÉDITO?', size=18, bold=True, color=NAVY)
merge_put(ws, 'B3:H3', 'Recomendación de Riesgo para ayudar a decidir. No es la aprobación del crédito. '
          'Memorando del 27/10/2026 · base al 30/09/2026.', size=10, italic=True, color=RED)
# números grandes
cards = [('B', 'C', TOTAL, 'asociados en la base', NAVY, 'BD FINAL'),
         ('D', 'D', cnt(n1), 'cumplen el memorando', '2E75B6', None),
         ('F', 'F', cnt(SI), 'RECOMENDADOS', GREEN, 'RECOMENDADOS'),
         ('G', 'H', cnt(NO), 'NO recomendados', RED, 'NO RECOMENDADOS')]
for a, b, n, lab, col, sh in cards:
    for rr in (5, 6):
        for cc in range(ws[f'{a}{rr}'].column, ws[f'{b}{rr}'].column + 1):
            ws.cell(rr, cc).fill = fill(col)
        if a != b:
            ws.merge_cells(f'{a}{rr}:{b}{rr}')
    c = put(ws, f'{a}5', n, 30, True, 'FFFFFF', col, 'center', fmt='#,##0')
    put(ws, f'{a}6', lab, 11, True, 'FFFFFF', col, 'center')
    if sh:
        c.hyperlink = f"#'{sh}'!A1"
        c.font = font(30, True, 'FFFFFF', u='single')
ws.row_dimensions[5].height = 48
ws.row_dimensions[6].height = 22
merge_put(ws, 'B7:H7', 'Haga clic en un número para ver la lista.', size=9, italic=True, color=GRAY, h='center')

# cómo se decide
r = 9
merge_put(ws, f'B{r}:H{r}', '¿CÓMO SE DECIDE?  (una sola regla, sin excepciones)', size=13, bold=True, color='FFFFFF', bg=NAVY)
pasos = [('1', 'Cumple el memorando', 'Tiene ingreso y fondo, edad de 30 a 55 años, categoría A1 o A2. '
          f'De {TOTAL:,} asociados, {cnt(n1):,} cumplen.'),
         ('2', 'Pasa TODAS las revisiones', 'Gana $1,500 o más · debe menos de $10,000 · no debe más de lo que gana en un año · '
          'tiene fondo para lo que debe · al día en aportaciones · más de 1 año de asociado · categoría A1 · '
          'el vehículo y la tarjeta le tocan según su tabla · sus datos están bien escritos.'),
         ('3', 'Resultado', f'Si pasa todo = RECOMENDADO ({cnt(SI)}).  Si falla aunque sea una = NO ({cnt(NO)}). '
          'La hoja NO RECOMENDADOS dice cuál falló, con el dato exacto.')]
r += 1
for n, t, d in pasos:
    put(ws, f'B{r}', f'{n}.  {t}', 11, True, NAVY, 'EAF1FB')
    merge_put(ws, f'C{r}:H{r}', d, size=10, bg='EAF1FB')
    ws.row_dimensions[r].height = 34
    r += 1

# productos
r += 1
merge_put(ws, f'B{r}:H{r}', f'¿QUÉ SE LES OFRECE A LOS {cnt(SI)} RECOMENDADOS?', size=13, bold=True, color='FFFFFF', bg=GREEN)
r += 1
for j, h in zip('BCD', ['Producto', 'Asociados', 'Monto total']):
    put(ws, f'{j}{r}', h, 11, True, bg='E2EFDA', h='center')
r += 1
for lab, col in (('Hipotecario', 'HIPO'), ('Vehículo', 'VEH'), ('Tarjeta de crédito', 'TC')):
    put(ws, f'B{r}', lab, 11, True)
    put(ws, f'C{r}', cnt(x[col].notna()), h='center', fmt='#,##0')
    put(ws, f'D{r}', float(x[col].sum()), h='right', fmt='$#,##0')
    for j in 'BCD':
        ws[f'{j}{r}'].border = box
    r += 1
merge_put(ws, f'B{r}:H{r}', 'Hipotecario: se toma como viene en la base (si dice N/A, no aplica). Un asociado puede tener más de un producto. '
          'Cada producto se revisó por separado: si se le da más de uno, Créditos debe ver que pueda pagar todas las cuotas.',
          size=9, italic=True, color=GRAY)
ws.row_dimensions[r].height = 26

# por qué no
r += 2
merge_put(ws, f'B{r}:H{r}', f'¿POR QUÉ {cnt(NO)} NO SE RECOMIENDAN?  (uno puede tener varios motivos)', size=13, bold=True,
          color='FFFFFF', bg=RED)
r += 1
put(ws, f'B{r}', 'Motivo', 11, True, bg='FBE5D6')
ws.merge_cells(f'B{r}:F{r}')
put(ws, f'G{r}', 'Asociados', 11, True, bg='FBE5D6', h='center')
r += 1
motivos = []
for (falla, _f), lab in zip(REVISIONES, MOTIVO_CORTO):
    m = NO & x.apply(falla, axis=1)
    if m.any():
        motivos.append((cnt(m), lab))
for k, lab in sorted(motivos, reverse=True):
    put(ws, f'B{r}', lab)
    ws.merge_cells(f'B{r}:F{r}')
    put(ws, f'G{r}', k, h='center', fmt='#,##0')
    for j in 'BG':
        ws[f'{j}{r}'].border = box
    r += 1

# antes de otorgar
r += 1
merge_put(ws, f'B{r}:H{r}', 'ANTES DE OTORGAR (esto no viene en la base)', size=13, bold=True, color='FFFFFF', bg=GRAY)
r += 1
for t in ['• Revisar lista negra, créditos castigados y que no sea relacionado ni empleado.',
          '• Hipotecario: solvencia en Hacienda (memorando, monto mayor de $30,000).',
          '• Vehículo mayor de $25,000: solvencia en Hacienda (propuesta de Riesgo).',
          '• Tarjeta mayor de $4,000: si no está solvente en Hacienda, se da con $4,000 (memorando).',
          '• La base vence 3 meses después de que la apruebe el Comité.',
          '• La columna "QUÉ FALTA" de la hoja RECOMENDADOS lo dice para cada asociado.']:
    merge_put(ws, f'B{r}:H{r}', t, size=10)
    r += 1

# ================================================================ listas


def lista(name, color, title, nota, df, cols):
    w = wb.create_sheet(name)
    w.sheet_view.showGridLines = False
    w.sheet_properties.tabColor = color
    put(w, 'A1', f'{title}: {len(df)} asociados', 15, True, color, wrap=False)
    put(w, 'A2', nota, 10, italic=True, color=GRAY, wrap=False)
    link(w, 'A3', 'RESUMEN', '◄ Volver al RESUMEN', 10)
    for j, (h, src, fmt, wd) in enumerate(cols, 1):
        c = w.cell(5, j, h)
        c.font = font(10, True, 'FFFFFF')
        c.fill = fill(color)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = box
        w.column_dimensions[L(j)].width = wd
    w.row_dimensions[5].height = 34
    for i, (_, e) in enumerate(df.iterrows()):
        rr = 6 + i
        for j, (h, src, fmt, wd) in enumerate(cols, 1):
            v = e[src]
            v = None if isinstance(v, float) and pd.isna(v) else v
            c = w.cell(rr, j, v)
            c.font = font(10, bold=(j == 1))
            c.border = box
            c.alignment = Alignment(vertical='center', wrap_text=wd >= 30,
                                    horizontal='center' if src in ('ASOCIADO', 'PUNTAJE', 'CALIFICACION') else None)
            if fmt:
                c.number_format = fmt
            if i % 2:
                c.fill = fill('F2F5FA')
    end = 5 + len(df)
    for j, (h, src, fmt, wd) in enumerate(cols, 1):
        if src in ('HIPO', 'VEH', 'TC', 'PRECA HIPO', 'PRECA VEHICULO', 'PRECA TC'):
            c = w.cell(end + 1, j, float(df[src].sum()))
            c.number_format = '$#,##0'
            c.font = font(10, True)
    w.cell(end + 1, 2, 'TOTAL').font = font(10, True)
    w.freeze_panes = 'C6'
    w.auto_filter.ref = f'A5:{L(len(cols))}{end}'


USD = '$#,##0'
base = [('N° ASOCIADO', 'ASOCIADO', '@', 11), ('NOMBRE', 'NOMBRE', None, 36), ('DUI', 'DUI', '@', 11),
        ('AGENCIA', 'AGENCIA', None, 16), ('GANA AL MES', 'INGRESO', USD, 12), ('DEBE EN OTROS BANCOS', 'SALDO_SSF', USD, 13)]
lista('RECOMENDADOS', GREEN, 'RECOMENDADOS', 'Pasaron todas las revisiones. Ordenados del mejor al peor puntaje '
      '(pesos: antigüedad 13%, ingreso 30%, deuda 20%, aportaciones 12%, categoría SSF 25%).',
      x[SI].sort_values(['PUNTAJE', 'NOMBRE'], ascending=[False, True]),
      base + [('HIPOTECARIO', 'HIPO', USD, 13), ('VEHÍCULO', 'VEH', USD, 12), ('TARJETA', 'TC', USD, 11),
              ('PUNTAJE (1 a 3)', 'PUNTAJE', '0.00', 10), ('CALIFICACIÓN', 'CALIFICACION', None, 12),
              ('QUÉ FALTA ANTES DE OTORGAR', 'FALTA', None, 60)])
lista('NO RECOMENDADOS', RED, 'NO RECOMENDADOS', 'Cumplen el memorando, pero fallan al menos una revisión. '
      'La última columna dice cuál, con el dato exacto. Los montos son los que trae la base.',
      x[NO].sort_values('ASOCIADO', key=lambda s: s.astype(int)),
      base + [('HIPOTECARIO (base)', 'PRECA HIPO', USD, 13), ('VEHÍCULO (base)', 'PRECA VEHICULO', USD, 12),
              ('TARJETA (base)', 'PRECA TC', USD, 11), ('POR QUÉ NO', 'POR_QUE_NO', None, 70)])

# ================================================================ BD FINAL (copia sin cambios)
src = openpyxl.load_workbook(BASE)['BD FINAL']
wbd = wb.create_sheet('BD FINAL')
wbd.sheet_properties.tabColor = '404040'
for row in src.iter_rows():
    for c in row:
        if c.value is not None:
            d = wbd.cell(c.row, c.column, c.value)
            d.number_format = c.number_format
            d.font = font(9, c.row == 1, 'FFFFFF' if c.row == 1 else '000000')
            if c.row == 1:
                d.fill = fill('404040')
for j in range(1, src.max_column + 1):
    wbd.column_dimensions[L(j)].width = 14
wbd.freeze_panes = 'C2'
wbd.auto_filter.ref = f'A1:{L(src.max_column)}{src.max_row}'

wb.save(OUT)
print({'base': TOTAL, 'memorando': cnt(n1), 'SI': cnt(SI), 'NO': cnt(NO),
       'hipo': cnt(x.HIPO.notna()), 'veh': cnt(x.VEH.notna()), 'tc': cnt(x.TC.notna())})
