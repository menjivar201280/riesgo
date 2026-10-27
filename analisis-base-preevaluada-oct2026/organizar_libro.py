# -*- coding: utf-8 -*-
"""Ordena el libro para lectura: hoja INICIO con cifras clave y enlaces, pestañas por color,
hojas técnicas al final y columnas auxiliares de BD_CLASIFICADA agrupadas (plegadas).

Uso: python3 -I organizar_libro.py <libro.xlsx>   (paso final; después recalcular el libro)
No elimina datos ni cambia fórmulas: solo reemplaza la hoja LEEME por INICIO.
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

OUT = sys.argv[1]
wb = openpyxl.load_workbook(OUT)
F = 'Arial'
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)

bd = wb['BD_CLASIFICADA']
H = {bd.cell(1, j).value: L(j) for j in range(1, bd.max_column + 1)}
last = bd.max_row
RNG = lambda c: f"BD_CLASIFICADA!${H[c]}$2:${H[c]}${last}"
CUMPLE = 'Cumple criterios verificables'

# grupos de hojas: (titulo, color de pestaña, [(hoja, pregunta que responde)])
GROUPS = [
    ('1. RESPUESTA AL MEMORANDO', 'C55A11', [
        ('Cumplimiento_Memorando', '¿La base cumple las condiciones del memorando? Criterio por criterio y cifras memo vs base.'),
        ('Pendientes_Aprobacion', '¿Qué debe decidir Riesgos antes de usar los rangos?'),
    ]),
    ('2. RANGOS BAJO / MEDIO / ALTO', '2E75B6', [
        ('Variables_Rangos', '¿Cuáles son los rangos y cuántos asociados hay en cada nivel?'),
        ('Parametros', '¿Quiere cambiar un límite? Edite las celdas amarillas y todo se recalcula.'),
    ]),
    ('3. LISTAS DE ASOCIADOS QUE CUMPLEN', '548235', [
        ('Cumplen_Memorando', 'Todos los que cumplen los criterios verificables (todos los productos).'),
        ('Cumplen_Hipotecario', 'Los que cumplen y tienen monto hipotecario.'),
        ('Cumplen_Vehiculo', 'Los que cumplen y tienen monto de vehículo.'),
        ('Cumplen_TC', 'Los que cumplen, no tienen tarjeta y tienen límite de TC.'),
        ('Consumo', '¿Por qué no hay lista de consumo y qué se necesita?'),
    ]),
    ('4. DETALLE TÉCNICO (para Riesgos; no es necesario abrirlas)', '7F7F7F', [
        ('BD_CLASIFICADA', 'Base completa con el nivel de cada asociado. Columnas auxiliares plegadas: botón "+" arriba.'),
        ('Analisis_Producto', 'Exposición por producto y montos que no coinciden con las tablas.'),
        ('Evaluacion_Conjunta', 'Cruces entre variables y conteo de alertas.'),
        ('Matriz_Segmentacion', 'Matriz completa: límites, justificación, faltantes y limitaciones.'),
        ('Consolidado', 'Conteo por cada categoría con control de totales.'),
        ('Estadisticas', 'Mínimos, máximos, promedios y percentiles.'),
        ('Calidad_Conciliacion', 'Conciliación con la base original y errores de datos.'),
        ('Metodologia', 'Fórmulas, supuestos y restricciones.'),
    ]),
    ('5. ARCHIVO ORIGINAL (sin cambios)', '404040', [
        ('BD FINAL', 'Base original recibida.'),
        ('Variables', 'Hoja original de variables.'),
        ('RANGOS VEHICULO', 'Tabla original de montos de vehículo.'),
        ('RANGOS TC', 'Tabla original de límites de tarjeta.'),
    ]),
]

if 'LEEME' in wb.sheetnames:
    del wb['LEEME']
if 'INICIO' in wb.sheetnames:
    del wb['INICIO']
ws = wb.create_sheet('INICIO', 0)
ws.sheet_properties.tabColor = '1F3864'
ws.sheet_view.showGridLines = False
for col, w in zip('ABCD', [3, 34, 22, 90]):
    ws.column_dimensions[col].width = w

ws['B1'] = 'EVALUACIÓN DE LA BASE PRE EVALUADA — EMPIECE AQUÍ'
ws['B1'].font = Font(name=F, size=16, bold=True, color='1F3864')
ws['B2'] = ('Memorando del 27/10/2026 (hipotecario, vehículo y tarjeta). Rangos PROPUESTOS, pendientes de aprobación de Riesgos. '
            'No es una decisión de crédito.')
ws['B2'].font = Font(name=F, size=10, italic=True, color='C00000')

# ---- cifras clave (formulas)
ws['B4'] = 'CIFRAS CLAVE'
ws['B4'].font = Font(name=F, size=12, bold=True, color='1F3864')
EG, ETC = RNG('ELEGIBILIDAD_GENERAL'), RNG('ELEGIBILIDAD_TC')
TOTAL = f'COUNTA(BD_CLASIFICADA!$A$2:$A${last})'
kpis = [
    ('Asociados en la base', f'={TOTAL}', '#,##0', None),
    ('Cumplen todos los criterios verificables', f'=COUNTIF({EG},"{CUMPLE}")', '#,##0', 'E2EFDA'),
    ('Información incompleta (sin dato)', f'=COUNTIF({EG},"Información incompleta")', '#,##0', 'FFF2CC'),
    ('No cumplen al menos un criterio', f'=COUNTIF({EG},"No cumple criterio verificable")', '#,##0', 'FCE4D6'),
    ('   · Cumplen y tienen hipotecario', f'=COUNTIFS({EG},"{CUMPLE}",{RNG("PRECA HIPO")},">0")', '#,##0', None),
    ('   · Cumplen y tienen vehículo', f'=COUNTIFS({EG},"{CUMPLE}",{RNG("PRECA VEHICULO")},">0")', '#,##0', None),
    ('   · Cumplen y tienen tarjeta (sin TC actual)', f'=COUNTIFS({ETC},"{CUMPLE}",{RNG("PRECA TC")},">0")', '#,##0', None),
    ('   · Consumo', 'No incluido en la base ni en el memorando', None, None),
]
r = 5
for lab, f, fmt, fill in kpis:
    a = ws.cell(r, 2, lab)
    b = ws.cell(r, 3, f)
    a.font = Font(name=F, size=10, bold=not lab.startswith('   '))
    b.font = Font(name=F, size=11, bold=True)
    b.alignment = Alignment(horizontal='right' if fmt else 'left')
    if fmt:
        b.number_format = fmt
    else:
        b.font = Font(name=F, size=10, italic=True)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    for c in (a, b):
        c.border = box
        if fill:
            c.fill = PatternFill('solid', fgColor=fill)
    r += 1
ws.cell(r, 2, '¿La base aplica con el memorando?').font = Font(name=F, size=10, bold=True)
v = ws.cell(r, 3, f'=IF(C6=C5,"SÍ, en su totalidad","NO: solo "&TEXT(C6,"#,##0")&" de "&TEXT(C5,"#,##0")&" ("&TEXT(C6/C5,"0%")&")")')
v.font = Font(name=F, size=11, bold=True, color='C00000')
ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
for c in (ws.cell(r, 2), v):
    c.border = box
r += 2

# ---- indice con enlaces
ws.cell(r, 2, '¿QUÉ QUIERE VER?  (haga clic en el nombre de la hoja)').font = Font(name=F, size=12, bold=True, color='1F3864')
r += 1
order = ['INICIO']
for title, color, sheets in GROUPS:
    for j, t in enumerate([title, '', ''], 2):
        c = ws.cell(r, j, t)
        c.fill = PatternFill('solid', fgColor=color)
        c.font = Font(name=F, size=10, bold=True, color='FFFFFF')
    r += 1
    for name, question in sheets:
        if name not in wb.sheetnames:
            continue
        wb[name].sheet_properties.tabColor = color
        order.append(name)
        c = ws.cell(r, 2, name)
        c.hyperlink = f"#'{name}'!A1"
        c.font = Font(name=F, size=10, bold=True, color='0563C1', underline='single')
        c.border = box
        d = ws.cell(r, 3, question)
        d.font = Font(name=F, size=10)
        d.alignment = Alignment(wrap_text=True, vertical='center')
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        d.border = box
        r += 1
    r += 1

ws.cell(r, 2, 'CÓMO LEER LOS NIVELES').font = Font(name=F, size=12, bold=True, color='1F3864')
r += 1
for lvl, fill, txt in [('Alto', 'F8CBAD', 'Más riesgo (ej. ingreso bajo, deuda alta, poca antigüedad).'),
                       ('Medio', 'FFE699', 'Riesgo intermedio.'),
                       ('Bajo', 'C6E0B4', 'Menos riesgo.'),
                       ('Sin información', 'F2F2F2', 'Falta el dato en la base: no se clasifica como cero.')]:
    a = ws.cell(r, 2, lvl)
    a.fill = PatternFill('solid', fgColor=fill)
    a.font = Font(name=F, size=10, bold=True)
    a.border = box
    b = ws.cell(r, 3, txt)
    b.font = Font(name=F, size=10)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    r += 1
ws.cell(r, 2, 'En fondo compensatorio y aportaciones el nivel indica respaldo, no riesgo.').font = Font(name=F, size=9, italic=True)
r += 1
ws.cell(r, 2, 'Celdas amarillas (hoja Parametros) = se pueden editar. Todo lo demás son fórmulas.').font = Font(name=F, size=9, italic=True)

# ---- orden de hojas
order += [s.title for s in wb._sheets if s.title not in order]
wb._sheets = [wb[n] for n in order]
wb.active = 0
for s in wb._sheets:
    s.sheet_view.tabSelected = (s.title == 'INICIO')

# ---- BD_CLASIFICADA: plegar columnas auxiliares y de contacto
VISIBLE = {'ASOCIADO', 'NOMBRE', 'EDAD', 'DUI', 'ANTIGUEDAD_ANIOS', 'INGRESO', 'FONDO_COMPENSATORIO', 'PRECA HIPO',
           'PRECA VEHICULO', 'PRECA TC', 'SEGMENTO', 'CUOTAS_VENC_APORT', 'AGENCIA', 'TIENE_CREDITO_VIGENTE', 'MAX_CAT_RIESGO',
           'APORTACIONES_AL_DIA', 'MONTO_TOTAL_APORTACIONES', 'TIENE_TARJETA', 'CALIFICACION_SSF', 'SALDO_SSF', 'INSTITUCION',
           'EDAD_RECALCULADA', 'NIVEL_ANTIGUEDAD', 'NIVEL_INGRESO', 'NIVEL_DEUDA', 'NIVEL_DEUDA_INGRESO',
           'NIVEL_RESPALDO_FONDO', 'NIVEL_COBERTURA_FONDO', 'NIVEL_APORTACIONES', 'NIVEL_MORA_APORTACIONES', 'TRAMO_EDAD',
           'NIVEL_CAT_INTERNA', 'NIVEL_CAT_EXTERNA', 'ELEGIBILIDAD_GENERAL', 'ELEGIBILIDAD_TC', 'MOTIVOS', 'CONDICIONES_MEMO',
           'TC_LIMITE_SIN_SOLVENCIA', 'TOTAL_ALERTAS_RIESGO', 'TOTAL_ALERTAS_DATOS', 'ESTADO_DESCRIPTIVO'}
hidden = {CI(H[h]) for h in H if h not in VISIBLE}
# se redefine cada columna por separado: rangos superpuestos (p. ej. 1-31 y 4-8) se ignoran al abrir el libro
for key in list(bd.column_dimensions):
    del bd.column_dimensions[key]
for j in range(1, bd.max_column + 1):
    cd = bd.column_dimensions[L(j)]
    cd.width = 14 if j <= 31 else 22
    if j in hidden:
        cd.outline_level = 1
        cd.hidden = True
runs = sum(1 for j in hidden if j - 1 not in hidden)
bd.sheet_properties.outlinePr.summaryRight = False
level_fill = PatternFill('solid', fgColor='548235')
for h in ('ELEGIBILIDAD_GENERAL', 'ELEGIBILIDAD_TC', 'MOTIVOS', 'ESTADO_DESCRIPTIVO') + tuple(x for x in H if x.startswith('NIVEL_')):
    bd[f'{H[h]}1'].fill = level_fill

# ---- colores Bajo / Medio / Alto donde aparecen niveles
LEVEL_FILLS = {'Alto': 'F8CBAD', 'Medio': 'FFE699', 'Bajo': 'C6E0B4'}


def paint_levels(sheet, rng):
    for lvl, color in LEVEL_FILLS.items():
        sheet.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=[f'"{lvl}"'],
                                                         fill=PatternFill('solid', fgColor=color, bgColor=color)))


# solo niveles de RIESGO: en fondo y aportaciones un nivel "Alto" es respaldo, no riesgo
NOT_RISK = {'NIVEL_RESPALDO_FONDO', 'NIVEL_COBERTURA_FONDO', 'NIVEL_APORTACIONES'}
for h in H:
    if h.startswith('NIVEL_') and h not in NOT_RISK:
        paint_levels(bd, f'{H[h]}2:{H[h]}{last}')

vr = wb['Variables_Rangos']
r = 1
while r <= vr.max_row:
    if vr.cell(r, 2).value == 'Nivel de riesgo':
        end = r + 1
        while end <= vr.max_row and vr.cell(end, 1).value not in (None, 'Total'):
            end += 1
        paint_levels(vr, f'B{r + 1}:B{end - 1}')
        r = end
    r += 1

# ---- enlace "◄ INICIO" en las hojas con título (en la primera fila libre entre 2 y 3)
for name in order[1:]:
    s = wb[name]
    if name.startswith('Cumplen_') or name in ('Consumo', 'BD_CLASIFICADA', 'BD FINAL', 'Variables', 'RANGOS VEHICULO', 'RANGOS TC'):
        continue
    for row in (3, 2):
        if all(s.cell(row, k).value in (None, '') for k in range(1, s.max_column + 1)) and \
                not any(row in range(m.min_row, m.max_row + 1) for m in s.merged_cells.ranges):
            c = s.cell(row, 1, '◄ INICIO')
            c.hyperlink = "#'INICIO'!A1"
            c.font = Font(name=F, size=10, bold=True, color='0563C1', underline='single')
            break

wb.save(OUT)
print('hojas:', order)
print('columnas plegadas en BD_CLASIFICADA:', len(hidden), 'en', runs, 'grupos')
