# -*- coding: utf-8 -*-
"""Genera el libro corto para decisión: Resumen_Ejecutivo_Base_PreEvaluada.xlsx.

Uso: python3 -I crear_resumen_ejecutivo.py <libro_tecnico_recalculado.xlsx> <salida.xlsx> [<base_original.xlsx>]
El libro resultante no lleva fórmulas (solo valores y enlaces), así que no necesita recalcularse.
Toma los resultados ya calculados (y verificados) de BD_CLASIFICADA del libro técnico y asigna a cada
asociado un nivel de seguridad acumulativo (1 a 5). El libro técnico se conserva como respaldo.
"""
import re
import sys
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L

SRC, OUT = sys.argv[1], sys.argv[2]
CUMPLE = 'Cumple criterios verificables'
x = pd.read_excel(SRC, sheet_name='BD_CLASIFICADA', dtype={'ASOCIADO': str, 'DUI': str})
num = lambda c: pd.to_numeric(x[c], errors='coerce')
for c in ('PRECA HIPO', 'PRECA VEHICULO', 'PRECA TC'):
    x[c] = num(c)
TOTAL = len(x)

# ---------------------------------------------------------------- niveles de seguridad (acumulativos)
RISK = ['NIVEL_ANTIGUEDAD', 'NIVEL_INGRESO', 'NIVEL_DEUDA', 'NIVEL_DEUDA_INGRESO', 'NIVEL_MORA_APORTACIONES',
        'NIVEL_CAT_INTERNA', 'NIVEL_CAT_EXTERNA']
ALL_LOW = ['NIVEL_ANTIGUEDAD', 'NIVEL_INGRESO', 'NIVEL_DEUDA', 'NIVEL_DEUDA_INGRESO', 'NIVEL_MORA_APORTACIONES']
n1 = x.ELEGIBILIDAD_GENERAL == CUMPLE
n2 = n1 & (x.TOTAL_ALERTAS_RIESGO == 0)
n3 = n2 & ~(x[RISK] == 'Alto').any(axis=1)
n4 = n3 & (x.TOTAL_ALERTAS_DATOS == 0) & (x.NIVEL_CAT_INTERNA == 'Bajo') & (x.NIVEL_CAT_EXTERNA == 'Bajo')
n5 = n4 & x[ALL_LOW].isin(['Bajo', 'Sin deuda']).all(axis=1)
LEVELS = [
    (1, 'Cumple el memorando', 'Pasa los requisitos del memorando que se pueden revisar con la base.', n1),
    (2, 'Sin señales de riesgo', 'Además: no debe demasiado para lo que gana, no está atrasado en aportaciones y no es categoría A2.', n2),
    (3, 'Nada en rojo', 'Además: ningún dato en rojo (ingreso, deuda, antigüedad, aportaciones).', n3),
    (4, 'Máxima seguridad (RECOMENDADO)', 'Además: sus datos están bien, el monto que le dieron sí le corresponde según la tabla y es categoría A1.', n4),
    (5, 'Máxima seguridad, todo en verde', 'Parte del nivel 4 con todos sus datos en verde. Muy estricto: casi no deja tarjetas.', n5),
]
x['NIVEL_SEGURIDAD'] = 0
for k, *_rest, m in LEVELS:
    x.loc[m, 'NIVEL_SEGURIDAD'] = k

def motivo(r):
    """Motivo de "no recomendado" en palabras simples, con el dato que lo comprueba."""
    if r.NIVEL_SEGURIDAD >= 4:
        return ''
    out = []
    if r.NIVEL_INGRESO == 'Alto':
        out.append(f'Gana poco: ${r.INGRESO:,.0f} al mes (menos de $1,500)')
    if r.NIVEL_DEUDA == 'Alto':
        out.append(f'Debe mucho en bancos y otras entidades: ${r.SALDO_SSF:,.0f} ($10,000 o más)')
    if r.NIVEL_DEUDA_INGRESO == 'Alto':
        out.append(f'Lo que debe equivale a {r.MULT_DEUDA_INGRESO:,.0f} meses de su sueldo (12 o más)')
    if r.AL_FONDO_BAJO_DEUDA_ELEVADA == 1:
        out.append(f'Tiene poco fondo compensatorio (${r.FONDO_COMPENSATORIO:,.0f}) para lo que debe')
    if r.AL_APORT_VENCIDAS == 1:
        out.append(f'Está atrasado en sus aportaciones ({int(r.CUOTAS_VENC_APORT)} cuotas)')
    if r.NIVEL_ANTIGUEDAD == 'Alto':
        out.append(f'Es asociado nuevo ({int(r.ANTIGUEDAD_ANIOS)} año(s))')
    if r.AL_CATEGORIA_A2 == 1:
        out.append('Su categoría es A2, no A1')
    if r.CONSIST_VEHICULO == 'No coincide con tabla':
        esp = r.ESPERADO_VEHICULO if isinstance(r.ESPERADO_VEHICULO, str) else f'${r.ESPERADO_VEHICULO:,.0f}'
        out.append(f'Le dieron ${r["PRECA VEHICULO"]:,.0f} de vehículo, pero con su ingreso la tabla dice {esp}')
    if r.CONSIST_HIPO == 'No coincide con tabla':
        esp = r.ESPERADO_HIPO if isinstance(r.ESPERADO_HIPO, str) else f'${r.ESPERADO_HIPO:,.0f}'
        out.append(f'Le dieron ${r["PRECA HIPO"]:,.0f} de hipotecario, pero con su ingreso corresponde: {esp}')
    if r.CONSIST_TC == 'No coincide con tabla':
        out.append('El límite de tarjeta no coincide con la tabla')
    if r.CHK_EDAD != 'Coincide':
        out.append('La edad anotada no coincide con su fecha de nacimiento')
    if r.ALERTA_ING_ATIPICO == 1:
        out.append(f'Su ingreso es muy alto (${r.INGRESO:,.0f}): hay que confirmarlo')
    if len(str(r.DUI)) != 9:
        out.append('El DUI está mal escrito en la base')
    if r.ALERTA_ANT_EDAD == 1:
        out.append('Los años de asociado no cuadran con su edad')
    if str(r.ESTADO_DEUDA_SSF).startswith('Saldo negativo'):
        out.append('Su deuda aparece en negativo')
    return '; '.join(out) + '.' if out else 'Revisar caso'


x['MOTIVO_NO_MAXIMA'] = x.apply(motivo, axis=1)

T_VEH, T_TC, T_ANT = 'Tabla 1 (vehículo)', 'Tabla 2 (tarjeta)', 'Tabla 3 (antigüedad)'
T_HIPO, T_LIM, T_DATOS = 'Tabla 4 (hipotecario)', 'Tabla 5 (límites de riesgo)', 'Tabla 6 (revisión de datos)'


def tablas(r):
    """Tablas de la hoja 8 con las que se comparó al asociado para decidir el NO."""
    if r.NIVEL_SEGURIDAD >= 4:
        return ''
    t = []
    if r.CONSIST_VEHICULO == 'No coincide con tabla':
        t.append(T_VEH)
    if r.CONSIST_TC == 'No coincide con tabla':
        t.append(T_TC)
    if r.NIVEL_ANTIGUEDAD == 'Alto':
        t.append(T_ANT)
    if r.CONSIST_HIPO == 'No coincide con tabla':
        t.append(T_HIPO)
    if (r.NIVEL_INGRESO == 'Alto' or r.NIVEL_DEUDA == 'Alto' or r.NIVEL_DEUDA_INGRESO == 'Alto'
            or r.AL_FONDO_BAJO_DEUDA_ELEVADA == 1 or r.AL_APORT_VENCIDAS == 1 or r.AL_CATEGORIA_A2 == 1):
        t.append(T_LIM)
    if (r.CHK_EDAD != 'Coincide' or r.ALERTA_ING_ATIPICO == 1 or len(str(r.DUI)) != 9 or r.ALERTA_ANT_EDAD == 1
            or str(r.ESTADO_DEUDA_SSF).startswith('Saldo negativo')):
        t.append(T_DATOS)
    return ', '.join(t)


x['TABLA_USADA'] = x.apply(tablas, axis=1)


def pendientes(r):
    p = []
    if pd.notna(r['PRECA HIPO']):
        p.append('Hipotecario: solvencia en Hacienda')
    if pd.notna(r['PRECA TC']) and r['PRECA TC'] > 4000:
        p.append('TC: solvencia en Hacienda o bajar a $4,000')
    if pd.notna(r['PRECA TC']) and r['APORTACIONES_AL_DIA'] == 'NO':
        p.append('TC: ponerse al día en aportaciones')
    return '; '.join(p)


x['PENDIENTE_ANTES_DE_OTORGAR'] = x.apply(pendientes, axis=1)

# ---------------------------------------------------------------- estilos
F = 'Arial'
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
NAVY, BLUE, GREEN, ORANGE, GRAY, RED = '1F3864', '2E75B6', '548235', 'C55A11', '7F7F7F', 'C00000'
USD, INT = '"$"#,##0.00', '"$"#,##0'
LEVEL_FILLS = {'Alto': 'F8CBAD', 'Medio': 'FFE699', 'Bajo': 'C6E0B4'}
wb = openpyxl.Workbook()
wb.remove(wb.active)


def font(size=10, bold=False, color='000000', italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)


def band(ws, r, text, color, ncols=6):
    for j in range(2, 2 + ncols):
        ws.cell(r, j).fill = PatternFill('solid', fgColor=color)
    c = ws.cell(r, 2, text)
    c.font = font(11, True, 'FFFFFF')
    return r + 1


def cellw(ws, r, c, v, fmt=None, bold=False, fill=None, align=None, color='000000', size=10, wrap=True):
    cc = ws.cell(r, c, v)
    cc.font = font(size, bold, color)
    cc.border = box
    cc.alignment = Alignment(wrap_text=wrap, vertical='center', horizontal=align)
    if fmt:
        cc.number_format = fmt
    if fill:
        cc.fill = PatternFill('solid', fgColor=fill)
    return cc


def link(ws, cell, sheet, text=None):
    c = ws[cell]
    c.value = text or sheet
    c.hyperlink = f"#'{sheet}'!A1"
    c.font = font(10, True, '0563C1')
    c.font = Font(name=F, size=10, bold=True, color='0563C1', underline='single')


cnt = lambda m: int(m.sum())
pct = lambda m: m.sum() / TOTAL
prod = {'Hipotecario': 'PRECA HIPO', 'Vehículo': 'PRECA VEHICULO', 'Tarjeta de crédito': 'PRECA TC'}
tc_ok = x.ELEGIBILIDAD_TC == CUMPLE

# ================================================================ HOJA 1: RESUMEN
ws = wb.create_sheet('1_RESUMEN')
ws.sheet_properties.tabColor = NAVY
ws.sheet_view.showGridLines = False
for col, w in zip('ABCDEFG', [2, 46, 16, 16, 16, 16, 40]):
    ws.column_dimensions[col].width = w
ws['B1'] = 'RESUMEN: ¿QUÉ ASOCIADOS DE LA BASE PRE EVALUADA SE PUEDEN OFRECER CON SEGURIDAD?'
ws['B1'].font = font(15, True, NAVY)
ws['B2'] = ('Memorando del 27/10/2026 (Créditos, TC y Mercadeo → Riesgos). Base con corte al 30/09/2026. '
            'Este archivo es un apoyo para decidir: no es una aprobación de crédito.')
ws['B2'].font = font(9, italic=True, color=RED)
r = 4

r = band(ws, r, '1. RESPUESTA CORTA', NAVY)
short = [
    f'• La base NO cumple completa con el memorando: de {TOTAL:,} asociados, solo {cnt(n1):,} ({pct(n1):.0%}) cumplen todos los '
    f'criterios que se pueden verificar.',
    f'• De esos {cnt(n1):,}, los más seguros son {cnt(n4):,} asociados (nivel 4 – máxima seguridad): sin alertas, sin niveles '
    f'altos de riesgo, sin errores de datos y con categoría A1.',
    f'• Recomendación: ofrecer primero a los {cnt(n4):,} del nivel 4 (hoja 3), después de hacer las consultas externas pendientes '
    f'(punto 6).',
]
for t in short:
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    c = ws.cell(r, 2, t)
    c.font = font(11, True)
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[r].height = 32
    r += 1
r += 1

r = band(ws, r, '2. QUÉ SE REVISÓ', BLUE)
for t in [f'• Base recibida: {TOTAL:,} asociados pre evaluados para hipotecario, vehículo y tarjeta (el memorando dice 2,263; '
          f'el archivo tiene 2,262 personas).',
          '• Se revisaron los 7 criterios del memorando que se pueden comprobar: ingreso, fondo compensatorio, sin tarjeta (solo TC), '
          'antigüedad, categoría interna A1/A2, categoría SSF A1/A2 y edad de 30 a 55 años.',
          '• Se clasificó a cada asociado en niveles Bajo / Medio / Alto de riesgo (ingreso, deuda, deuda frente al ingreso, '
          'antigüedad, aportaciones vencidas).',
          '• No se pueden comprobar con la base: lista negra, créditos castigados, personas relacionadas o empleados y solvencia '
          'en Hacienda.']:
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    c = ws.cell(r, 2, t)
    c.font = font(10)
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[r].height = 28
    r += 1
r += 1

r = band(ws, r, '3. RESULTADO CONTRA EL MEMORANDO', ORANGE)
for j, h in enumerate(['Resultado', 'Asociados', '% de la base', '', '', 'Qué significa'], 2):
    if h:
        cellw(ws, r, j, h, bold=True, fill='FBE5D6')
r += 1
inc = x.ELEGIBILIDAD_GENERAL == 'Información incompleta'
noc = x.ELEGIBILIDAD_GENERAL == 'No cumple criterio verificable'
for lab, m, fill, mean in [('Cumplen todos los criterios', n1, 'E2EFDA', 'Pasan a la evaluación de seguridad (punto 4).'),
                           ('Información incompleta', inc, 'FFF2CC', 'Les falta un dato (sobre todo el ingreso): no se pueden evaluar.'),
                           ('No cumplen algún criterio', noc, 'FCE4D6', 'Ingreso $0, fondo $0 o 56 años al corte.')]:
    cellw(ws, r, 2, lab, bold=True, fill=fill)
    cellw(ws, r, 3, cnt(m), '#,##0', align='right')
    cellw(ws, r, 4, pct(m), '0.0%', align='right')
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
    cellw(ws, r, 5, mean)
    r += 1
cellw(ws, r, 2, 'Total', bold=True, fill='F2F2F2')
cellw(ws, r, 3, TOTAL, '#,##0', bold=True, align='right', fill='F2F2F2')
cellw(ws, r, 4, 1, '0.0%', bold=True, align='right', fill='F2F2F2')
r += 1
ing_vacio = int(x.INGRESO.isna().sum())
ing_cero = int((num('INGRESO') == 0).sum())
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
c = ws.cell(r, 2, f'Causa principal: {ing_vacio:,} asociados no tienen ingreso registrado y {ing_cero:,} tienen ingreso $0, '
                  f'aunque todos recibieron un monto pre evaluado.')
c.font = font(9, italic=True)
r += 2

r = band(ws, r, '4. NIVELES DE SEGURIDAD (cada nivel es más estricto que el anterior)', GREEN)
for j, h in enumerate(['Nivel', 'Asociados', 'Hipotecario', 'Vehículo', 'Tarjeta', 'Condición que se agrega'], 2):
    cellw(ws, r, j, h, bold=True, fill='E2EFDA', align='center')
r += 1
LEVEL_ROW = {}
for k, name, cond, m in LEVELS:
    LEVEL_ROW[k] = r
    rec = k == 4
    fill = 'C6E0B4' if rec else None
    cellw(ws, r, 2, f'{k}. {name}', bold=True, fill=fill)
    cellw(ws, r, 3, cnt(m), '#,##0', bold=rec, align='right', fill=fill)
    for j, (pl, col) in enumerate(prod.items(), 4):
        mm = m & x[col].notna() & (tc_ok if col == 'PRECA TC' else True)
        cellw(ws, r, j, cnt(mm), '#,##0', bold=rec, align='right', fill=fill)
    cellw(ws, r, 7, cond, fill=fill, size=9)
    ws.row_dimensions[r].height = 40
    r += 1
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
c = ws.cell(r, 2, 'Un asociado puede tener más de un producto, por eso los productos suman más que el total.')
c.font = font(9, italic=True)
r += 2

r = band(ws, r, f'5. RECOMENDADO: NIVEL 4 – MÁXIMA SEGURIDAD ({cnt(n4)} asociados)', GREEN)
for j, h in enumerate(['Producto', 'Asociados', 'Monto total', 'Monto promedio', '', 'Dónde ver la lista'], 2):
    if h:
        cellw(ws, r, j, h, bold=True, fill='E2EFDA', align='center')
r += 1
sheet_of = {'Hipotecario': '4_Hipotecario', 'Vehículo': '5_Vehiculo', 'Tarjeta de crédito': '6_Tarjeta'}
for pl, col in prod.items():
    mm = n4 & x[col].notna() & (tc_ok if col == 'PRECA TC' else True)
    cellw(ws, r, 2, pl, bold=True)
    cellw(ws, r, 3, cnt(mm), '#,##0', align='right')
    cellw(ws, r, 4, float(x.loc[mm, col].sum()), INT, align='right')
    cellw(ws, r, 5, float(x.loc[mm, col].mean()) if mm.any() else 0, INT, align='right')
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    cellw(ws, r, 6, '')
    link(ws, f'F{r}', sheet_of[pl], f'Hoja {sheet_of[pl]}')
    r += 1
cellw(ws, r, 2, 'Todos (sin repetir)', bold=True, fill='F2F2F2')
cellw(ws, r, 3, cnt(n4), '#,##0', bold=True, align='right', fill='F2F2F2')
ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
cellw(ws, r, 6, '')
link(ws, f'F{r}', '3_Maxima_Seguridad', 'Hoja 3_Maxima_Seguridad')
r += 2

r = band(ws, r, '6. ANTES DE OTORGAR (obligatorio, no está en la base)', RED)
hip4 = cnt(n4 & x['PRECA HIPO'].notna())
tc4_big = cnt(n4 & tc_ok & (x['PRECA TC'] > 4000))
for t in ['• Consultar lista negra y créditos castigados con la Cooperativa.',
          '• Verificar que no sea persona relacionada ni empleado (memorando, criterio 8).',
          f'• Solvencia en Hacienda: obligatoria para los {hip4} hipotecarios del nivel 4 (monto > $30,000).',
          f'• Tarjeta con límite mayor de $4,000 ({tc4_big} en el nivel 4): si no está solvente en Hacienda, se otorga con $4,000.',
          '• La base vence 3 meses después de la aprobación del Comité: actualizar edad, aportaciones y Central de Riesgo.']:
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    c = ws.cell(r, 2, t)
    c.font = font(10)
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[r].height = 20
    r += 1
r += 1

r = band(ws, r, '7. PROBLEMAS ENCONTRADOS EN LA BASE (para corregir con Créditos)', GRAY)
veh = num('PRECA VEHICULO')
age56 = cnt(x.EDAD_RECALCULADA > 55)
for t in [f'• {ing_vacio + ing_cero:,} asociados tienen monto pre evaluado sin ingreso válido: no se sabe con qué ingreso se calculó.',
          f'• {age56} asociados ya tienen 56 años al 30/09/2026 (el límite es 55).',
          f'• {cnt(num("FONDO_COMPENSATORIO") == 0)} asociados tienen fondo compensatorio $0 (el memorando exige fondo).',
          f'• {cnt(x.TIENE_TARJETA == "SI")} asociados ya tienen tarjeta: el memorando dice que nadie la tiene (solo aplica a TC).',
          f'• Vehículo: {cnt(veh == 150000)} asociado con $150,000, monto que no existe en la tabla (máximo $40,000).',
          '• El memorando dice 61 sin deuda en la SSF y 54 no al día en aportaciones; la base muestra 342 sin registro SSF y 186 no al día.',
          '• Crédito de consumo: no viene en la base ni en el memorando; se necesita su tabla de montos para evaluarlo.']:
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    c = ws.cell(r, 2, t)
    c.font = font(10)
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[r].height = 20
    r += 1
r += 1

r = band(ws, r, '8. QUÉ HAY EN ESTE ARCHIVO', NAVY)
for sh, desc in [('2_Niveles_y_Rangos', 'Rangos Bajo / Medio / Alto de cada variable y cuántos asociados hay en cada nivel.'),
                 ('3_Maxima_Seguridad', f'Los {cnt(n4)} asociados recomendados (nivel 4), todos los productos.'),
                 ('4_Hipotecario', 'Nivel 4 con monto hipotecario.'),
                 ('5_Vehiculo', 'Nivel 4 con monto de vehículo.'),
                 ('6_Tarjeta', 'Nivel 4 sin tarjeta actual y con límite de tarjeta.'),
                 ('7_Los_528', f'Los {cnt(n1)} que cumplen el memorando: columna ¿RECOMENDADO? (SÍ {cnt(n4)} / NO {cnt(n1) - cnt(n4)}) '
                               f'y el motivo de cada NO.'),
                 ('8_Tablas_usadas', 'Las tablas completas con las que se comparó a cada asociado (Tabla 1 a 6).'),
                 ('9_Como_explicarlo', 'Ejemplos reales paso a paso para explicar cada "NO" (y un "SÍ").'),
                 ('10_Asociados_por_tabla', 'Los asociados "NO" agrupados por tabla, con el dato que lo explica (aquí llevan los enlaces).'),
                 ('BD FINAL', 'Su base original, tal como la envió (sin cambios).')]:
    cellw(ws, r, 2, '')
    link(ws, f'B{r}', sh)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    cellw(ws, r, 3, desc)
    r += 1
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
c = ws.cell(r, 2, 'Respaldo técnico (fórmulas, metodología, conciliación): archivo Segmentacion_Riesgos_Base_PreEvaluada.xlsx.')
c.font = font(9, italic=True)

# ================================================================ HOJA 2: NIVELES Y RANGOS
wr = wb.create_sheet('2_Niveles_y_Rangos')
wr.sheet_properties.tabColor = BLUE
wr.sheet_view.showGridLines = False
for col, w in zip('ABCDEFG', [2, 34, 22, 22, 22, 16, 44]):
    wr.column_dimensions[col].width = w
wr['B1'] = 'RANGOS BAJO / MEDIO / ALTO POR VARIABLE'
wr['B1'].font = font(15, True, NAVY)
link(wr, 'B2', '1_RESUMEN', '◄ Volver al RESUMEN')
wb_src = openpyxl.load_workbook(SRC, data_only=True)
p = {row[0].value: row[2].value for row in wb_src['Parametros'].iter_rows(min_row=5, max_row=60) if row[0].value}
money = lambda v: f'${v:,.0f}'
money99 = lambda v: f'${v - 0.01:,.2f}'
cnt_lvl = lambda col, lv: cnt(x[col] == lv)
VARS = [
    ('Antigüedad como asociado', 'NIVEL_ANTIGUEDAD', f'{p["ANT_MEDIO_MAX"] + 1} años o más', f'{p["ANT_ALTO_MAX"] + 1} a {p["ANT_MEDIO_MAX"]} años',
     f'0 a {p["ANT_ALTO_MAX"]} año', ['Sin dato', 'Inconsistente - revisar'], 'Más antigüedad = menos riesgo.'),
    ('Ingreso mensual', 'NIVEL_INGRESO', f'{money(p["ING_L2"])} o más', f'{money(p["ING_L1"])} a {money99(p["ING_L2"])}',
     f'Menos de {money(p["ING_L1"])}', ['Sin información'], 'Más ingreso = menos riesgo.'),
    ('Deuda en la Central de Riesgo (SSF)', 'NIVEL_DEUDA', f'Menos de {money(p["DEU_L1"])}', f'{money(p["DEU_L1"])} a {money99(p["DEU_L2"])}',
     f'{money(p["DEU_L2"])} o más', ['Sin deuda', 'Sin información'], 'Deuda alta no significa que no pague.'),
    ('Deuda ÷ ingreso mensual', 'NIVEL_DEUDA_INGRESO', f'Menos de {p["DTI_L1"]:g} veces', f'{p["DTI_L1"]:g} a menos de {p["DTI_L2"]:g} veces',
     f'{p["DTI_L2"]:g} veces o más', ['Sin deuda', 'Sin información'], 'Cuántos meses de ingreso equivale la deuda.'),
    ('Cuotas de aportación vencidas', 'NIVEL_MORA_APORTACIONES', '0 (al día)', f'{p["MORA_L1"]} a {p["MORA_L2"] - 1}',
     f'{p["MORA_L2"]} o más', [], 'Atraso en aportaciones con la Cooperativa.'),
    ('Categoría interna', 'NIVEL_CAT_INTERNA', 'A1', 'A2', 'B, C, D o E', ['Sin información'], 'Sin información = sin historial interno.'),
    ('Categoría SSF (externa)', 'NIVEL_CAT_EXTERNA', 'A1', 'A2', 'B, C, D o E', ['Sin información'], ''),
]
r = 4
for j, h in enumerate(['Variable (nivel de RIESGO)', 'BAJO', 'MEDIO', 'ALTO', 'Sin dato / sin deuda', 'Nota'], 2):
    cellw(wr, r, j, h, bold=True, fill=NAVY, color='FFFFFF', align='center')
r += 1
for name, col, lo, me, hi, other, note in VARS:
    cellw(wr, r, 2, name, bold=True)
    for j, (rng, lv) in enumerate([(lo, 'Bajo'), (me, 'Medio'), (hi, 'Alto')], 3):
        cellw(wr, r, j, f'{rng}\n{cnt_lvl(col, lv):,} asociados', fill=LEVEL_FILLS[lv], align='center')
    cellw(wr, r, 6, f'{sum(cnt_lvl(col, o) for o in other):,}' if other else '0', align='center')
    cellw(wr, r, 7, note, size=9)
    wr.row_dimensions[r].height = 34
    r += 1
r += 1
for j, h in enumerate(['Variable (nivel de RESPALDO, no es riesgo)', 'BAJO', 'MEDIO', 'ALTO', 'Sin fondo', 'Nota'], 2):
    cellw(wr, r, j, h, bold=True, fill=GRAY, color='FFFFFF', align='center')
r += 1
for name, col, lo, me, hi, other, note in [
        ('Fondo compensatorio', 'NIVEL_RESPALDO_FONDO', f'Menos de {money(p["FON_L1"])}', f'{money(p["FON_L1"])} a {money99(p["FON_L2"])}',
         f'{money(p["FON_L2"])} o más', ['Sin fondo'], 'No se sabe si el fondo está disponible para pagar.'),
        ('Aportaciones acumuladas', 'NIVEL_APORTACIONES', f'Menos de {money(p["APO_L1"])}', f'{money(p["APO_L1"])} a {money99(p["APO_L2"])}',
         f'{money(p["APO_L2"])} o más', [], 'Muestra el vínculo con la Cooperativa.')]:
    cellw(wr, r, 2, name, bold=True)
    for j, (rng, lv) in enumerate([(lo, 'Bajo'), (me, 'Medio'), (hi, 'Alto')], 3):
        cellw(wr, r, j, f'{rng}\n{cnt_lvl(col, lv):,} asociados', fill='DDEBF7', align='center')
    cellw(wr, r, 6, f'{sum(cnt_lvl(col, o) for o in other):,}' if other else '0', align='center')
    cellw(wr, r, 7, note, size=9)
    wr.row_dimensions[r].height = 34
    r += 1
r += 1
wr.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
c = wr.cell(r, 2, 'Rangos propuestos, pendientes de aprobación de Riesgos. Se calcularon con la distribución de esta base '
                  '(cortes cercanos a tercios). No son límites de la normativa.')
c.font = font(9, italic=True, color=RED)
c.alignment = Alignment(wrap_text=True)
wr.row_dimensions[r].height = 28

# ================================================================ HOJAS DE LISTAS
BASE_COLS = [('N° ASOCIADO', 'ASOCIADO', '@', 11), ('NOMBRE', 'NOMBRE', None, 36), ('DUI', 'DUI', '@', 11),
             ('AGENCIA', 'AGENCIA', None, 16), ('INGRESO MENSUAL', 'INGRESO', USD, 13),
             ('FONDO COMPENSATORIO', 'FONDO_COMPENSATORIO', USD, 14), ('DEUDA SSF', 'SALDO_SSF', USD, 14)]
LVL_COLS = [('RIESGO POR ANTIGÜEDAD', 'NIVEL_ANTIGUEDAD', None, 11), ('RIESGO POR INGRESO', 'NIVEL_INGRESO', None, 11),
            ('RIESGO POR DEUDA', 'NIVEL_DEUDA', None, 11), ('RIESGO POR DEUDA vs SUELDO', 'NIVEL_DEUDA_INGRESO', None, 12)]
PEND = [('PENDIENTE ANTES DE OTORGAR', 'PENDIENTE_ANTES_DE_OTORGAR', None, 44)]


def list_sheet(name, color, title, note, df, cols):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = color
    ws.sheet_view.showGridLines = False
    ws['A1'] = f'{title} — {len(df)} asociados'
    ws['A1'].font = font(14, True, NAVY)
    ws['A2'] = note
    ws['A2'].font = font(9, italic=True, color=RED)
    link(ws, 'A3', '1_RESUMEN', '◄ Volver al RESUMEN')
    for j, (h, src, fmt, w) in enumerate(cols, 1):
        c = ws.cell(5, j, h)
        c.font = font(9, True, 'FFFFFF')
        c.fill = PatternFill('solid', fgColor=color)
        c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        c.border = box
        ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[5].height = 32
    df = df.sort_values('ASOCIADO', key=lambda s: s.astype(int))
    for i, (_, row) in enumerate(df.iterrows()):
        rr = 6 + i
        for j, (h, src, fmt, w) in enumerate(cols, 1):
            v = row[src]
            if isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(rr, j, v)
            c.font = font(9, bold=(j == 1))
            c.border = box
            if fmt:
                c.number_format = fmt
            if src.startswith('NIVEL_') and src != 'NIVEL_SEGURIDAD' or src == 'RECOMENDADO':
                c.alignment = Alignment(horizontal='center', vertical='center')
            elif src in ('MOTIVO_NO_MAXIMA', 'PENDIENTE_ANTES_DE_OTORGAR', 'TABLA_USADA'):
                c.alignment = Alignment(wrap_text=True, vertical='center')
            if i % 2:
                c.fill = PatternFill('solid', fgColor='F7F9FC')
    end = 5 + len(df)
    for j, (h, src, fmt, w) in enumerate(cols, 1):
        if src.startswith('NIVEL_') and src != 'NIVEL_SEGURIDAD':
            for lv, fc in LEVEL_FILLS.items():
                ws.conditional_formatting.add(f'{L(j)}6:{L(j)}{end}', CellIsRule(
                    operator='equal', formula=[f'"{lv}"'], fill=PatternFill('solid', fgColor=fc, bgColor=fc)))
    for j, (h, src, fmt, w) in enumerate(cols, 1):
        if src == 'RECOMENDADO':
            for txt, fc in (('SÍ', 'C6E0B4'), ('NO', 'F8CBAD')):
                ws.conditional_formatting.add(f'{L(j)}6:{L(j)}{end}', CellIsRule(
                    operator='equal', formula=[f'"{txt}"'], fill=PatternFill('solid', fgColor=fc, bgColor=fc),
                    font=Font(name=F, bold=True)))
    ws.freeze_panes = 'C6'
    ws.auto_filter.ref = f'A5:{L(len(cols))}{end}'
    # totales de montos
    rr = end + 2
    monto = [j for j, (h, src, fmt, w) in enumerate(cols, 1) if src in ('PRECA HIPO', 'PRECA VEHICULO', 'PRECA TC')]
    if monto:
        ws.cell(rr, monto[0] - 1, 'Total:').font = font(10, True)
    for j in monto:
        c = ws.cell(rr, j, float(df[cols[j - 1][1]].sum()))
        c.number_format = INT
        c.font = font(10, True)
    return ws


M_H = ('HIPOTECARIO', 'PRECA HIPO', INT, 13)
M_V = ('VEHÍCULO', 'PRECA VEHICULO', INT, 12)
M_T = ('TARJETA', 'PRECA TC', INT, 11)
EXT = 'Falta: lista negra, créditos castigados, relacionados/empleados y solvencia en Hacienda. No es aprobación de crédito.'
list_sheet('3_Maxima_Seguridad', GREEN, 'NIVEL 4 – MÁXIMA SEGURIDAD (RECOMENDADO)',
           'Cumplen el memorando, sin alertas, sin niveles Alto, sin errores de datos y con A1 interna y externa. ' + EXT,
           x[n4], BASE_COLS + [M_H, M_V, M_T] + LVL_COLS + PEND)
list_sheet('4_Hipotecario', GREEN, 'HIPOTECARIO – NIVEL 4', 'Todos requieren solvencia en Hacienda. ' + EXT,
           x[n4 & x['PRECA HIPO'].notna()], BASE_COLS + [M_H] + LVL_COLS + PEND)
list_sheet('5_Vehiculo', GREEN, 'VEHÍCULO – NIVEL 4', EXT, x[n4 & x['PRECA VEHICULO'].notna()], BASE_COLS + [M_V] + LVL_COLS + PEND)
list_sheet('6_Tarjeta', GREEN, 'TARJETA DE CRÉDITO – NIVEL 4',
           'No tienen tarjeta actualmente. Límite > $4,000 sin solvencia en Hacienda: se otorga con $4,000. ' + EXT,
           x[n4 & tc_ok & x['PRECA TC'].notna()],
           BASE_COLS + [M_T, ('SEGMENTO', 'SEGMENTO', None, 20)] + LVL_COLS + PEND)
x['NIVEL_SEGURIDAD_TXT'] = x.NIVEL_SEGURIDAD.map({k: f'{k}. {nm}' for k, nm, *_ in LEVELS})
x['RECOMENDADO'] = x.NIVEL_SEGURIDAD.ge(4).map({True: 'SÍ', False: 'NO'})
ws7 = list_sheet('7_Los_528', ORANGE, 'LOS 528 QUE CUMPLEN EL MEMORANDO: ¿SE RECOMIENDAN?',
                 'SÍ = se puede ofrecer con seguridad. NO = cumple el memorando, pero la columna "¿POR QUÉ NO?" dice qué le falta. ' + EXT,
                 x[n1], BASE_COLS[:4] + [('¿RECOMENDADO?', 'RECOMENDADO', None, 13),
                                         ('¿POR QUÉ NO?', 'MOTIVO_NO_MAXIMA', None, 70),
                                         ('TABLA CON LA QUE SE COMPARÓ (ver hoja 8)', 'TABLA_USADA', None, 28),
                                         M_H, M_V, M_T])

# ================================================================ HOJA 8: TABLAS USADAS
w8 = wb.create_sheet('8_Tablas_usadas')
w8.sheet_properties.tabColor = GRAY
w8.sheet_view.showGridLines = False
for col, w in zip('ABCDEF', [2, 40, 30, 30, 34, 2]):
    w8.column_dimensions[col].width = w
w8['B1'] = 'TABLAS CON LAS QUE SE COMPARÓ A CADA ASOCIADO'
w8['B1'].font = font(15, True, NAVY)
link(w8, 'B2', '1_RESUMEN', '◄ Volver al RESUMEN')
w8['B3'] = ('La hoja 7 dice, para cada asociado, con qué tabla se comparó. Aquí están las tablas completas. '
            'Las tablas 1, 2 y 3 son de SU archivo; la 4 y la 5 son propuestas hechas con sus datos.')
w8['B3'].font = font(10, True)
w8['B3'].alignment = Alignment(wrap_text=True)
w8.merge_cells('B3:E3')
w8.row_dimensions[3].height = 30
r = 5
noc = lambda col: cnt(n1 & ~n4 & (x[col] == 'No coincide con tabla'))


T_ROW = {}  # fila del título de cada tabla en la hoja 8 (para los enlaces)


def table_title(r, num, title, source, color, nota):
    T_ROW[num] = r
    w8.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    c = w8.cell(r, 2, f'TABLA {num}: {title}')
    c.font = font(12, True, 'FFFFFF')
    c.fill = PatternFill('solid', fgColor=color)
    w8.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=5)
    c = w8.cell(r + 1, 2, f'De dónde sale: {source}.  {nota}')
    c.font = font(9, italic=True)
    c.alignment = Alignment(wrap_text=True)
    w8.row_dimensions[r + 1].height = 26
    return r + 2


def table_rows(r, head, rows, fills=None):
    for j, h in enumerate(head, 2):
        cellw(w8, r, j, h, bold=True, fill='D9E1F2', align='center')
    r += 1
    for i, row in enumerate(rows):
        for j, v in enumerate(row, 2):
            fill = fills[i] if fills and j == 1 + len(row) else None
            cellw(w8, r, j, v, align='center' if j > 2 else None, fill=fill)
        r += 1
    return r + 1


OWN, PROP = '2E75B6', 'BF8F00'
NA = 'F8CBAD'
r = table_title(r, 1, 'VEHÍCULO (monto según ingreso)', 'su hoja "RANGOS VEHICULO" (copiada tal cual)', OWN,
                f'{cnt(n1 & ~n4 & (x.CONSIST_VEHICULO == "No coincide con tabla") & (x.INGRESO < 3000))} de los NO ganan menos de $3,000 '
                f'y aun así la base les dio vehículo (esta tabla dice NO APLICA). '
                f'{cnt(n1 & ~n4 & (x.CONSIST_VEHICULO == "No coincide con tabla") & (x.INGRESO >= 3000))} tiene $150,000, monto que no existe en la tabla.')
r = table_rows(r, ['Ingreso mensual', 'Monto de vehículo'],
               [('$15,000 o más', '$40,000'), ('$10,000 a $14,999.99', '$30,000'), ('$5,000 a $9,999.99', '$25,000'),
                ('$3,000 a $4,999.99', '$20,000'), ('Menos de $3,000', 'NO APLICA')], [None] * 4 + [NA])
r = table_title(r, 2, 'TARJETA DE CRÉDITO (límite según ingreso)', 'su hoja "RANGOS TC" (copiada tal cual)', OWN,
                'Ningún asociado de los 528 tiene un límite de tarjeta distinto al de esta tabla.')
r = table_rows(r, ['Ingreso mensual', 'Límite', 'Segmento'],
               [('$500 a $749.99', '$500', 'Clásica local'), ('$750 a $999.99', '$1,500', 'Clásica internacional'),
                ('$1,000 a $1,499.99', '$2,000', 'Clásica internacional'), ('$1,500 a $1,699.99', '$2,500', 'Dorada'),
                ('$1,700 a $1,999.99', '$3,000', 'Dorada'), ('$2,000 a $2,499.99', '$3,500', 'Platinum'),
                ('$2,500 a $2,999.99', '$4,500', 'Platinum'), ('$3,000 a $3,499.99', '$5,000', 'Platinum'),
                ('$3,500 a $7,499.99', '$7,000', 'Black'), ('$7,500 a $7,999.99', '$8,000', 'Black'),
                ('$8,000 o más', '$10,000', 'Black')])
r = table_title(r, 3, 'ANTIGÜEDAD COMO ASOCIADO', 'su hoja "Variables" (la que usted pidió)', OWN,
                f'{cnt(n1 & ~n4 & (x.NIVEL_ANTIGUEDAD == "Alto"))} de los NO son asociados nuevos (0 a 1 año).')
r = table_rows(r, ['Años como asociado', 'Nivel de riesgo'],
               [('0 a 1', 'Alto'), ('2 a 4', 'Medio'), ('5 o más', 'Bajo')], [NA, 'FFE699', 'C6E0B4'])
r = table_title(r, 4, 'HIPOTECARIO (monto según ingreso)',
                'PROPUESTA. Su archivo NO trae tabla de hipotecario; esta se armó viendo cómo se asignaron los montos en su base',
                PROP, f'{noc("CONSIST_HIPO")} de los NO ganan menos de $3,000 y aun así la base les dio hipotecario. Créditos debe dar la tabla oficial.')
r = table_rows(r, ['Ingreso mensual', 'Monto hipotecario'],
               [('$9,000 o más', '$205,000'), ('$5,000 a $8,999.99', '$150,000'), ('$3,000 a $4,999.99', '$110,000'),
                ('Menos de $3,000', 'NO APLICA')], [None] * 3 + [NA])
# evidencia: como se asignaron los montos hipotecarios en la base
hh = x[x['PRECA HIPO'].notna() & (num('INGRESO') > 0)]
tramos = [('Menos de $3,000', 0, 3000), ('$3,000 a $4,999.99', 3000, 5000), ('$5,000 a $8,999.99', 5000, 9000), ('$9,000 o más', 9000, 1e12)]
w8.merge_cells(start_row=r - 1, start_column=2, end_row=r - 1, end_column=5)
c = w8.cell(r - 1, 2, f'¿Cómo se armó? Se contó en SU base a los {len(hh)} asociados con hipotecario e ingreso registrado:')
c.font = font(10, True, NAVY)
ev = []
for lab, lo, hi in tramos:
    m = (hh.INGRESO >= lo) & (hh.INGRESO < hi)
    ev.append((lab,) + tuple(f'{int(((hh["PRECA HIPO"] == amt) & m).sum()):,}' for amt in (110000, 150000, 205000)))
r = table_rows(r, ['Ingreso mensual', 'Les dieron $110,000', 'Les dieron $150,000', 'Les dieron $205,000'], ev)
sin_ing = cnt(x['PRECA HIPO'].notna() & ~(num('INGRESO') > 0))
w8.merge_cells(start_row=r - 1, start_column=2, end_row=r - 1, end_column=5)
c = w8.cell(r - 1, 2, f'Se ve una regla clara: cada tramo de ingreso recibió un solo monto. En toda la base, solo {ev[0][1]} + {ev[0][2]} + {ev[0][3]} asociados '
                      f'con menos de $3,000 se salen de la regla ({noc("CONSIST_HIPO")} de ellos están entre los 528). Otros {sin_ing} con hipotecario no tienen ingreso registrado (vacío o $0).')
c.font = font(9, italic=True)
c.alignment = Alignment(wrap_text=True)
w8.row_dimensions[r - 1].height = 30
r += 1
r = table_title(r, 5, 'LÍMITES DE RIESGO', 'PROPUESTA hecha con los datos de su base (Riesgos debe aprobarla o cambiarla)', PROP,
                'Si el asociado cae en cualquier fila de esta tabla, se marca NO recomendado.')
r = table_rows(r, ['Se marca NO si…', 'Dato de su base', 'Por qué ese límite'],
               [('Gana menos de $1,500 al mes', 'INGRESO', 'Es el corte de $1,500 de su tabla de tarjeta'),
                ('Debe $10,000 o más', 'SALDO_SSF', 'Es el tercio de asociados que más debe'),
                ('Debe 12 meses o más de su sueldo', 'SALDO_SSF ÷ INGRESO', 'Debe un año de sueldo o más'),
                ('Tiene fondo menor de $4,000 y debe $10,000 o más', 'FONDO_COMPENSATORIO y SALDO_SSF', 'Poco respaldo para lo que debe'),
                ('Está atrasado en aportaciones', 'APORTACIONES_AL_DIA', 'El memorando pide ponerse al día'),
                ('Su categoría es A2 (no A1)', 'MAX_CAT_RIESGO y CALIFICACION_SSF', 'Para máxima seguridad se pide A1')])
r = table_title(r, 6, 'REVISIÓN DE DATOS', 'errores encontrados en su propia base', GRAY,
                'No es riesgo del asociado: es un dato mal anotado. Si Créditos lo corrige, el asociado puede pasar a SÍ.')
r = table_rows(r, ['Se marca NO si…', 'Dato de su base'],
               [('La edad no coincide con la fecha de nacimiento', 'EDAD y CLI_FECHA_NACIMIENTO'),
                ('El DUI no tiene 9 dígitos', 'DUI'),
                ('El ingreso es muy alto y hay que confirmarlo', 'INGRESO'),
                ('Los años de asociado no cuadran con la edad', 'ANTIGUEDAD_ANIOS y EDAD'),
                ('La deuda aparece en negativo', 'SALDO_SSF')])

# ================================================================ HOJA 9: CÓMO EXPLICARLO (ejemplos paso a paso)
w9 = wb.create_sheet('9_Como_explicarlo')
w9.sheet_properties.tabColor = NAVY
w9.sheet_view.showGridLines = False
w9.column_dimensions['A'].width = 2
w9.column_dimensions['B'].width = 16
w9.column_dimensions['C'].width = 110
w9['B1'] = 'CÓMO EXPLICAR POR QUÉ UN ASOCIADO ES "NO RECOMENDADO" (ejemplos reales, paso a paso)'
w9['B1'].font = font(15, True, NAVY)
link(w9, 'B2', '1_RESUMEN', '◄ Volver al RESUMEN')
r = 4


EX = []  # (fila del paso 1, fila de "Cuántos", asociado) de cada ejemplo


def ejemplo(r, titulo, pasos, resultado, frase, total, ok=False):
    EX.append([r + 1, r + len(pasos) + 3, re.search(r'(\d+) –', pasos[0]).group(1)])
    w9.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    c = w9.cell(r, 2, titulo)
    c.font = font(12, True, 'FFFFFF')
    c.fill = PatternFill('solid', fgColor=NAVY)
    r += 1
    for i, p in enumerate(pasos, 1):
        cellw(w9, r, 2, f'Paso {i}', bold=True, fill='D9E1F2', align='center')
        cellw(w9, r, 3, p)
        r += 1
    rf = 'C6E0B4' if ok else 'F8CBAD'
    cellw(w9, r, 2, 'Resultado', bold=True, fill=rf, align='center')
    cellw(w9, r, 3, resultado, bold=True, fill=rf)
    r += 1
    cellw(w9, r, 2, 'Cómo decirlo', bold=True, fill='E2EFDA', align='center')
    cellw(w9, r, 3, frase, fill='E2EFDA')
    r += 1
    cellw(w9, r, 2, 'Cuántos', bold=True, fill='F2F2F2', align='center')
    cellw(w9, r, 3, total, fill='F2F2F2')
    return r + 2


nb = n1 & ~n4
e = x[nb & (x.CONSIST_VEHICULO == 'No coincide con tabla') & (x.INGRESO < 3000)].iloc[0]
r = ejemplo(r, 'EJEMPLO 1 – VEHÍCULO: "Le dieron vehículo, pero la tabla dice NO APLICA"', [
    f'Busque al asociado en su base: {e.ASOCIADO} – {e.NOMBRE}.',
    f'Mire cuánto gana (columna INGRESO): ${e.INGRESO:,.2f} al mes.',
    'Busque ese ingreso en la Tabla 1 (su hoja RANGOS VEHICULO): la última fila dice "Menor a $3,000 → NO APLICA".',
    f'Mire cuánto le asignaron en la base (columna PRECA VEHICULO): ${e["PRECA VEHICULO"]:,.0f}.'],
    f'Su propia tabla dice que no le corresponde vehículo, pero la base le dio ${e["PRECA VEHICULO"]:,.0f}. No cuadra.',
    '"Según nuestra tabla de vehículo, con lo que gana no le corresponde ese crédito. Hay que confirmar con Créditos '
    'por qué se le asignó."',
    f'{cnt(nb & (x.CONSIST_VEHICULO == "No coincide con tabla") & (x.INGRESO < 3000))} asociados de los 528 están en este caso.')
e = x[nb & (x.CONSIST_HIPO == 'No coincide con tabla')].iloc[0]
r = ejemplo(r, 'EJEMPLO 2 – HIPOTECARIO: "Le dieron hipotecario con ingreso menor de $3,000"', [
    f'Busque al asociado: {e.ASOCIADO} – {e.NOMBRE}.',
    f'Mire cuánto gana: ${e.INGRESO:,.2f} al mes.',
    'Su archivo no trae tabla de hipotecario. Pero en toda su base, a quienes se les dio hipotecario con un ingreso '
    'registrado, ganan $3,000 o más (Tabla 4), salvo los casos como este.',
    f'Mire cuánto le asignaron (columna PRECA HIPO): ${e["PRECA HIPO"]:,.0f}.'],
    f'Gana menos de $3,000 y le dieron ${e["PRECA HIPO"]:,.0f} de hipotecario: es distinto a como se asignó al resto.',
    '"Su ingreso está por debajo del que tienen los demás asociados con hipotecario. Hay que confirmar la regla con Créditos."',
    f'{cnt(nb & (x.CONSIST_HIPO == "No coincide con tabla"))} asociados de los 528 están en este caso.')
e = x[nb & (x.NIVEL_DEUDA == 'Alto') & (x.NIVEL_INGRESO != 'Alto') & (x.CONSIST_VEHICULO != 'No coincide con tabla')].iloc[0]
r = ejemplo(r, 'EJEMPLO 3 – DEUDA: "Debe mucho en bancos y otras entidades"', [
    f'Busque al asociado: {e.ASOCIADO} – {e.NOMBRE}.',
    f'Mire cuánto debe en la Central de Riesgo (columna SALDO_SSF): ${e.SALDO_SSF:,.2f} con {e.INSTITUCION}.',
    'Compare con la Tabla 5 (límites de riesgo): deber $10,000 o más se marca como riesgo Alto.'],
    'Debe $10,000 o más: no se recomienda para máxima seguridad.',
    '"Ya tiene una deuda alta; prestarle más aumenta el riesgo."',
    f'{cnt(nb & (x.NIVEL_DEUDA == "Alto"))} asociados de los 528 deben $10,000 o más.')
e = x[nb & (x.AL_APORT_VENCIDAS == 1)].iloc[0]
r = ejemplo(r, 'EJEMPLO 4 – APORTACIONES: "Está atrasado en sus aportaciones"', [
    f'Busque al asociado: {e.ASOCIADO} – {e.NOMBRE}.',
    f'Mire la columna APORTACIONES_AL_DIA: dice "{e.APORTACIONES_AL_DIA}", y CUOTAS_VENC_APORT: {int(e.CUOTAS_VENC_APORT)} cuota(s) vencidas.',
    'El memorando dice que quien no está al día debe ponerse al día antes de darle el producto.'],
    'Está atrasado con la Cooperativa: no se recomienda hasta que se ponga al día.',
    '"Primero debe ponerse al día con sus aportaciones."',
    f'{cnt(nb & (x.AL_APORT_VENCIDAS == 1))} asociados de los 528 están atrasados.')
e = x[n4 & x['PRECA VEHICULO'].notna()].iloc[0]
prods_e = [f'{t}: ${e[col]:,.0f} (correcto según {tb})' for t, col, tb in
           [('hipotecario', 'PRECA HIPO', 'la Tabla 4'), ('vehículo', 'PRECA VEHICULO', 'la Tabla 1'), ('tarjeta', 'PRECA TC', 'la Tabla 2')]
           if pd.notna(e[col])]
r = ejemplo(r, 'EJEMPLO 5 – UN "SÍ": por qué sí se recomienda', [
    f'Asociado: {e.ASOCIADO} – {e.NOMBRE}.',
    f'Gana ${e.INGRESO:,.2f} al mes; debe ${e.SALDO_SSF:,.2f}; tiene ${e.FONDO_COMPENSATORIO:,.2f} de fondo; está al día en aportaciones; '
    f'categoría A1.',
    'Lo que le dieron: ' + '; '.join(prods_e) + '. Ningún dato cae en la Tabla 5 ni en la Tabla 6.'],
    'Pasa todo: es recomendado (nivel 4, máxima seguridad).',
    '"Cumple el memorando, no tiene ninguna señal de riesgo y sus datos están correctos."',
    f'{cnt(n4)} asociados están en este caso.', ok=True)
for row in w9.iter_rows(min_row=4, max_row=r):
    for c in row:
        if c.column == 3:
            c.alignment = Alignment(wrap_text=True, vertical='center')

# ================================================================ HOJA 10: ASOCIADOS POR TABLA (destino de los enlaces)
w10 = wb.create_sheet('10_Asociados_por_tabla')
w10.sheet_properties.tabColor = GRAY
w10.sheet_view.showGridLines = False
for col, w in zip('ABCDEFGHI', [2, 12, 38, 16, 18, 18, 18, 18, 30]):
    w10.column_dimensions[col].width = w
w10['B1'] = 'ASOCIADOS "NO RECOMENDADOS", AGRUPADOS POR LA TABLA QUE LOS DEJÓ FUERA'
w10['B1'].font = font(15, True, NAVY)
link(w10, 'B2', '1_RESUMEN', '◄ Volver al RESUMEN')
w10['B3'] = ('Cada bloque explica qué se revisó y muestra el dato de cada asociado. Haga clic en el N° de asociado para ver su fila '
             'completa en la hoja 7. Un asociado puede aparecer en más de un bloque.')
w10['B3'].font = font(10, italic=True)
ROW7 = {a: 6 + i for i, a in enumerate(sorted(x.loc[n1, 'ASOCIADO'].astype(str), key=int))}
SEC = {}
r = 5
nb = n1 & ~n4
money0 = lambda v: '' if pd.isna(v) else f'${v:,.0f}'
money2 = lambda v: '' if pd.isna(v) else f'${v:,.2f}'
to_txt = lambda v: v if isinstance(v, str) else money0(v)


def section(r, key, title, explica, table_num, mask, cols):
    """cols: [(encabezado, funcion(fila)->valor)]"""
    d = x[mask].sort_values('ASOCIADO', key=lambda s_: s_.astype(int))
    SEC[key] = (r, len(d))
    w10.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
    c = w10.cell(r, 2, f'{title} — {len(d)} asociados')
    c.font = font(12, True, 'FFFFFF')
    c.fill = PatternFill('solid', fgColor=NAVY)
    w10.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=8)
    c = w10.cell(r + 1, 2, explica)
    c.font = font(9)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    w10.row_dimensions[r + 1].height = 30
    c = w10.cell(r + 1, 9, f'◄ Ver la Tabla {table_num}')
    c.hyperlink = f"#'8_Tablas_usadas'!B{T_ROW[table_num]}"
    c.font = Font(name=F, size=10, bold=True, color='0563C1', underline='single')
    r += 2
    heads = ['N° ASOCIADO', 'NOMBRE', 'AGENCIA'] + [h for h, _ in cols]
    for j, h in enumerate(heads, 2):
        cellw(w10, r, j, h, bold=True, fill='D9E1F2', align='center', size=9)
    r += 1
    for _, e in d.iterrows():
        a = str(e.ASOCIADO)
        c = cellw(w10, r, 2, a, size=9, bold=True)
        c.hyperlink = f"#'7_Los_528'!A{ROW7[a]}"
        c.font = Font(name=F, size=9, bold=True, color='0563C1', underline='single')
        cellw(w10, r, 3, e.NOMBRE, size=9)
        cellw(w10, r, 4, e.AGENCIA, size=9)
        for j, (h, fn) in enumerate(cols, 5):
            cellw(w10, r, j, fn(e), size=9, align='center')
        r += 1
    return r + 2


r = section(r, 'T1', 'TABLA 1 – VEHÍCULO: le dieron vehículo pero no cuadra con su tabla',
            'Qué se revisó: se buscó lo que gana cada asociado en su tabla RANGOS VEHICULO y se comparó con el vehículo que le asignó '
            'la base. Si gana menos de $3,000, la tabla dice NO APLICA.', 1, nb & (x.CONSIST_VEHICULO == 'No coincide con tabla'),
            [('GANA AL MES', lambda e: money2(e.INGRESO)), ('LE DIERON DE VEHÍCULO', lambda e: money0(e['PRECA VEHICULO'])),
             ('SEGÚN LA TABLA 1 LE TOCA', lambda e: to_txt(e.ESPERADO_VEHICULO)), ('¿CUADRA?', lambda e: 'NO')])
r = section(r, 'T3', 'TABLA 3 – ANTIGÜEDAD: asociados nuevos (0 a 1 año)',
            'Qué se revisó: los años como asociado (columna ANTIGUEDAD_ANIOS). Según su hoja Variables, 0 a 1 año es riesgo Alto.', 3,
            nb & (x.NIVEL_ANTIGUEDAD == 'Alto'),
            [('AÑOS DE ASOCIADO', lambda e: int(e.ANTIGUEDAD_ANIOS)), ('NIVEL', lambda e: 'Alto')])
r = section(r, 'T4', 'TABLA 4 – HIPOTECARIO: le dieron hipotecario con ingreso menor de $3,000',
            'Qué se revisó: lo que gana cada asociado contra la regla que siguen los demás hipotecarios de su base '
            '(nadie más recibió hipotecario ganando menos de $3,000).', 4, nb & (x.CONSIST_HIPO == 'No coincide con tabla'),
            [('GANA AL MES', lambda e: money2(e.INGRESO)), ('LE DIERON DE HIPOTECARIO', lambda e: money0(e['PRECA HIPO'])),
             ('SEGÚN LA TABLA 4 LE TOCA', lambda e: to_txt(e.ESPERADO_HIPO)), ('¿CUADRA?', lambda e: 'NO')])
r = section(r, 'T5a', 'TABLA 5 – GANA POCO (menos de $1,500 al mes)', 'Qué se revisó: la columna INGRESO.', 5,
            nb & (x.NIVEL_INGRESO == 'Alto'), [('GANA AL MES', lambda e: money2(e.INGRESO))])
r = section(r, 'T5b', 'TABLA 5 – DEBE MUCHO ($10,000 o más)',
            'Qué se revisó: lo que debe en bancos y otras entidades según la Central de Riesgo (columna SALDO_SSF).', 5,
            nb & (x.NIVEL_DEUDA == 'Alto'), [('DEBE', lambda e: money2(e.SALDO_SSF)), ('INSTITUCIÓN', lambda e: e.INSTITUCION)])
r = section(r, 'T5c', 'TABLA 5 – DEBE 12 MESES O MÁS DE SU SUELDO',
            'Qué se revisó: se dividió lo que debe entre lo que gana al mes. 12 o más quiere decir que debe un año de sueldo o más.', 5,
            nb & (x.NIVEL_DEUDA_INGRESO == 'Alto'),
            [('DEBE', lambda e: money0(e.SALDO_SSF)), ('GANA AL MES', lambda e: money0(e.INGRESO)),
             ('MESES DE SUELDO', lambda e: f'{e.MULT_DEUDA_INGRESO:,.0f}')])
r = section(r, 'T5d', 'TABLA 5 – POCO FONDO PARA LO QUE DEBE (fondo menor de $4,000 y deuda de $10,000 o más)',
            'Qué se revisó: el fondo compensatorio contra lo que debe.', 5, nb & (x.AL_FONDO_BAJO_DEUDA_ELEVADA == 1),
            [('FONDO COMPENSATORIO', lambda e: money2(e.FONDO_COMPENSATORIO)), ('DEBE', lambda e: money0(e.SALDO_SSF))])
r = section(r, 'T5e', 'TABLA 5 – ATRASADO EN APORTACIONES',
            'Qué se revisó: las columnas APORTACIONES_AL_DIA y CUOTAS_VENC_APORT. El memorando pide que se ponga al día.', 5,
            nb & (x.AL_APORT_VENCIDAS == 1),
            [('¿AL DÍA?', lambda e: e.APORTACIONES_AL_DIA), ('CUOTAS VENCIDAS', lambda e: int(e.CUOTAS_VENC_APORT))])
r = section(r, 'T5f', 'TABLA 5 – CATEGORÍA A2 (no A1)', 'Qué se revisó: la categoría interna y la de la Central de Riesgo.', 5,
            nb & (x.AL_CATEGORIA_A2 == 1),
            [('CATEGORÍA INTERNA', lambda e: '' if pd.isna(e.MAX_CAT_RIESGO) else e.MAX_CAT_RIESGO),
             ('CATEGORÍA SSF', lambda e: '' if pd.isna(e.CALIFICACION_SSF) else e.CALIFICACION_SSF)])


def dato_malo(e):
    out = []
    if e.CHK_EDAD != 'Coincide':
        out.append(f'Edad anotada {int(e.EDAD)}, por su fecha de nacimiento ({e.CLI_FECHA_NACIMIENTO}) tiene {int(e.EDAD_RECALCULADA)}')
    if len(str(e.DUI)) != 9:
        out.append(f'DUI "{e.DUI}" no tiene 9 dígitos')
    if e.ALERTA_ING_ATIPICO == 1:
        out.append(f'Ingreso muy alto: ${e.INGRESO:,.0f}')
    if e.ALERTA_ANT_EDAD == 1:
        out.append(f'{int(e.ANTIGUEDAD_ANIOS)} años de asociado con {int(e.EDAD)} años de edad')
    if str(e.ESTADO_DEUDA_SSF).startswith('Saldo negativo'):
        out.append('Deuda en negativo')
    return '; '.join(out)


w10.column_dimensions['E'].width = 24
r = section(r, 'T6', 'TABLA 6 – DATOS MAL ANOTADOS EN LA BASE',
            'No es riesgo del asociado: es un dato que hay que corregir o confirmar con Créditos. Si se corrige, puede pasar a SÍ.', 6,
            nb & (x.TOTAL_ALERTAS_DATOS > 0) & (x.MOTIVO_NO_MAXIMA.str.contains('edad anotada|DUI|ingreso es muy alto|no cuadran|negativo', case=False)),
            [('QUÉ DATO ESTÁ MAL', dato_malo)])
# la hoja 7 se ordena por N° de asociado: sus filas sirven de destino para los enlaces
w7 = wb['7_Los_528']
hdr7 = {w7.cell(5, j).value: j for j in range(1, w7.max_column + 1)}
tcol = hdr7['TABLA CON LA QUE SE COMPARÓ (ver hoja 8)']
for rr in range(6, w7.max_row + 1):
    c = w7.cell(rr, tcol)
    mm = re.search(r'Tabla (\d)', str(c.value or ''))
    if mm:
        c.hyperlink = f"#'8_Tablas_usadas'!B{T_ROW[int(mm.group(1))]}"
        c.font = Font(name=F, size=9, color='0563C1', underline='single')

# ---- enlaces desde la hoja 8 hacia las listas
w8.column_dimensions['F'].width = 30


def link_cell(ws_, cell, target, text):
    c = ws_[cell]
    c.value = text
    c.hyperlink = target
    c.font = Font(name=F, size=10, bold=True, color='0563C1', underline='single')


SEC_LINK = lambda key: f"#'10_Asociados_por_tabla'!B{SEC[key][0]}"
for num, key in ((1, 'T1'), (3, 'T3'), (4, 'T4'), (6, 'T6')):
    link_cell(w8, f'F{T_ROW[num]}', SEC_LINK(key), f'▶ Ver los {SEC[key][1]} asociados')
link_cell(w8, f'F{T_ROW[5]}', SEC_LINK('T5a'), '▶ Ver los asociados (por fila)')
for i, key in enumerate(['T5a', 'T5b', 'T5c', 'T5d', 'T5e', 'T5f']):
    link_cell(w8, f'F{T_ROW[5] + 3 + i}', SEC_LINK(key), f'▶ Ver los {SEC[key][1]}')
w8.cell(T_ROW[2], 6, 'Nadie: 0 asociados').font = font(10, True, GRAY)

# ---- enlaces desde la hoja 9 (ejemplos)
for (paso1, cuantos, aso), key in zip(EX, ['T1', 'T4', 'T5b', 'T5e', None]):
    c = w9.cell(paso1, 3)
    c.hyperlink = f"#'7_Los_528'!A{ROW7[aso]}"
    c.font = Font(name=F, size=10, color='0563C1', underline='single')
    c = w9.cell(cuantos, 3)
    if key:
        c.value = f'{c.value}  ▶ Clic aquí para ver la lista.'
        c.hyperlink = SEC_LINK(key)
    else:
        c.value = f'{c.value}  ▶ Clic aquí para ver la lista.'
        c.hyperlink = "#'3_Maxima_Seguridad'!A1"
    c.font = Font(name=F, size=10, bold=True, color='0563C1', underline='single')

# ---- enlaces desde el RESUMEN (niveles)
for k, target in ((1, "#'7_Los_528'!A1"), (4, "#'3_Maxima_Seguridad'!A1"), (5, "#'3_Maxima_Seguridad'!A1")):
    c = ws.cell(LEVEL_ROW[k], 3)
    c.hyperlink = target
    c.font = Font(name=F, size=10, bold=True, color='0563C1', underline='single')

# ================================================================ BD FINAL original (copia sin cambios)
BD_SRC = sys.argv[3] if len(sys.argv) > 3 else None
if BD_SRC:
    src_ws = openpyxl.load_workbook(BD_SRC)['BD FINAL']
    wbd = wb.create_sheet('BD FINAL')
    wbd.sheet_properties.tabColor = '404040'
    for row in src_ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            d = wbd.cell(c.row, c.column, c.value)
            d.number_format = c.number_format
            d.font = font(9, bold=(c.row == 1), color='FFFFFF' if c.row == 1 else '000000')
            if c.row == 1:
                d.fill = PatternFill('solid', fgColor='404040')
    for j in range(1, src_ws.max_column + 1):
        wbd.column_dimensions[L(j)].width = 14
    wbd.freeze_panes = 'C2'
    wbd.auto_filter.ref = f'A1:{L(src_ws.max_column)}{src_ws.max_row}'

wb.active = 0
wb.save(OUT)
print({f'nivel {k}': cnt(m) for k, _, _, m in LEVELS})
print('hojas:', wb.sheetnames)
