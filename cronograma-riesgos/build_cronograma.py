# -*- coding: utf-8 -*-
"""
Genera el libro Excel «Cronograma_Dashboard_Riesgos.xlsx».

Uso:
    python build_cronograma.py [ruta_salida.xlsx]

El script también expone build(...) con parámetros opcionales de prueba
(actividades temporales, semana actual, fecha de referencia) que usa
test_cronograma.py. El archivo entregable se genera SIN datos de prueba.

Las fórmulas se escriben con nombres de función en inglés porque así las
almacena el formato .xlsx; Excel en español las muestra traducidas
(LOOKUP = BUSCAR, IFERROR = SI.ERROR, SUMPRODUCT = SUMAPRODUCTO, etc.).
"""
import re
import sys
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.formatting.rule import FormulaRule, DataBarRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo, TableFormula
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

# --------------------------------------------------------------------------
# Paleta y estilos
# --------------------------------------------------------------------------
NAVY = "1F3864"       # azul oscuro: encabezados principales
MID = "2E75B6"        # azul medio: indicadores
LIGHT = "DDEBF7"      # azul muy claro: apoyo
INPUT = "FFF7D6"      # amarillo pálido: celdas de entrada
CALC = "F2F2F2"       # gris muy claro: celdas calculadas (protegidas)
GREEN_F, GREEN_T = "C6EFCE", "006100"
DKGREEN = "1E7B34"
AMBER_F, AMBER_T = "FFE699", "7F6000"
ORANGE_F = "F8CBAD"
RED_F, RED_T = "FFC7CE", "9C0006"
BLUE_F, BLUE_T = "BDD7EE", "1F3864"
TXT = "1A1A1A"

FONT = "Arial"


def font(size=10, bold=False, color=TXT, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)

PUESTOS = [
    "Riesgo Operacional",
    "Riesgo Legal y Normativo",
    "Riesgo Financiero",
    "Asistente Técnico de Riesgo",
]
ESTADOS = ["Pendiente", "En curso", "Completado"]
SEMANAS = ["Semana 1", "Semana 2", "Semana 3", "Semana 4"]

S_DASH = "Dashboard - Consolidado"
S_CRON = "Cronograma y Actividades"
S_HITO = "Linea de Tiempo - Hitos"
S_CONF = "Configuración"
S_INST = "Instrucciones"
S_CALC = "Calculos"

FIRST = 5            # primera fila de datos de las tablas
N_ACT_ROWS = 100     # filas preformateadas en tblActividades (25 por puesto)
LAST_RNG = 1000      # límite de los rangos con nombre que leen el Dashboard

COLS_ACT = [
    "ID Actividad", "Área / Puesto", "Macro Actividad", "Responsable",
    "Fecha de inicio", "Fecha límite", "Peso de actividad (%)",
    "Avance Semana 1 (%)", "Avance Semana 2 (%)", "Avance Semana 3 (%)",
    "Avance Semana 4 (%)", "% de Avance Total Actual", "Estado",
    "Última actualización", "Evidencia / Observaciones", "Validación de datos",
]
COLS_HITO = [
    "ID Hito", "Semana", "Fecha objetivo", "Área / Puesto", "Entregable / Hito",
    "Criterio de aceptación", "Responsable", "Estado",
    "Fecha real de cumplimiento", "Evidencia / Enlace", "Observaciones",
]

# Propuesta inicial de hitos (editable, sujeta a validación de responsables)
HITOS = [
    # (código, semana, puesto, entregable, criterio)
    ("RO", 1, 0, "Diagnóstico inicial: inventario de procesos críticos y revisión de la matriz de riesgo operacional vigente",
     "Documento con el listado de procesos críticos, fuentes consultadas y brechas identificadas, revisado por el titular del puesto."),
    ("RO", 2, 0, "Borrador de matriz de riesgo operacional actualizada (eventos, probabilidad, impacto y controles)",
     "Matriz con proceso, evento de riesgo, probabilidad, impacto, control y dueño para cada proceso crítico del inventario."),
    ("RO", 3, 0, "Validación de controles con dueños de proceso y propuesta de planes de mitigación",
     "Planes de mitigación con responsable y fecha para todos los riesgos calificados como altos; validación documentada (acta o correo)."),
    ("RO", 4, 0, "Informe final de riesgo operacional y presentación ejecutiva",
     "Informe y presentación entregados en la reunión de cierre y con conformidad registrada de la jefatura."),
    ("RLN", 1, 1, "Inventario de normativa aplicable y diagnóstico preliminar de cumplimiento",
     "Listado de normas aplicables con fuente, vigencia y estado preliminar de cumplimiento por requisito."),
    ("RLN", 2, 1, "Matriz de cumplimiento normativo y de riesgo legal documentada",
     "Matriz con requisito, área responsable, evidencia de cumplimiento y nivel de riesgo para cada norma del inventario."),
    ("RLN", 3, 1, "Análisis de brechas normativas y propuestas de adecuación",
     "Cada brecha identificada tiene propuesta de adecuación, responsable sugerido y plazo; revisión del titular documentada."),
    ("RLN", 4, 1, "Informe de cumplimiento normativo y presentación ejecutiva",
     "Informe y presentación entregados en la reunión de cierre y con conformidad registrada de la jefatura."),
    ("RF", 1, 2, "Levantamiento de información financiera y definición de indicadores de riesgo (liquidez, mercado, crédito)",
     "Fuentes de datos identificadas y ficha de cada indicador con fórmula, periodicidad, fuente y responsable."),
    ("RF", 2, 2, "Cálculo y documentación de indicadores y límites de riesgo financiero",
     "Indicadores calculados con datos de corte definidos, memoria de cálculo y comparación contra límites vigentes."),
    ("RF", 3, 2, "Análisis de sensibilidad / escenarios y propuesta de límites o alertas tempranas",
     "Al menos un escenario base y uno adverso documentados, con propuesta de umbrales de alerta justificada."),
    ("RF", 4, 2, "Informe de riesgo financiero y presentación ejecutiva",
     "Informe y presentación entregados en la reunión de cierre y con conformidad registrada de la jefatura."),
    ("ATR", 1, 3, "Plan de trabajo consolidado del área y estructura del repositorio de evidencias",
     "Actividades de los cuatro puestos registradas en el cronograma con pesos que suman 100% y repositorio con carpeta por puesto."),
    ("ATR", 2, 3, "Organización de bases de datos y documentación de soporte; primer reporte de seguimiento",
     "Evidencias de la Semana 1 y 2 enlazadas en el cronograma y reporte de avance semanal emitido."),
    ("ATR", 3, 3, "Control de calidad de entregables y consolidación de la información de los tres puestos",
     "Lista de verificación de calidad aplicada a cada entregable; Dashboard sin alertas de datos pendientes de corrección."),
    ("ATR", 4, 3, "Consolidación del informe final del área y soporte a la presentación ejecutiva",
     "Informe consolidado y anexos entregados; todos los hitos con evidencia enlazada y estado actualizado."),
]


# --------------------------------------------------------------------------
# Utilidades
# --------------------------------------------------------------------------
def q(sheet):
    return "'" + sheet.replace("'", "''") + "'"


def style_range(ws, ref, **kw):
    for row in ws[ref]:
        for c in row:
            for k, v in kw.items():
                setattr(c, k, v)


def merge_set(ws, ref, value=None, **kw):
    ws.merge_cells(ref)
    first = ws[ref.split(":")[0]]
    if value is not None:
        first.value = value
    style_range(ws, ref, **kw)
    return first


def add_name(wb, name, ref):
    dn = DefinedName(name, attr_text=ref)
    wb.defined_names[name] = dn


# --------------------------------------------------------------------------
# Fórmulas de la tabla de actividades (fila r)
# --------------------------------------------------------------------------
def reg(r):
    """Fila registrada: hay algún dato de entrada."""
    return f"(COUNTA($A{r}:$K{r})+COUNTA($N{r}:$O{r}))>0"


def f_avance(r):
    return (f'=IF(NOT({reg(r)}),"",'
            f'IFERROR(LOOKUP(2,1/(H{r}:K{r}<>""),H{r}:K{r}),0))')


def f_estado(r):
    return (f'=IF(L{r}="","",IF(L{r}>=1,"Completado",'
            f'IF(L{r}>0,"En curso","Pendiente")))')


def checks(r):
    A, B, C, D, E, F, G = (f"{c}{r}" for c in "ABCDEFG")
    H, I, J, K, L, O = (f"{c}{r}" for c in "HIJKLO")
    sum_w = f"SUMIFS($G${FIRST}:$G${LAST_RNG},$B${FIRST}:$B${LAST_RNG},{B})"
    return [
        (f'{A}=""', "Falta ID"),
        (f'AND({A}<>"",COUNTIF($A${FIRST}:$A${LAST_RNG},{A})>1)', "ID duplicado"),
        (f'{B}=""', "Área sin seleccionar"),
        (f'AND({B}<>"",ISNA(MATCH({B},ListaPuestos,0)))', "Área no válida"),
        (f'{C}=""', "Macroactividad vacía"),
        (f'{D}=""', "Sin responsable"),
        (f'{E}=""', "Sin fecha de inicio"),
        (f'{F}=""', "Sin fecha límite"),
        (f'OR(AND({E}<>"",NOT(ISNUMBER({E}))),AND({F}<>"",NOT(ISNUMBER({F}))))', "Fecha no válida"),
        (f'AND(ISNUMBER({E}),ISNUMBER({F}),{E}>{F})', "Inicio posterior a fecha límite"),
        (f'{G}=""', "Peso vacío"),
        (f'AND({G}<>"",OR(NOT(ISNUMBER({G})),{G}<=0,{G}>1))', "Peso fuera de rango"),
        (f'AND({B}<>"",ABS({sum_w}-1)>0.0001)',
         f'"Pesos del puesto suman "&TEXT({sum_w},"0%")'),
        (f'(COUNTIF({H}:{K},">1")+COUNTIF({H}:{K},"<0")+COUNTA({H}:{K})-COUNT({H}:{K}))>0',
         "Avance fuera de rango 0-100%"),
        (f'OR(AND({I}<>"",{I}<{H}),AND({J}<>"",{J}<MAX({H}:{I})),AND({K}<>"",{K}<MAX({H}:{J})))',
         "Avance semanal decreciente"),
        (f'AND(ISNUMBER({F}),{F}<FechaRef,N({L})<1)', "Vencida"),
        (f'AND(EvidenciaObligatoria="Sí",N({L})>=1,{O}="")', "Completada sin evidencia"),
    ]


def f_validacion(r):
    parts = []
    for cond, msg in checks(r):
        m = msg if msg.startswith('"') else f'"{msg}"'
        parts.append(f'IF({cond},{m}&" · ","")')
    s = "&".join(parts)
    return (f'=IF(NOT({reg(r)}),"",IF(({s})="","✔ OK",'
            f'"⚠ "&LEFT({s},LEN({s})-3)))')


# Traducción a referencias estructuradas para la definición de columna
# calculada de la tabla (permite que Excel extienda las fórmulas al agregar filas).
def to_structured(formula, r=FIRST, table="tblActividades"):
    f = formula.lstrip("=")

    def colname(letter):
        return COLS_ACT[ord(letter) - ord("A")]

    # Rangos de la misma fila: $A5:$K5 / H5:K5
    def rng(m):
        c1, c2 = m.group(1), m.group(2)
        return f"{table}[[#This Row],[{colname(c1)}]:[{colname(c2)}]]"
    f = re.sub(rf"\$?([A-P]){r}:\$?([A-P]){r}(?!\d)", rng, f)

    # Celdas sueltas de la fila (no absolutas en fila)
    def cell(m):
        return f"{table}[[#This Row],[{colname(m.group(2))}]]"
    f = re.sub(rf"(?<![\$A-Za-z0-9_])(\$?)([A-P]){r}(?!\d)", cell, f)
    return f


# --------------------------------------------------------------------------
# Construcción
# --------------------------------------------------------------------------
def build(path, test_rows=None, semana=1, fecha_ref=None, extra_rows=0,
          inicio_ciclo=dt.date(2026, 10, 12), evidencia="Sí"):
    wb = Workbook()
    ws_d = wb.active
    ws_d.title = S_DASH
    ws_c = wb.create_sheet(S_CRON)
    ws_h = wb.create_sheet(S_HITO)
    ws_i = wb.create_sheet(S_INST)
    ws_k = wb.create_sheet(S_CONF)
    ws_x = wb.create_sheet(S_CALC)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90

    build_config(wb, ws_k, semana, fecha_ref, inicio_ciclo, evidencia)
    last_act = build_cronograma(wb, ws_c, test_rows or [], extra_rows)
    build_hitos(wb, ws_h)
    build_calculos(wb, ws_x)
    build_dashboard(wb, ws_d)
    build_instrucciones(ws_i)

    ws_x.sheet_state = "hidden"
    wb.active = 0
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(path)
    return path


# --------------------------------------------------------------------------
def build_config(wb, ws, semana, fecha_ref, inicio_ciclo, evidencia):
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 70

    merge_set(ws, "B1:D1", "Configuración y listas del libro",
              font=font(14, True, "FFFFFF"), fill=fill(NAVY), alignment=LEFT)
    ws.row_dimensions[1].height = 28
    ws["B2"] = "Celdas amarillas = parámetros editables. El resto se usa en listas desplegables y fórmulas; no cambie su posición."
    ws["B2"].font = font(9, italic=True)

    hdr = ["Parámetro", "Valor", "Uso / nota"]
    for i, h in enumerate(hdr):
        c = ws.cell(row=3, column=2 + i, value=h)
        c.font = font(10, True, "FFFFFF"); c.fill = fill(MID); c.alignment = CENTER; c.border = BORDER

    params = [
        ("Semana actual del ciclo (1 a 4)", semana,
         "Controla los mensajes de finalización anticipada y qué semanas se muestran en el seguimiento semanal.", "0"),
        ("Fecha de referencia (corte)", fecha_ref if fecha_ref else "=TODAY()",
         "Por defecto =HOY(). Puede reemplazarla por una fecha fija para emitir un reporte de corte. Determina actividades vencidas.", "dd/mm/yyyy"),
        ("Fecha de inicio del ciclo (lunes de la Semana 1)", inicio_ciclo,
         "VALOR INICIAL EDITABLE (supuesto). Calcula las fechas objetivo propuestas de los hitos (viernes de cada semana).", "dd/mm/yyyy"),
        ("Días de anticipación para alerta «próxima a vencer»", 3,
         "Actividades no completadas cuya fecha límite cae dentro de este número de días se señalan en ámbar.", "0"),
        ("¿Evidencia obligatoria para actividades completadas?", evidencia,
         "Si es «Sí», una actividad al 100% sin texto en Evidencia / Observaciones genera alerta.", "@"),
        ("Meta final del proyecto", 1, "Meta de cumplimiento al cierre de la Semana 4.", "0%"),
    ]
    names = ["SemanaActual", "FechaRef", "InicioCiclo", "DiasAlerta", "EvidenciaObligatoria", "MetaFinal"]
    for i, (lbl, val, note, fmt) in enumerate(params):
        r = 4 + i
        ws.cell(row=r, column=2, value=lbl).font = font(10, True)
        v = ws.cell(row=r, column=3, value=val)
        v.number_format = fmt; v.font = font(10, True, NAVY); v.fill = fill(INPUT)
        v.alignment = CENTER; v.protection = Protection(locked=False)
        ws.cell(row=r, column=4, value=note).font = font(9)
        for col in range(2, 5):
            ws.cell(row=r, column=col).border = BORDER
            if col != 3:
                ws.cell(row=r, column=col).alignment = LEFT
        ws.row_dimensions[r].height = 30
        add_name(wb, names[i], f"{q(S_CONF)}!$C${r}")

    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="4",
                        showErrorMessage=True, errorTitle="Semana no válida",
                        error="Ingrese un número entero entre 1 y 4.")
    dv.add("C4"); ws.add_data_validation(dv)
    dvd = DataValidation(type="date", operator="greaterThan", formula1="36526",
                         showErrorMessage=True, error="Ingrese una fecha válida (dd/mm/aaaa).")
    dvd.add("C5"); dvd.add("C6"); ws.add_data_validation(dvd)
    dvn = DataValidation(type="whole", operator="between", formula1="0", formula2="30",
                         showErrorMessage=True, error="Ingrese un número de días entre 0 y 30.")
    dvn.add("C7"); ws.add_data_validation(dvn)
    dvs = DataValidation(type="list", formula1="ListaSiNo", showErrorMessage=True)
    dvs.add("C8"); ws.add_data_validation(dvs)
    ws["C9"].protection = Protection(locked=True)
    ws["C9"].fill = fill(CALC)

    # Plan lineal de referencia
    r0 = 11
    merge_set(ws, f"B{r0}:D{r0}", "Plan de referencia (avance esperado acumulado por semana) – supuesto editable",
              font=font(10, True, "FFFFFF"), fill=fill(MID), alignment=LEFT)
    for i in range(4):
        r = r0 + 1 + i
        ws.cell(row=r, column=2, value=f"Semana {i + 1}").font = font(10, True)
        v = ws.cell(row=r, column=3, value=(i + 1) * 0.25)
        v.number_format = "0%"; v.fill = fill(INPUT); v.font = font(10, True, NAVY)
        v.alignment = CENTER; v.protection = Protection(locked=False)
        ws.cell(row=r, column=4, value="Referencia lineal solo para comparación visual; no afecta el cálculo del avance real.").font = font(9)
        for col in range(2, 5):
            ws.cell(row=r, column=col).border = BORDER
    add_name(wb, "PlanRef", f"{q(S_CONF)}!$C${r0 + 1}:$C${r0 + 4}")

    # Listas
    r1 = 17
    merge_set(ws, f"B{r1}:D{r1}", "Listas desplegables (no editar sin actualizar las fórmulas)",
              font=font(10, True, "FFFFFF"), fill=fill(NAVY), alignment=LEFT)
    lists = [("Áreas / Puestos", PUESTOS, "ListaPuestos"),
             ("Estados", ESTADOS, "ListaEstados"),
             ("Semanas", SEMANAS, "ListaSemanas"),
             ("Sí / No", ["Sí", "No"], "ListaSiNo")]
    r = r1 + 1
    for title, items, nm in lists:
        ws.cell(row=r, column=2, value=title).font = font(10, True, MID)
        r += 1
        start = r
        for it in items:
            c = ws.cell(row=r, column=2, value=it)
            c.font = font(10); c.border = BORDER
            r += 1
        add_name(wb, nm, f"{q(S_CONF)}!$B${start}:$B${r - 1}")
        r += 1

    # Metodología
    r += 1
    merge_set(ws, f"B{r}:D{r}", "Metodología de cálculo (documentación)",
              font=font(10, True, "FFFFFF"), fill=fill(NAVY), alignment=LEFT)
    notes = [
        "1. Avance de una actividad = último avance semanal acumulado registrado (Semana 4 → 1), ignorando semanas vacías. Los avances semanales NO se suman.",
        "2. Estado automático: 0% = Pendiente; >0% y <100% = En curso; 100% = Completado.",
        "3. Avance ponderado del puesto = Σ (peso de la actividad × avance actual). Solo es válido si el puesto tiene ≥1 actividad, ningún peso vacío o fuera de 0–100% y los pesos suman 100% (tolerancia ±0,01%).",
        "4. Avance general del proyecto = promedio simple de los 4 puestos (cada puesto pesa 25%, sin importar cuántas actividades tenga). Si algún puesto no es válido, el Dashboard muestra «No definitivo» y un valor provisional rotulado como tal.",
        "5. Avance semanal acumulado (Semana n) = Σ peso × último avance registrado hasta la Semana n. Incremento semanal = acumulado Semana n − acumulado Semana n−1.",
        "6. «¡Enhorabuena, completado anticipadamente!»: el puesto (o el proyecto) es válido, llega al 100% y la semana en que lo alcanzó (la menor entre la primera semana con 100% registrada y la Semana actual) es anterior a la Semana 4. Si lo alcanza en la Semana 4: «Completado». Si no: «En ejecución».",
        "7. Una semana se muestra en el seguimiento si es ≤ Semana actual o si ya tiene algún avance registrado.",
        "8. Los rangos que alimentan el Dashboard cubren las filas 5 a 1000 de la hoja Cronograma y Actividades.",
    ]
    for n in notes:
        r += 1
        merge_set(ws, f"B{r}:D{r}", n, font=font(9), alignment=LEFT_TOP)
        ws.row_dimensions[r].height = 30

    ws.freeze_panes = "A4"
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# --------------------------------------------------------------------------
def build_cronograma(wb, ws, test_rows, extra_rows):
    widths = [12, 27, 44, 22, 12, 12, 11, 11, 11, 11, 11, 12, 13, 13, 36, 58]
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w

    merge_set(ws, "A1:P1", "Cronograma y Actividades – Ciclo de 4 semanas del área de Riesgos",
              font=font(14, True, "FFFFFF"), fill=fill(NAVY), alignment=LEFT)
    ws.row_dimensions[1].height = 28
    merge_set(ws, "A2:P2",
              "Celdas amarillas: entrada de datos.  Celdas grises: cálculo automático (protegidas).  "
              "Avances semanales = % ACUMULADO al cierre de cada semana (deje vacía la semana aún no actualizada).  "
              "Pesos: % relativo dentro del puesto; deben sumar 100% por puesto.",
              font=font(9, False, NAVY), fill=fill(LIGHT), alignment=LEFT)
    ws.row_dimensions[2].height = 30
    ws["A3"] = "Entrada"; ws["A3"].fill = fill(INPUT); ws["A3"].font = font(8); ws["A3"].border = BORDER
    ws["B3"] = "Cálculo automático"; ws["B3"].fill = fill(CALC); ws["B3"].font = font(8); ws["B3"].border = BORDER

    hr = FIRST - 1
    calc_cols = {12, 13, 16}
    for i, h in enumerate(COLS_ACT):
        c = ws.cell(row=hr, column=i + 1, value=h)
        c.font = font(10, True, "FFFFFF")
        c.fill = fill(NAVY if (i + 1) not in calc_cols else MID)
        c.alignment = CENTER; c.border = BORDER
    ws.row_dimensions[hr].height = 42

    last = FIRST + N_ACT_ROWS - 1 + extra_rows
    for r in range(FIRST, last + 1):
        for col in range(1, 17):
            c = ws.cell(row=r, column=col)
            c.border = BORDER
            c.font = font(10)
            if col in calc_cols:
                c.fill = fill(CALC)
                c.protection = Protection(locked=True)
            else:
                c.fill = fill(INPUT)
                c.protection = Protection(locked=False)
            if col in (5, 6, 14):
                c.number_format = "dd/mm/yyyy"; c.alignment = CENTER
            elif col == 7:
                c.number_format = "0.0%"; c.alignment = CENTER
            elif col in (8, 9, 10, 11, 12):
                c.number_format = "0%"; c.alignment = CENTER
            elif col in (1, 13):
                c.alignment = CENTER
            else:
                c.alignment = LEFT
        ws.cell(row=r, column=12, value=f_avance(r))
        ws.cell(row=r, column=13, value=f_estado(r))
        ws.cell(row=r, column=16, value=f_validacion(r))
        ws.cell(row=r, column=16).font = font(9)

    # Datos de prueba (solo en ejecuciones de test)
    for i, row in enumerate(test_rows):
        r = row.get("_row", FIRST + i)
        for key, col in [("id", 1), ("area", 2), ("macro", 3), ("resp", 4), ("ini", 5),
                         ("fin", 6), ("peso", 7), ("s1", 8), ("s2", 9), ("s3", 10),
                         ("s4", 11), ("act", 14), ("evid", 15)]:
            if key in row and row[key] is not None:
                ws.cell(row=r, column=col, value=row[key])

    # Tabla estructurada
    tab = Table(displayName="tblActividades", ref=f"A{hr}:P{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    tab._initialise_columns()
    for i, col in enumerate(tab.tableColumns):
        col.name = COLS_ACT[i]
    for idx, fn in [(11, f_avance), (12, f_estado), (15, f_validacion)]:
        tab.tableColumns[idx].calculatedColumnFormula = TableFormula(
            attr_text=to_structured(fn(FIRST)))
    ws.add_table(tab)

    # Nombres de rango para Dashboard / Cálculos
    for nm, col in [("rID", "A"), ("rArea", "B"), ("rResp", "D"), ("rFIni", "E"), ("rFLim", "F"),
                    ("rPeso", "G"), ("rS1", "H"), ("rS2", "I"), ("rS3", "J"), ("rS4", "K"),
                    ("rAvance", "L"), ("rEstado", "M"), ("rEvid", "O"), ("rValid", "P")]:
        add_name(wb, nm, f"{q(S_CRON)}!${col}${FIRST}:${col}${LAST_RNG}")

    # Validaciones de datos
    V = f"{FIRST}:{LAST_RNG}"
    dv_area = DataValidation(type="list", formula1="ListaPuestos", allow_blank=True,
                             showErrorMessage=True, errorTitle="Área no válida",
                             error="Seleccione una de las cuatro áreas de la lista.",
                             showInputMessage=True, promptTitle="Área / Puesto",
                             prompt="Seleccione el puesto responsable de la actividad.")
    dv_area.add(f"B{FIRST}:B{LAST_RNG}")
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                             showErrorMessage=True, errorTitle="Fecha no válida",
                             error="Ingrese una fecha válida (dd/mm/aaaa).")
    for col in "EFN":
        dv_date.add(f"{col}{FIRST}:{col}{LAST_RNG}")
    dv_peso = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                             allow_blank=True, showErrorMessage=True, errorTitle="Peso no válido",
                             error="Ingrese un porcentaje entre 0% y 100% (por ejemplo 25%).",
                             showInputMessage=True, promptTitle="Peso de actividad",
                             prompt="Peso relativo dentro del puesto. La suma por puesto debe ser 100%.")
    dv_peso.add(f"G{FIRST}:G{LAST_RNG}")
    dv_av = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                           allow_blank=True, showErrorMessage=True, errorTitle="Avance no válido",
                           error="Ingrese un porcentaje entre 0% y 100%. Deje la celda vacía si la semana no se ha actualizado.",
                           showInputMessage=True, promptTitle="Avance acumulado",
                           prompt="% ACUMULADO al cierre de la semana (no el incremento).")
    dv_av.add(f"H{FIRST}:K{LAST_RNG}")
    dv_est = DataValidation(type="list", formula1="ListaEstados", allow_blank=True,
                            showErrorMessage=True,
                            showInputMessage=True, promptTitle="Estado automático",
                            prompt="Se calcula a partir del % de avance. No se edita manualmente.")
    dv_est.add(f"M{FIRST}:M{LAST_RNG}")
    for d in (dv_area, dv_date, dv_peso, dv_av, dv_est):
        ws.add_data_validation(d)

    # Formato condicional
    R = lambda cols: f"{cols[0]}{FIRST}:{cols[1]}{LAST_RNG}"
    regf = f"(COUNTA($A{FIRST}:$K{FIRST})+COUNTA($N{FIRST}:$O{FIRST}))>0"
    red = dict(fill=fill(RED_F), font=Font(color=RED_T))
    amb = dict(fill=fill(AMBER_F), font=Font(color=AMBER_T))
    org = dict(fill=fill(ORANGE_F), font=Font(color="833C0B", bold=True))
    cf = ws.conditional_formatting
    # Estado
    cf.add(R("MM"), FormulaRule(formula=[f'$M{FIRST}="Completado"'], fill=fill(GREEN_F), font=Font(color=GREEN_T, bold=True)))
    cf.add(R("MM"), FormulaRule(formula=[f'$M{FIRST}="En curso"'], fill=fill(BLUE_F), font=Font(color=BLUE_T, bold=True)))
    cf.add(R("MM"), FormulaRule(formula=[f'$M{FIRST}="Pendiente"'], fill=fill(RED_F), font=Font(color=RED_T)))
    # Avance total
    cf.add(R("LL"), FormulaRule(formula=[f'AND(ISNUMBER($L{FIRST}),$L{FIRST}>=1)'], fill=fill(GREEN_F), font=Font(color=GREEN_T, bold=True)))
    cf.add(R("LL"), DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="5B9BD5", showValue=True))
    # Vencida / próxima a vencer (fecha límite)
    cf.add(R("FF"), FormulaRule(formula=[f'AND(ISNUMBER($F{FIRST}),$F{FIRST}<FechaRef,N($L{FIRST})<1)'], stopIfTrue=True, **org))
    cf.add(R("FF"), FormulaRule(formula=[f'AND(ISNUMBER($F{FIRST}),$F{FIRST}>=FechaRef,$F{FIRST}-FechaRef<=DiasAlerta,N($L{FIRST})<1)'], **amb))
    # Datos obligatorios faltantes
    cf.add(R("AG"), FormulaRule(formula=[f'AND({regf},A{FIRST}="")'], **red))
    # ID duplicado
    cf.add(R("AA"), FormulaRule(formula=[f'AND($A{FIRST}<>"",COUNTIF($A${FIRST}:$A${LAST_RNG},$A{FIRST})>1)'], **org))
    # Fechas incoherentes
    cf.add(R("EF"), FormulaRule(formula=[f'AND(ISNUMBER($E{FIRST}),ISNUMBER($F{FIRST}),$E{FIRST}>$F{FIRST})'], **org))
    # Ponderación incorrecta del puesto
    cf.add(R("GG"), FormulaRule(
        formula=[f'AND($B{FIRST}<>"",ABS(SUMIFS($G${FIRST}:$G${LAST_RNG},$B${FIRST}:$B${LAST_RNG},$B{FIRST})-1)>0.0001)'], **amb))
    # Porcentajes semanales fuera de rango / decrecientes
    cf.add(R("HK"), FormulaRule(formula=[f'AND(H{FIRST}<>"",OR(NOT(ISNUMBER(H{FIRST})),H{FIRST}>1,H{FIRST}<0))'], stopIfTrue=True, **red))
    cf.add(R("IK"), FormulaRule(formula=[f'AND(I{FIRST}<>"",I{FIRST}<MAX($H{FIRST}:H{FIRST}))'], **org))
    # Evidencia faltante en completadas
    cf.add(R("OO"), FormulaRule(formula=[f'AND(EvidenciaObligatoria="Sí",N($L{FIRST})>=1,$O{FIRST}="")'], **red))
    # Columna de validación
    cf.add(R("PP"), FormulaRule(formula=[f'LEFT($P{FIRST},1)="⚠"'], fill=fill(RED_F), font=Font(color=RED_T)))
    cf.add(R("PP"), FormulaRule(formula=[f'LEFT($P{FIRST},1)="✔"'], font=Font(color=GREEN_T, bold=True)))

    ws.freeze_panes = f"D{FIRST}"
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hr}:{hr}"
    ws.print_area = f"A1:P{last}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "Página &P de &N"
    return last


# --------------------------------------------------------------------------
def build_hitos(wb, ws):
    widths = [11, 11, 13, 26, 48, 52, 18, 13, 14, 30, 34]
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    merge_set(ws, "A1:K1", "Línea de Tiempo – Hitos y entregables (propuesta inicial editable)",
              font=font(14, True, "FFFFFF"), fill=fill(NAVY), alignment=LEFT)
    ws.row_dimensions[1].height = 28
    merge_set(ws, "A2:K2",
              "PROPUESTA INICIAL EDITABLE, sujeta a validación de los responsables de cada puesto. No constituye una "
              "obligación institucional aprobada. Semana 1: diagnóstico y planificación · Semana 2: desarrollo y documentación · "
              "Semana 3: validación, análisis y propuestas · Semana 4: cierre y presentación ejecutiva.",
              font=font(9, True, "7F6000"), fill=fill("FFF2CC"), alignment=LEFT)
    ws.row_dimensions[2].height = 32
    ws["A3"] = "Entrada"; ws["A3"].fill = fill(INPUT); ws["A3"].font = font(8); ws["A3"].border = BORDER
    ws["B3"] = "Fecha objetivo propuesta = viernes de la semana según «Fecha de inicio del ciclo» (Configuración); puede sobrescribirse."
    ws["B3"].font = font(8, italic=True)

    hr = FIRST - 1
    for i, h in enumerate(COLS_HITO):
        c = ws.cell(row=hr, column=i + 1, value=h)
        c.font = font(10, True, "FFFFFF"); c.fill = fill(NAVY); c.alignment = CENTER; c.border = BORDER
    ws.row_dimensions[hr].height = 34

    hitos_sorted = sorted(HITOS, key=lambda h: (h[1], h[2]))
    n_rows = 40
    last = FIRST + n_rows - 1
    for idx in range(n_rows):
        r = FIRST + idx
        for col in range(1, 12):
            c = ws.cell(row=r, column=col)
            c.border = BORDER; c.font = font(10); c.fill = fill(INPUT)
            c.protection = Protection(locked=False)
            c.alignment = CENTER if col in (1, 2, 3, 8, 9) else LEFT
            if col in (3, 9):
                c.number_format = "dd/mm/yyyy"
        if idx < len(hitos_sorted):
            code, sem, p, ent, crit = hitos_sorted[idx]
            ws.cell(row=r, column=1, value=f"H-{code}-{sem:02d}")
            ws.cell(row=r, column=2, value=f"Semana {sem}")
            ws.cell(row=r, column=3,
                    value=f'=IF(OR(InicioCiclo="",B{r}=""),"",InicioCiclo+7*VALUE(RIGHT(B{r},1))-3)')
            ws.cell(row=r, column=4, value=PUESTOS[p])
            ws.cell(row=r, column=5, value=ent)
            ws.cell(row=r, column=6, value=crit)
            ws.cell(row=r, column=7, value="Por asignar")
            ws.cell(row=r, column=8, value="Pendiente")
            ws.cell(row=r, column=11, value="Propuesta inicial editable – sujeta a validación del responsable.")
            ws.row_dimensions[r].height = 45

    tab = Table(displayName="tblHitos", ref=f"A{hr}:K{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(tab)

    for nm, col in [("hSemana", "B"), ("hFechaObj", "C"), ("hArea", "D"), ("hEstado", "H"),
                    ("hFechaReal", "I")]:
        add_name(wb, nm, f"{q(S_HITO)}!${col}${FIRST}:${col}$500")

    dv_s = DataValidation(type="list", formula1="ListaSemanas", allow_blank=True, showErrorMessage=True,
                          error="Seleccione Semana 1, Semana 2, Semana 3 o Semana 4.")
    dv_s.add(f"B{FIRST}:B500")
    dv_a = DataValidation(type="list", formula1="ListaPuestos", allow_blank=True, showErrorMessage=True,
                          error="Seleccione una de las cuatro áreas.")
    dv_a.add(f"D{FIRST}:D500")
    dv_e = DataValidation(type="list", formula1="ListaEstados", allow_blank=True, showErrorMessage=True,
                          error="Seleccione Pendiente, En curso o Completado.")
    dv_e.add(f"H{FIRST}:H500")
    dv_d = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                          showErrorMessage=True, error="Ingrese una fecha válida (dd/mm/aaaa).")
    dv_d.add(f"C{FIRST}:C500"); dv_d.add(f"I{FIRST}:I500")
    for d in (dv_s, dv_a, dv_e, dv_d):
        ws.add_data_validation(d)

    cf = ws.conditional_formatting
    rng = lambda c: f"{c}{FIRST}:{c}500"
    cf.add(rng("H"), FormulaRule(formula=[f'$H{FIRST}="Completado"'], fill=fill(GREEN_F), font=Font(color=GREEN_T, bold=True)))
    cf.add(rng("H"), FormulaRule(formula=[f'$H{FIRST}="En curso"'], fill=fill(BLUE_F), font=Font(color=BLUE_T, bold=True)))
    cf.add(rng("H"), FormulaRule(formula=[f'$H{FIRST}="Pendiente"'], fill=fill(RED_F), font=Font(color=RED_T)))
    cf.add(rng("C"), FormulaRule(formula=[f'AND(ISNUMBER($C{FIRST}),$C{FIRST}<FechaRef,$H{FIRST}<>"Completado")'],
                                 stopIfTrue=True, fill=fill(ORANGE_F), font=Font(color="833C0B", bold=True)))
    cf.add(rng("C"), FormulaRule(formula=[f'AND(ISNUMBER($C{FIRST}),$C{FIRST}>=FechaRef,$C{FIRST}-FechaRef<=DiasAlerta,$H{FIRST}<>"Completado")'],
                                 fill=fill(AMBER_F), font=Font(color=AMBER_T)))
    cf.add(rng("I"), FormulaRule(formula=[f'AND(ISNUMBER($I{FIRST}),ISNUMBER($C{FIRST}),$I{FIRST}>$C{FIRST})'],
                                 fill=fill(AMBER_F), font=Font(color=AMBER_T)))
    cf.add(rng("J"), FormulaRule(formula=[f'AND($H{FIRST}="Completado",$J{FIRST}="")'], fill=fill(RED_F), font=Font(color=RED_T)))

    ws.freeze_panes = f"C{FIRST}"
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hr}:{hr}"
    ws.print_area = f"A1:K{last}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "Página &P de &N"


# --------------------------------------------------------------------------
# Hoja oculta de cálculos por puesto
# --------------------------------------------------------------------------
V = {
    1: "rS1",
    2: "(rS2+(rS2=\"\")*rS1)",
}
V[3] = f"(rS3+(rS3=\"\")*{V[2]})"
V[4] = f"(rS4+(rS4=\"\")*{V[3]})"

CALC_HDR = ["Puesto", "N° actividades", "Suma de pesos", "Pesos vacíos o inválidos",
            "Configuración válida", "Avance ponderado actual",
            "Acum. Semana 1", "Acum. Semana 2", "Acum. Semana 3", "Acum. Semana 4",
            "Primera semana con 100%", "Completado (100%)", "Completado anticipadamente",
            "Completadas", "En curso", "Pendientes", "Vencidas", "Con alertas"]


def build_calculos(wb, ws):
    ws["A1"] = ("Hoja auxiliar OCULTA: cálculos intermedios por puesto que alimentan el Dashboard. "
                "No editar. Ver metodología en la hoja Configuración.")
    ws["A1"].font = font(9, True)
    for i, h in enumerate(CALC_HDR):
        c = ws.cell(row=3, column=i + 1, value=h)
        c.font = font(9, True, "FFFFFF"); c.fill = fill(NAVY); c.alignment = CENTER
        ws.column_dimensions[get_column_letter(i + 1)].width = 14
    ws.column_dimensions["A"].width = 28
    for i, p in enumerate(PUESTOS):
        r = 4 + i
        ws[f"A{r}"] = f"=INDEX(ListaPuestos,{i + 1})"
        ws[f"B{r}"] = f"=COUNTIFS(rArea,$A{r})"
        ws[f"C{r}"] = f"=SUMIFS(rPeso,rArea,$A{r})"
        ws[f"D{r}"] = (f'=COUNTIFS(rArea,$A{r},rPeso,"")+COUNTIFS(rArea,$A{r},rPeso,"<=0")'
                       f'+COUNTIFS(rArea,$A{r},rPeso,">1")')
        ws[f"E{r}"] = f"=AND(B{r}>0,D{r}=0,ABS(C{r}-1)<=0.0001)"
        ws[f"F{r}"] = f"=IFERROR(SUMPRODUCT(--(rArea=$A{r}),rPeso,rAvance),0)"
        for w in range(1, 5):
            col = get_column_letter(6 + w)
            ws[f"{col}{r}"] = f"=IFERROR(SUMPRODUCT((rArea=$A{r})*rPeso*{V[w]}),0)"
        ws[f"K{r}"] = (f'=IF(ROUND(G{r},4)>=1,1,IF(ROUND(H{r},4)>=1,2,'
                       f'IF(ROUND(I{r},4)>=1,3,IF(ROUND(J{r},4)>=1,4,""))))')
        ws[f"L{r}"] = f"=AND(E{r},ROUND(F{r},4)>=1)"
        ws[f"M{r}"] = f"=AND(L{r},MIN(IF(K{r}=\"\",4,K{r}),SemanaActual)<4)"
        ws[f"N{r}"] = f'=COUNTIFS(rArea,$A{r},rEstado,"Completado")'
        ws[f"O{r}"] = f'=COUNTIFS(rArea,$A{r},rEstado,"En curso")'
        ws[f"P{r}"] = f'=COUNTIFS(rArea,$A{r},rEstado,"Pendiente")'
        ws[f"Q{r}"] = f'=SUMPRODUCT((rArea=$A{r})*ISNUMBER(SEARCH("Vencida",rValid)))'
        ws[f"R{r}"] = f'=SUMPRODUCT((rArea=$A{r})*(LEFT(rValid,1)="⚠"))'
    r = 8
    ws[f"A{r}"] = "Proyecto (promedio de 4 puestos)"
    ws[f"B{r}"] = "=SUM(B4:B7)"
    ws[f"C{r}"] = '=COUNTIF(E4:E7,TRUE)'          # n° de puestos válidos
    ws[f"D{r}"] = '=IF(C8>0,SUMPRODUCT(F4:F7*(E4:E7=TRUE))/C8,"")'  # provisional
    ws[f"E{r}"] = "=AND(E4,E5,E6,E7)"
    ws[f"F{r}"] = '=IF(E8,AVERAGE(F4:F7),"")'
    for w in range(1, 5):
        col = get_column_letter(6 + w)
        ws[f"{col}{r}"] = f'=IF($E$8,AVERAGE({col}4:{col}7),"")'
    ws[f"K{r}"] = (f'=IF(NOT(E8),"",IF(ROUND(G8,4)>=1,1,IF(ROUND(H8,4)>=1,2,'
                   f'IF(ROUND(I8,4)>=1,3,IF(ROUND(J8,4)>=1,4,"")))))')
    ws[f"L{r}"] = "=AND(E8,ROUND(N(F8),4)>=1)"
    ws[f"M{r}"] = "=AND(L8,MIN(IF(K8=\"\",4,K8),SemanaActual)<4)"
    ws[f"N{r}"] = '=COUNTIF(rEstado,"Completado")'
    ws[f"O{r}"] = '=COUNTIF(rEstado,"En curso")'
    ws[f"P{r}"] = '=COUNTIF(rEstado,"Pendiente")'
    ws[f"Q{r}"] = '=SUMPRODUCT(--ISNUMBER(SEARCH("Vencida",rValid)))'
    ws[f"R{r}"] = '=SUMPRODUCT(--(LEFT(rValid,1)="⚠"))'
    ws["A9"] = "Fila 8: C = n° puestos válidos; D = avance provisional (promedio solo de puestos válidos)."
    ws["A9"].font = font(8, italic=True)

    ws["A10"] = "Semana visible en seguimiento"
    for w in range(1, 5):
        col = get_column_letter(6 + w)
        ws[f"{col}10"] = f"=OR({w}<=SemanaActual,COUNT(rS{w})>0)"
    ws["A11"] = "Total registros"
    ws["B11"] = '=SUMPRODUCT(--(LEN(rValid)>0))'
    ws["A12"] = "Puestos al 100%"
    ws["B12"] = "=COUNTIF(L4:L7,TRUE)"
    ws["A13"] = "Próximas a vencer"
    ws["B13"] = '=COUNTIFS(rFLim,">="&FechaRef,rFLim,"<="&(FechaRef+DiasAlerta),rAvance,"<1")'
    for rr in range(4, 9):
        for cc in "CDFGHIJ":
            ws[f"{cc}{rr}"].number_format = "0.0%"
    ws["C8"].number_format = "0"
    ws.protection.sheet = True


# --------------------------------------------------------------------------
# Dashboard
# --------------------------------------------------------------------------
def count_msgs(*msgs):
    parts = "+".join(f'ISNUMBER(SEARCH("{m}",rValid))' for m in msgs)
    return f"=SUMPRODUCT(--(({parts})>0))"


def build_dashboard(wb, ws):
    C = q(S_CALC) + "!"
    widths = {"A": 2, "B": 30, "C": 12, "D": 16, "E": 13, "F": 13, "G": 13, "H": 13,
              "I": 13, "J": 13, "K": 22, "L": 20, "M": 24, "N": 2}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v

    merge_set(ws, "B1:M1", "Dashboard Consolidado – Gestión de Riesgos · Ciclo de 4 semanas",
              font=font(16, True, "FFFFFF"), fill=fill(NAVY), alignment=LEFT)
    ws.row_dimensions[1].height = 34
    ws["B2"] = "Semana actual:"; ws["B2"].font = font(10, True, NAVY); ws["B2"].alignment = Alignment(horizontal="right")
    ws["C2"] = "=SemanaActual"; ws["C2"].font = font(11, True, NAVY); ws["C2"].alignment = CENTER
    ws["C2"].number_format = '0" de 4"'
    ws["D2"] = "Fecha de corte:"; ws["D2"].font = font(10, True, NAVY); ws["D2"].alignment = Alignment(horizontal="right")
    ws["E2"] = "=FechaRef"; ws["E2"].number_format = "dd/mm/yyyy"; ws["E2"].font = font(11, True, NAVY)
    ws["E2"].alignment = CENTER
    merge_set(ws, "F2:M2", "Se actualiza automáticamente desde «Cronograma y Actividades». Semana y fecha de corte: hoja Configuración.",
              font=font(9, italic=True, color=NAVY), alignment=LEFT)
    ws.row_dimensions[2].height = 22

    # --- Tarjetas KPI ---
    cards = [
        ("Avance general del proyecto", f'=IF({C}$E$8,{C}$F$8,"No definitivo")', "0%",
         f'=IF({C}$E$8,"Promedio simple de los 4 puestos",IF({C}$C$8=0,"Ningún puesto con datos válidos",'
         f'"Provisional "&TEXT({C}$D$8,"0%")&" ("&{C}$C$8&" de 4 puestos válidos)"))'),
        ("Meta de cumplimiento", "=MetaFinal", "0%", "Al cierre de la Semana 4"),
        ("Macroactividades registradas", f"={C}$B$11", "0", f'="Con área asignada: "&{C}$B$8'),
        ("Actividades completadas", f"={C}$N$8", "0",
         f'=IF({C}$B$11=0,"—",TEXT({C}$N$8/{C}$B$11,"0%")&" del total")'),
        ("Actividades en curso", f"={C}$O$8", "0", "Avance > 0% y < 100%"),
        ("Actividades pendientes", f"={C}$P$8", "0", f'="De todas las actividades, vencidas: "&{C}$Q$8'),
        ("Puestos que alcanzaron el 100%", f"={C}$B$12", '0" de 4"', "Con ponderación válida"),
        ("Actividades con alertas", f"={C}$R$8", "0", f'="Próximas a vencer: "&{C}$B$13'),
    ]
    groups = [("B", "D"), ("E", "G"), ("H", "J"), ("K", "M")]
    for idx, (lbl, val, fmt, sub) in enumerate(cards):
        g = groups[idx % 4]
        r = 4 if idx < 4 else 8
        a, b = g
        merge_set(ws, f"{a}{r}:{b}{r}", lbl, font=font(10, True, "FFFFFF"), fill=fill(MID),
                  alignment=CENTER, border=BORDER)
        v = merge_set(ws, f"{a}{r + 1}:{b}{r + 1}", val, font=font(24, True, NAVY), fill=fill("FFFFFF"),
                      alignment=CENTER, border=BORDER)
        v.number_format = fmt
        merge_set(ws, f"{a}{r + 2}:{b}{r + 2}", sub, font=font(9, False, NAVY), fill=fill(LIGHT),
                  alignment=CENTER, border=BORDER)
    for r in (4, 8):
        ws.row_dimensions[r].height = 20
        ws.row_dimensions[r + 1].height = 38
        ws.row_dimensions[r + 2].height = 18
    ws.row_dimensions[7].height = 8
    cf = ws.conditional_formatting
    cf.add("B5", FormulaRule(formula=['AND(ISNUMBER($B$5),$B$5>=1)'], fill=fill(GREEN_F), font=Font(color=GREEN_T, bold=True)))
    cf.add("B5", FormulaRule(formula=['NOT(ISNUMBER($B$5))'], fill=fill(AMBER_F), font=Font(color=AMBER_T, bold=True)))
    cf.add("K9", FormulaRule(formula=['$K$9>0'], fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    cf.add("K9", FormulaRule(formula=['$K$9=0'], fill=fill(GREEN_F), font=Font(color=GREEN_T, bold=True)))
    cf.add("H9", FormulaRule(formula=['$H$9=4'], fill=fill(GREEN_F), font=Font(color=GREEN_T, bold=True)))
    cf.add("E9", FormulaRule(formula=['$E$9>0'], fill=fill(GREEN_F), font=Font(color=GREEN_T, bold=True)))

    # Barra de avance general
    ws["B12"] = f'=IF({C}$E$8,"Avance general vs meta 100%","Avance PROVISIONAL vs meta")'
    ws["B12"].font = font(10, True, NAVY)
    merge_set(ws, "C12:F12", f'=IF({C}$E$8,{C}$F$8,N({C}$D$8))', font=font(10, True, NAVY), alignment=LEFT)
    ws["C12"].number_format = "0%"
    cf.add("C12:F12", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                  color="2E75B6", showValue=True))
    # Mensaje de proyecto
    proj_msg = (
        f'=IF({C}$B$11=0,"⚠ No hay actividades registradas. Registre actividades y pesos en «Cronograma y Actividades».",'
        f'IF(NOT({C}$E$8),"⚠ Avance global NO definitivo: "&(4-{C}$C$8)&" puesto(s) sin actividades válidas o con pesos que no suman 100%.",'
        f'IF({C}$M$8,"Proyecto: ¡Enhorabuena, completado anticipadamente!",'
        f'IF({C}$L$8,"Proyecto completado: meta del 100% alcanzada.",'
        f'"Proyecto en ejecución – Semana "&SemanaActual&" de 4."))))'
    )
    merge_set(ws, "G12:M12", proj_msg, font=font(11, True, NAVY), alignment=CENTER, border=BORDER)
    ws.row_dimensions[12].height = 30
    msg_cf(cf, "G12:M12", "$G$12")

    # --- Avance individual por puesto ---
    r = 14
    merge_set(ws, f"B{r}:M{r}", "Avance individual por puesto (ponderado por el peso de cada actividad)",
              font=font(12, True, "FFFFFF"), fill=fill(NAVY), alignment=LEFT)
    ws.row_dimensions[r].height = 24
    hdrs = ["Puesto", "Avance ponderado", "Progreso", "Total actividades", "Completadas", "En curso",
            "Pendientes", "Vencidas", "Con alertas", "Estado general", "Mensaje de finalización"]
    cols = ["B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"]
    for h, col in zip(hdrs, cols):
        c = ws[f"{col}{r + 1}"]
        c.value = h; c.font = font(10, True, "FFFFFF"); c.fill = fill(MID); c.alignment = CENTER; c.border = BORDER
    ws.merge_cells(f"L{r + 1}:M{r + 1}")
    ws[f"M{r + 1}"].border = BORDER
    ws.row_dimensions[r + 1].height = 30
    for i in range(4):
        rr = r + 2 + i
        k = 4 + i
        ws[f"B{rr}"] = f"={C}A{k}"
        ws[f"C{rr}"] = (f'=IF({C}B{k}=0,"Sin actividades",IF(NOT({C}E{k}),"Revisar pesos",{C}F{k}))')
        ws[f"D{rr}"] = f"=IF(ISNUMBER(C{rr}),C{rr},0)"
        ws[f"E{rr}"] = f"={C}B{k}"
        ws[f"F{rr}"] = f"={C}N{k}"
        ws[f"G{rr}"] = f"={C}O{k}"
        ws[f"H{rr}"] = f"={C}P{k}"
        ws[f"I{rr}"] = f"={C}Q{k}"
        ws[f"J{rr}"] = f"={C}R{k}"
        ws[f"K{rr}"] = (f'=IF({C}B{k}=0,"Sin actividades",IF(NOT({C}E{k}),"Pesos suman "&TEXT({C}C{k},"0%")'
                        f'&IF({C}D{k}>0," / "&{C}D{k}&" vacío(s)",""),'
                        f'IF({C}L{k},"Completado",IF({C}F{k}>0,"En curso","Pendiente"))))')
        ws[f"L{rr}"] = (f'=IF({C}B{k}=0,"⚠ Registrar actividades",IF(NOT({C}E{k}),"⚠ Configurar ponderación al 100%",'
                        f'IF({C}M{k},"¡Enhorabuena, completado anticipadamente!",IF({C}L{k},"Completado","En ejecución"))))')
        ws.merge_cells(f"L{rr}:M{rr}")
        for col in "BCDEFGHIJKLM":
            c = ws[f"{col}{rr}"]
            c.border = BORDER
            c.font = font(10, col == "B", TXT)
            c.alignment = LEFT if col == "B" else CENTER
        ws[f"C{rr}"].number_format = "0%"
        ws[f"D{rr}"].number_format = '0%;;""'
        ws[f"D{rr}"].font = font(9, color="FFFFFF")
        ws.row_dimensions[rr].height = 28
    rt = r + 6
    ws[f"B{rt}"] = "Total proyecto"
    ws[f"C{rt}"] = f'=IF({C}$E$8,{C}$F$8,"No definitivo")'
    ws[f"D{rt}"] = f"=IF(ISNUMBER(C{rt}),C{rt},0)"
    for col, src in zip("EFGHIJ", ["B8", "N8", "O8", "P8", "Q8", "R8"]):
        ws[f"{col}{rt}"] = f"={C}{src}"
    ws[f"E{rt}"] = f"={C}B11"
    ws[f"K{rt}"] = f'=IF({C}$E$8,"Promedio 4 puestos","No definitivo")'
    ws[f"L{rt}"] = (f'=IF(NOT({C}$E$8),"⚠ Completar configuración de todos los puestos",'
                    f'IF({C}$M$8,"¡Enhorabuena, completado anticipadamente!",IF({C}$L$8,"Completado","En ejecución")))')
    ws.merge_cells(f"L{rt}:M{rt}")
    for col in "BCDEFGHIJKLM":
        c = ws[f"{col}{rt}"]
        c.border = BORDER; c.font = font(10, True, NAVY); c.fill = fill(LIGHT)
        c.alignment = LEFT if col == "B" else CENTER
    ws[f"C{rt}"].number_format = "0%"
    ws[f"D{rt}"].number_format = '0%;;""'
    ws[f"D{rt}"].font = font(9, color=LIGHT)
    ws.row_dimensions[rt].height = 28

    a, b = r + 2, rt
    cf.add(f"D{a}:D{b}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                     color="2E75B6", showValue=False))
    cf.add(f"C{a}:C{b}", FormulaRule(formula=[f'AND(ISNUMBER(C{a}),C{a}>=1)'], fill=fill(GREEN_F), font=Font(color=GREEN_T, bold=True)))
    cf.add(f"C{a}:C{b}", FormulaRule(formula=[f'AND(ISNUMBER(C{a}),C{a}>0,C{a}<1)'], fill=fill(BLUE_F), font=Font(color=BLUE_T, bold=True)))
    cf.add(f"C{a}:C{b}", FormulaRule(formula=[f'OR(NOT(ISNUMBER(C{a})),C{a}=0)'], fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    cf.add(f"K{a}:K{b}", FormulaRule(formula=[f'K{a}="Completado"'], fill=fill(GREEN_F), font=Font(color=GREEN_T, bold=True)))
    cf.add(f"K{a}:K{b}", FormulaRule(formula=[f'K{a}="En curso"'], fill=fill(BLUE_F), font=Font(color=BLUE_T, bold=True)))
    cf.add(f"K{a}:K{b}", FormulaRule(formula=[f'OR(K{a}="Sin actividades",LEFT(K{a},11)="Pesos suman",K{a}="No definitivo")'],
                                     fill=fill(AMBER_F), font=Font(color=AMBER_T, bold=True)))
    cf.add(f"K{a}:K{b}", FormulaRule(formula=[f'K{a}="Pendiente"'], fill=fill(RED_F), font=Font(color=RED_T)))
    cf.add(f"I{a}:J{b}", FormulaRule(formula=[f'I{a}>0'], fill=fill(ORANGE_F), font=Font(color="833C0B", bold=True)))
    msg_cf(cf, f"L{a}:M{b}", f"$L{a}")

    # --- Seguimiento semanal ---
    r = 22
    merge_set(ws, f"B{r}:M{r}", "Seguimiento semanal: avance acumulado ponderado e incremento semanal",
              font=font(12, True, "FFFFFF"), fill=fill(NAVY), alignment=LEFT)
    ws.row_dimensions[r].height = 24
    merge_set(ws, f"C{r + 1}:F{r + 1}", "Avance ACUMULADO ponderado al cierre de cada semana",
              font=font(9, True, "FFFFFF"), fill=fill(MID), alignment=CENTER, border=BORDER)
    merge_set(ws, f"H{r + 1}:K{r + 1}", "INCREMENTO respecto a la semana anterior (puntos %)",
              font=font(9, True, "FFFFFF"), fill=fill(MID), alignment=CENTER, border=BORDER)
    merge_set(ws, f"L{r + 1}:M{r + 1}", "Brecha vs meta 100%",
              font=font(9, True, "FFFFFF"), fill=fill(MID), alignment=CENTER, border=BORDER)
    hdr_r = r + 2
    ws[f"B{hdr_r}"] = "Puesto / Indicador"
    for w in range(4):
        ws[f"{'CDEF'[w]}{hdr_r}"] = f"Semana {w + 1}"
        ws[f"{'HIJK'[w]}{hdr_r}"] = f"Semana {w + 1}"
    ws[f"L{hdr_r}"] = "Pendiente para la meta"
    ws.merge_cells(f"L{hdr_r}:M{hdr_r}")
    for col in "BCDEFHIJKLM":
        c = ws[f"{col}{hdr_r}"]
        c.font = font(9, True, NAVY); c.fill = fill(LIGHT); c.alignment = CENTER; c.border = BORDER
    labels = PUESTOS + ["Avance general del proyecto", "Plan de referencia (supuesto)", "Meta final"]
    for i, lbl in enumerate(labels):
        rr = hdr_r + 1 + i
        k = 4 + i
        if i < 4:
            ws[f"B{rr}"] = f"={C}A{k}"
        else:
            ws[f"B{rr}"] = lbl
        for w in range(4):
            ccol = "CDEF"[w]; icol = "HIJK"[w]; kcol = "GHIJ"[w]
            if i < 4:
                ws[f"{ccol}{rr}"] = f'=IF(NOT({C}{kcol}$10),"",IF({C}$E{k},{C}{kcol}{k},"–"))'
            elif i == 4:
                ws[f"{ccol}{rr}"] = f'=IF(NOT({C}{kcol}$10),"",IF({C}$E$8,{C}{kcol}8,"n/d"))'
            elif i == 5:
                ws[f"{ccol}{rr}"] = f"=INDEX(PlanRef,{w + 1})"
            else:
                ws[f"{ccol}{rr}"] = "=MetaFinal"
            if i <= 4:
                if w == 0:
                    ws[f"{icol}{rr}"] = f'=IF(ISNUMBER(C{rr}),C{rr},"")'
                else:
                    prev = "CDEF"[w - 1]
                    ws[f"{icol}{rr}"] = f'=IF(AND(ISNUMBER({ccol}{rr}),ISNUMBER({prev}{rr})),{ccol}{rr}-{prev}{rr},"")'
            ws[f"{ccol}{rr}"].number_format = "0%"
            ws[f"{icol}{rr}"].number_format = '+0%;-0%;0%'
        if i <= 4:
            ws[f"L{rr}"] = (f'=IFERROR(MetaFinal-LOOKUP(2,1/ISNUMBER(C{rr}:F{rr}),C{rr}:F{rr}),"")')
            ws[f"L{rr}"].number_format = "0%"
        ws.merge_cells(f"L{rr}:M{rr}")
        for col in "BCDEFHIJKLM":
            c = ws[f"{col}{rr}"]
            c.border = BORDER
            c.alignment = LEFT if col == "B" else CENTER
            c.font = font(10, i >= 4, NAVY if i >= 4 else TXT)
            if i >= 4:
                c.fill = fill(LIGHT if i == 4 else "FFFFFF")
        if i >= 5:
            for col in "BCDEF":
                ws[f"{col}{rr}"].font = font(9, False, "404040", italic=(i == 5))
    first_w, last_w = hdr_r + 1, hdr_r + 5
    cf.add(f"C{first_w}:F{last_w}", FormulaRule(formula=[f'AND(ISNUMBER(C{first_w}),C{first_w}>=1)'],
                                               fill=fill(GREEN_F), font=Font(color=GREEN_T, bold=True)))
    cf.add(f"C{first_w}:F{last_w}", FormulaRule(formula=[f'AND(ISNUMBER(C{first_w}),C{first_w}<INDEX(PlanRef,COLUMN(C{first_w})-2))'],
                                               fill=fill(AMBER_F), font=Font(color=AMBER_T)))
    cf.add(f"H{first_w}:K{last_w}", FormulaRule(formula=[f'AND(ISNUMBER(H{first_w}),H{first_w}<0)'],
                                               fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    note_r = hdr_r + 9
    merge_set(ws, f"B{note_r}:M{note_r}",
              "Lectura: el acumulado NO se suma entre semanas; el incremento es la diferencia entre semanas consecutivas. "
              "«–» = puesto sin configuración válida; «n/d» = avance global no definitivo; celda vacía = semana aún no alcanzada. "
              "Ámbar = acumulado por debajo del plan de referencia.",
              font=font(8, italic=True, color="404040"), alignment=LEFT)
    ws.row_dimensions[note_r].height = 26

    # --- Gráficos ---
    ch_r = note_r + 2
    bar = BarChart()
    bar.type = "bar"
    bar.style = 10
    bar.title = "Avance ponderado por puesto"
    bar.add_data(Reference(ws, min_col=4, min_row=16, max_row=19), titles_from_data=False)
    bar.set_categories(Reference(ws, min_col=2, min_row=16, max_row=19))
    bar.y_axis.scaling.min = 0; bar.y_axis.scaling.max = 1
    bar.y_axis.number_format = "0%"; bar.y_axis.majorUnit = 0.25
    bar.x_axis.scaling.orientation = "maxMin"
    bar.legend = None
    bar.series[0].graphicalProperties.solidFill = MID
    bar.series[0].graphicalProperties.line.solidFill = MID
    bar.dataLabels = DataLabelList(); bar.dataLabels.showVal = True; bar.dataLabels.numFmt = "0%"
    bar.dataLabels.showSerName = False; bar.dataLabels.showCatName = False
    bar.dataLabels.showLegendKey = False; bar.dataLabels.showPercent = False
    bar.x_axis.delete = False; bar.y_axis.delete = False
    bar.height = 7.2; bar.width = 15.5
    ws.add_chart(bar, f"B{ch_r}")

    col = BarChart()
    col.type = "col"
    col.title = "Avance general acumulado vs plan de referencia"
    gen_row = hdr_r + 5
    col.add_data(Reference(ws, min_col=2, max_col=6, min_row=gen_row), titles_from_data=True, from_rows=True)
    col.add_data(Reference(ws, min_col=2, max_col=6, min_row=gen_row + 1), titles_from_data=True, from_rows=True)
    col.set_categories(Reference(ws, min_col=3, max_col=6, min_row=hdr_r))
    col.y_axis.scaling.min = 0; col.y_axis.scaling.max = 1
    col.y_axis.number_format = "0%"; col.y_axis.majorUnit = 0.25
    col.series[0].graphicalProperties.solidFill = NAVY
    col.series[1].graphicalProperties.solidFill = "BFBFBF"
    col.legend.position = "b"
    col.x_axis.delete = False; col.y_axis.delete = False
    col.height = 7.2; col.width = 17.5
    ws.add_chart(col, f"H{ch_r}")

    # --- Control de calidad ---
    r = ch_r + 16
    merge_set(ws, f"B{r}:M{r}", "Control de calidad de datos",
              font=font(12, True, "FFFFFF"), fill=fill(NAVY), alignment=LEFT)
    ws.row_dimensions[r].height = 24
    hr = r + 1
    merge_set(ws, f"B{hr}:F{hr}", "Verificación", font=font(10, True, "FFFFFF"), fill=fill(MID), alignment=CENTER, border=BORDER)
    ws[f"G{hr}"] = "Casos"
    ws[f"H{hr}"] = "Resultado"
    for cc in "GH":
        c = ws[f"{cc}{hr}"]; c.font = font(10, True, "FFFFFF"); c.fill = fill(MID); c.alignment = CENTER; c.border = BORDER
    merge_set(ws, f"I{hr}:M{hr}", "Acción sugerida", font=font(10, True, "FFFFFF"), fill=fill(MID), alignment=CENTER, border=BORDER)
    qc = [
        ("Puestos sin actividades registradas", f"=COUNTIF({C}B4:B7,0)", "crit",
         "Registrar al menos una macroactividad por puesto."),
        ("Puestos cuya ponderación no suma 100% (o con pesos vacíos)", f"=COUNTIFS({C}B4:B7,\">0\",{C}E4:E7,FALSE)", "crit",
         "Ajustar los pesos del puesto hasta sumar 100%."),
        ("Actividades con peso vacío o fuera de rango", count_msgs("Peso vacío", "Peso fuera de rango"), "crit",
         "Ingresar un peso entre 0% y 100% (mayor que 0)."),
        ("Actividades sin responsable", count_msgs("Sin responsable"), "crit", "Asignar responsable."),
        ("Actividades sin fecha límite", count_msgs("Sin fecha límite"), "crit", "Registrar la fecha límite."),
        ("Fechas incoherentes o no válidas (incluye sin fecha de inicio)",
         count_msgs("Inicio posterior", "Fecha no válida", "Sin fecha de inicio"), "crit",
         "La fecha de inicio debe ser anterior o igual a la fecha límite."),
        ("Porcentajes de avance > 100%, < 0% o no numéricos", count_msgs("Avance fuera de rango"), "crit",
         "Corregir: los avances deben estar entre 0% y 100%."),
        ("Avance semanal menor que el de una semana anterior", count_msgs("decreciente"), "crit",
         "Revisar: el avance es acumulado y no debería disminuir."),
        ("Actividades vencidas no completadas", f"={C}Q8", "crit",
         "Actualizar avance o reprogramar con justificación en Observaciones."),
        ("Actividades próximas a vencer (no completadas)", f"={C}B13", "warn",
         "Alerta preventiva: priorizar seguimiento."),
        ("Datos obligatorios incompletos (ID, área o macroactividad)",
         count_msgs("Falta ID", "Área sin seleccionar", "Área no válida", "Macroactividad vacía"), "crit",
         "Completar los campos obligatorios."),
        ("ID de actividad duplicados", count_msgs("ID duplicado"), "crit", "Asignar un ID único a cada actividad."),
        ("Actividades completadas sin evidencia (si es obligatoria)", count_msgs("sin evidencia"), "crit",
         "Registrar evidencia o enlace en Evidencia / Observaciones."),
        ("Avance global no definitivo (puestos no válidos)", f"=4-{C}C8", "crit",
         "El avance general solo es definitivo con los 4 puestos válidos."),
    ]
    for i, (lbl, f, kind, act) in enumerate(qc):
        rr = hr + 1 + i
        merge_set(ws, f"B{rr}:F{rr}", lbl, font=font(10), alignment=LEFT, border=BORDER)
        ws[f"G{rr}"] = f
        ws[f"G{rr}"].number_format = "0"
        ws[f"G{rr}"].font = font(10, True); ws[f"G{rr}"].alignment = CENTER; ws[f"G{rr}"].border = BORDER
        ws[f"H{rr}"] = f'=IF(G{rr}=0,"✔ OK",IF("{kind}"="warn","● Atención","⚠ Revisar"))'
        ws[f"H{rr}"].font = font(10, True); ws[f"H{rr}"].alignment = CENTER; ws[f"H{rr}"].border = BORDER
        merge_set(ws, f"I{rr}:M{rr}", act, font=font(9), alignment=LEFT, border=BORDER)
        ws.row_dimensions[rr].height = 20
    q1, q2 = hr + 1, hr + len(qc)
    cf.add(f"H{q1}:H{q2}", FormulaRule(formula=[f'LEFT(H{q1},1)="✔"'], fill=fill(GREEN_F), font=Font(color=GREEN_T, bold=True)))
    cf.add(f"H{q1}:H{q2}", FormulaRule(formula=[f'LEFT(H{q1},1)="●"'], fill=fill(AMBER_F), font=Font(color=AMBER_T, bold=True)))
    cf.add(f"H{q1}:H{q2}", FormulaRule(formula=[f'LEFT(H{q1},1)="⚠"'], fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    note = q2 + 1
    merge_set(ws, f"B{note}:M{note}",
              "Detalle por actividad: columna «Validación de datos» de la hoja Cronograma y Actividades (filtrar por «⚠»). "
              "Pendiente normal = sin avance y dentro de plazo; Vencida = fecha límite anterior a la fecha de corte sin llegar a 100%.",
              font=font(8, italic=True, color="404040"), alignment=LEFT)
    ws.row_dimensions[note].height = 26

    # --- Hitos ---
    r = note + 2
    merge_set(ws, f"B{r}:M{r}", "Hitos por puesto y semana (completados / programados)",
              font=font(12, True, "FFFFFF"), fill=fill(NAVY), alignment=LEFT)
    ws.row_dimensions[r].height = 24
    hr = r + 1
    heads = {"B": "Puesto", "C": "Semana 1", "D": "Semana 2", "E": "Semana 3", "F": "Semana 4",
             "G": "Total", "H": "Vencidos"}
    for colh, t in heads.items():
        c = ws[f"{colh}{hr}"]; c.value = t; c.font = font(10, True, "FFFFFF"); c.fill = fill(MID)
        c.alignment = CENTER; c.border = BORDER
    merge_set(ws, f"I{hr}:M{hr}", "Lectura", font=font(10, True, "FFFFFF"), fill=fill(MID), alignment=CENTER, border=BORDER)
    for i in range(4):
        rr = hr + 1 + i
        ws[f"B{rr}"] = f"={C}A{4 + i}"
        for w in range(4):
            colw = "CDEF"[w]
            ws[f"{colw}{rr}"] = (f'=COUNTIFS(hArea,$B{rr},hSemana,"Semana {w + 1}",hEstado,"Completado")&" / "'
                                 f'&COUNTIFS(hArea,$B{rr},hSemana,"Semana {w + 1}")')
        ws[f"G{rr}"] = f'=COUNTIFS(hArea,$B{rr},hEstado,"Completado")&" / "&COUNTIFS(hArea,$B{rr})'
        ws[f"H{rr}"] = (f'=COUNTIFS(hArea,$B{rr},hEstado,"<>Completado",hFechaObj,"<"&FechaRef)')
        merge_set(ws, f"I{rr}:M{rr}",
                  f'=IF(H{rr}>0,"⚠ "&H{rr}&" hito(s) con fecha objetivo vencida sin completar","Sin hitos vencidos")',
                  font=font(9), alignment=LEFT, border=BORDER)
        for colw in "BCDEFGH":
            c = ws[f"{colw}{rr}"]; c.border = BORDER; c.font = font(10, colw == "B")
            c.alignment = LEFT if colw == "B" else CENTER
    h1, h2 = hr + 1, hr + 4
    cf.add(f"H{h1}:H{h2}", FormulaRule(formula=[f'H{h1}>0'], fill=fill(ORANGE_F), font=Font(color="833C0B", bold=True)))
    cf.add(f"I{h1}:M{h2}", FormulaRule(formula=[f'LEFT($I{h1},1)="⚠"'], fill=fill(AMBER_F), font=Font(color=AMBER_T, bold=True)))

    ws.freeze_panes = "A3"
    ws.protection.sheet = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:N{h2 + 1}"
    ws.print_title_rows = "1:2"
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.row_breaks.append(__import__("openpyxl").worksheet.pagebreak.Break(id=note_r + 1))
    ws.oddFooter.center.text = "Página &P de &N"


def msg_cf(cf, rng, first):
    cf.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("Enhorabuena",{first}))'], stopIfTrue=True,
                            fill=fill(DKGREEN), font=Font(color="FFFFFF", bold=True)))
    cf.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("ompletado",{first}))'], stopIfTrue=True,
                            fill=fill(GREEN_F), font=Font(color=GREEN_T, bold=True)))
    cf.add(rng, FormulaRule(formula=[f'LEFT({first},1)="⚠"'], stopIfTrue=True,
                            fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    cf.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("ejecución",{first}))'],
                            fill=fill(BLUE_F), font=Font(color=BLUE_T, bold=True)))


# --------------------------------------------------------------------------
def build_instrucciones(ws):
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 120
    merge_set(ws, "B1:B1", "Instrucciones de uso, colaboración y protección",
              font=font(14, True, "FFFFFF"), fill=fill(NAVY), alignment=LEFT)
    ws.row_dimensions[1].height = 28
    sections = [
        ("1. Propósito", [
            "Administrar el ciclo de cuatro semanas de los puestos Riesgo Operacional, Riesgo Legal y Normativo, Riesgo Financiero y Asistente Técnico de Riesgo: macroactividades, responsables, fechas, pesos, avances semanales, estados, evidencias e hitos. Meta: 100% al cierre de la Semana 4.",
        ]),
        ("2. Estructura del libro", [
            "• Dashboard - Consolidado: indicadores, avance por puesto, seguimiento semanal, control de calidad e hitos. Solo lectura (todo es fórmula).",
            "• Cronograma y Actividades: base de datos principal (tabla tblActividades). Aquí se registran los datos.",
            "• Linea de Tiempo - Hitos: hitos y entregables por semana (tabla tblHitos). Contiene una PROPUESTA INICIAL EDITABLE de 16 hitos, sujeta a validación de los responsables.",
            "• Configuración: semana actual, fecha de corte, fecha de inicio del ciclo, días de alerta, evidencia obligatoria, plan de referencia, listas desplegables y metodología.",
            "• Calculos (oculta): cálculos intermedios por puesto. Para verla: clic derecho en una pestaña > Mostrar. No editar.",
        ]),
        ("3. Paso a paso", [
            "1) En Configuración, verifique la Fecha de inicio del ciclo (valor inicial editable) y seleccione la Semana actual (1 a 4).",
            "2) En Cronograma y Actividades registre cada macroactividad en una fila: ID único (ej. RO-01), Área / Puesto (lista), Macro Actividad, Responsable, Fecha de inicio, Fecha límite y Peso (%).",
            "3) Los pesos son relativos dentro de cada puesto y deben sumar 100% por puesto (ej. 40% + 35% + 25%). Si aún no hay pesos aprobados, deje el campo vacío: el libro mostrará una alerta y no calculará el avance del puesto.",
            "4) Al cierre de cada semana registre el avance ACUMULADO de la actividad en la columna de esa semana (ej. S1 20%, S2 45%, S3 75%, S4 100%). No registre incrementos. Deje vacías las semanas futuras.",
            "5) Actualice «Última actualización» y registre la evidencia o enlace (SharePoint/OneDrive) en «Evidencia / Observaciones». Las incidencias excepcionales (bloqueo, cancelación, reprogramación) se anotan en Observaciones: no alteran el estado calculado.",
            "6) Revise la columna «Validación de datos»: «✔ OK» o «⚠» con la lista de problemas. Use el filtro de la columna para ver solo las filas con «⚠».",
            "7) Consulte el Dashboard. Si un puesto o el proyecto llega al 100% antes de la Semana 4 aparece «¡Enhorabuena, completado anticipadamente!».",
            "8) Actualice el Estado, Fecha real y Evidencia de los hitos en Linea de Tiempo - Hitos.",
        ]),
        ("4. Código de colores", [
            "• Amarillo pálido: celda de entrada (desbloqueada). • Gris: fórmula protegida. • Azul oscuro/medio: encabezados e indicadores.",
            "• Verde: 100% / Completado. • Azul claro: En curso. • Rojo suave: Pendiente, dato faltante o alerta crítica. • Ámbar: próxima a vencer, pesos ≠ 100% o por debajo del plan. • Naranja: vencida, ID duplicado, fechas incoherentes o avance decreciente. • Verde oscuro con texto blanco: finalización anticipada.",
        ]),
        ("5. Agregar filas", [
            "La tabla tblActividades trae 100 filas preformateadas (fórmulas, validaciones y formato). Las filas vacías no se cuentan.",
            "Para agregar más filas: Revisar > Desproteger hoja; ubíquese en la última celda de la tabla y presione Tab (o escriba justo debajo de la tabla). Excel extiende automáticamente las fórmulas de las columnas calculadas (% de Avance Total Actual, Estado y Validación). Vuelva a proteger la hoja al terminar. El Dashboard lee hasta la fila 1000.",
            "No elimine ni mueva las columnas de la tabla; no escriba sobre las columnas grises.",
        ]),
        ("6. Compartir y trabajar en colaboración (Excel para la Web / OneDrive / SharePoint)", [
            "1) Guarde el archivo en una biblioteca de SharePoint o en OneDrive para el trabajo (idealmente el sitio del área de Riesgos).",
            "2) Archivo > Compartir > Compartir con personas específicas y otorgue permiso de edición a los cuatro responsables; permiso de solo lectura a la jefatura que solo consulta el Dashboard.",
            "3) Varias personas pueden editar al mismo tiempo (coautoría) en Excel para la Web o Microsoft 365. Use Autoguardado activado. Las fórmulas se recalculan automáticamente.",
            "4) Use «Vistas de hoja» (Vista > Vista de hoja > Nueva) para que cada persona filtre su puesto sin alterar el filtro de los demás.",
            "5) Historial de versiones (Archivo > Información > Historial de versiones) permite recuperar versiones anteriores y ver quién cambió qué. También puede usar Revisar > Mostrar cambios.",
            "6) Las evidencias deben almacenarse en la biblioteca de SharePoint y pegarse como enlace en la columna de evidencia.",
        ]),
        ("7. Protección de fórmulas", [
            "Todas las hojas están protegidas SIN contraseña: las celdas con fórmulas están bloqueadas y las celdas de entrada (amarillas) desbloqueadas. Se permiten filtros, ordenar y ajustar anchos.",
            "Para establecer una contraseña: Revisar > Desproteger hoja, luego Revisar > Proteger hoja, marque las opciones «Seleccionar celdas desbloqueadas», «Usar Autofiltro» y «Ordenar» y escriba la contraseña. Guárdela en un lugar seguro. Opcional: Revisar > Proteger libro (estructura) para impedir mostrar/eliminar hojas.",
            "LIMITACIÓN: la protección de hoja evita sobrescribir fórmulas, pero NO impide que un usuario con permiso de edición modifique los registros de otro puesto dentro de las celdas de entrada. Excel para la Web tampoco admite «Permitir que los usuarios editen rangos» con permisos por usuario. La trazabilidad se apoya en el historial de versiones.",
        ]),
        ("8. Segregación real de permisos (si se requiere)", [
            "1) Cree cuatro archivos de captura idénticos (uno por puesto) con la misma estructura de tblActividades, cada uno en una carpeta de SharePoint con permisos solo para su responsable.",
            "2) En este archivo consolidado: Datos > Obtener datos > Desde archivo > Desde SharePoint (o Desde carpeta), seleccione la carpeta con los cuatro archivos y combine la tabla tblActividades de cada uno (Combinar y transformar).",
            "3) Cargue el resultado como tabla en la hoja Cronograma y Actividades (o reemplace tblActividades por la tabla de la consulta, con los mismos encabezados). Las columnas calculadas pueden recalcularse en la consulta o en la tabla.",
            "4) Configure Datos > Consultas y conexiones > Propiedades > «Actualizar al abrir el archivo». Nota: la actualización de Power Query desde orígenes de SharePoint se ejecuta en Excel de escritorio; en Excel para la Web el soporte depende del tipo de origen y del plan de Microsoft 365.",
            "5) Solo el coordinador tiene permiso de edición del consolidado; los demás, lectura.",
        ]),
        ("9. Compatibilidad técnica", [
            "Diseñado para Microsoft 365 en español (escritorio y Web). Archivo .xlsx sin macros ni complementos.",
            "Las fórmulas se guardan internamente en inglés y Excel en español las muestra traducidas automáticamente: LOOKUP = BUSCAR, IFERROR = SI.ERROR, IF = SI, SUMPRODUCT = SUMAPRODUCTO, COUNTIFS = CONTAR.SI.CONJUNTO, SUMIFS = SUMAR.SI.CONJUNTO, ISNUMBER = ESNUMERO, SEARCH = HALLAR, TODAY = HOY. El separador de argumentos también se adapta (; en configuración regional española).",
            "Fórmula del avance actual (fila 5, tal como se ve en Excel en español): =SI(NO((CONTARA($A5:$K5)+CONTARA($N5:$O5))>0);\"\";SI.ERROR(BUSCAR(2;1/(H5:K5<>\"\");H5:K5);0)). La primera parte deja la celda vacía en filas sin datos.",
            "Fórmula del estado: =SI(L5=\"\";\"\";SI(L5>=1;\"Completado\";SI(L5>0;\"En curso\";\"Pendiente\"))).",
        ]),
        ("10. Limitaciones conocidas", [
            "• El estado es automático (basado en el %); los estados excepcionales se documentan en Observaciones para no alterar el cálculo.",
            "• La regla de finalización anticipada usa la semana en que el puesto alcanzó el 100% según los avances registrados, acotada por la Semana actual de Configuración. Si se registran avances en una columna de semana futura, el mensaje se basa en la Semana actual.",
            "• Los rangos del Dashboard llegan hasta la fila 1000 de la tabla de actividades y hasta la fila 500 de hitos.",
            "• Agregar filas a la tabla requiere desproteger temporalmente la hoja (restricción de Excel para tablas en hojas protegidas).",
            "• Los hitos y fechas objetivo son propuestas iniciales; no son obligaciones institucionales aprobadas.",
        ]),
    ]
    r = 3
    for title, lines in sections:
        c = ws.cell(row=r, column=2, value=title)
        c.font = font(11, True, "FFFFFF"); c.fill = fill(MID); c.alignment = LEFT
        ws.row_dimensions[r].height = 20
        r += 1
        for ln in lines:
            c = ws.cell(row=r, column=2, value=ln)
            c.font = font(10); c.alignment = LEFT_TOP
            ws.row_dimensions[r].height = max(15, 15 * (len(ln) // 115 + 1))
            r += 1
        r += 1
    ws.protection.sheet = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "Cronograma_Dashboard_Riesgos.xlsx"
    build(out)
    print("Generado:", out)
